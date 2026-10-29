"""Тесты md2docx по SPEC.md §6 — конвертер как чёрный ящик: CLI `python3 md2docx.py …`
из корня проекта и функция convert() из README (в отдельном процессе). Сверка — независимым
tests/verify.py.

    python3 -m pytest tests -q --basetemp=<свой каталог>/pytest-tmp -p no:cacheprovider

Необязательные части (пропускаются, если нет данных или инструментов):
  * свои реальные документы для регрессии — пути в tests/real_docs.local.txt (файл не в git,
    результат конвертации удаляется сразу после сверки);
  * эталон стиля (docx, с которого снято оформление) для красной пробы сверщика — путь
    в MD2DOCX_REFERENCE_DOCX; известные отклонения эталона — в REFERENCE_DEVIATIONS;
  * рендер LibreOffice → PDF (soffice; poppler: pdfinfo, pdftotext, pdffonts, pdftoppm;
    Pillow): MD2DOCX_SOFFICE — путь к soffice, MD2DOCX_LO_PROFILE — каталог профиля
    LibreOffice (по умолчанию временный).
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from pathlib import Path

import pytest
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
TESTS = ROOT / "tests"
MD2DOCX = ROOT / "md2docx.py"
VERIFY = TESTS / "verify.py"
SAMPLES = ROOT / "samples"
REAL_DOCS_LIST = TESTS / "real_docs.local.txt"

sys.path.insert(0, str(TESTS))
import verify as V  # noqa: E402  (свой сверщик, не md2docx)

q, wa, NS = V.q, V.wa, V.NS

# Ожидаемая ориентация разделов по образцам (P — книжный, L — альбомный).
EXPECT_ORIENT = {
    "01-headings": "P",
    "02-tables": "P",
    "03-orientation-markers": "P,L,P",
    "04-auto-landscape": "L,P",
    "05-lists": "P",
    "06-blocks": "P",
    "07-pagebreak": "P,L",
    "08-showcase": "P,L,P",
}
SAMPLE_FILES = sorted(SAMPLES.glob("*.md"))

# --------------------------------------------------------------------------
# Запуск конвертера и сверщика
# --------------------------------------------------------------------------


def _require_impl():
    if not MD2DOCX.is_file():
        pytest.fail(f"нет {MD2DOCX.name} в корне проекта — реализации ещё нет, тест ждёт её",
                    pytrace=False)


def cli(*args, cwd=ROOT, timeout=300) -> subprocess.CompletedProcess:
    _require_impl()
    return subprocess.run([sys.executable, str(MD2DOCX), *map(str, args)], cwd=cwd,
                          capture_output=True, text=True, timeout=timeout)


def run_verify(*args) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(VERIFY), *map(str, args)], cwd=ROOT,
                          capture_output=True, text=True, timeout=300)


_CACHE: dict = {}


@pytest.fixture(scope="session")
def outdir(tmp_path_factory):
    return tmp_path_factory.mktemp("md2docx-out")


def convert(outdir: Path, md: Path, *flags: str, stem: str | None = None):
    """Конвертирует md → docx через CLI (с кэшем на сессию). Возвращает (docx, proc)."""
    key = (str(md), flags)
    if key in _CACHE:
        return _CACHE[key]
    suffix = "".join("_" + f.strip("-") for f in flags)
    out = outdir / f"{stem or md.stem}{suffix}.docx"
    proc = cli(md, "-o", out, *flags)
    assert proc.returncode == 0, (f"md2docx упал на {md.name} (код {proc.returncode})\n"
                                  f"stdout: {proc.stdout}\nstderr: {proc.stderr[-3000:]}")
    assert out.is_file(), f"md2docx вернул 0, но {out.name} не создан"
    _CACHE[key] = (out, proc)
    return out, proc


def sample(name: str) -> Path:
    return SAMPLES / f"{name}.md"


def orient(dx: V.Docx) -> list:
    return [s.orient for s in dx.sections]


def paragraphs(dx: V.Docx):
    return list(dx.body.iter(q("p")))


def find_p(dx: V.Docx, needle: str):
    for p in dx.body.iter(q("p")):
        if needle in V.para_text(p):
            return p
    raise AssertionError(f"в docx нет абзаца с текстом «{needle}»")


def find_table(dx: V.Docx, needle: str) -> V.DTable:
    for t in dx.tables:
        if needle in V.cell_text(t.rows[0][0]) or any(
                needle in V.cell_text(tc) for row in t.rows for tc in row):
            return t
    raise AssertionError(f"в docx нет таблицы с текстом «{needle}»")


def run_with_text(p, needle: str):
    for r in p.iter(q("r")):
        if needle in V.para_text(r):
            return r
    raise AssertionError(f"в абзаце нет прогона с текстом «{needle}»")


def spacing(dx, p):
    return {a: V.to_int(dx.eff_ppr(p, "spacing", a)) for a in ("before", "after", "line")} | {
        "lineRule": dx.eff_ppr(p, "spacing", "lineRule")}


def write_md(tmp_path: Path, name: str, text: str) -> Path:
    p = tmp_path / f"{name}.md"
    p.write_text(text, encoding="utf-8")
    return p


def run_md(tmp_path: Path, name: str, text: str, *flags: str):
    """md из строки → docx через CLI, код 0 обязателен. → (md, docx, proc, V.Docx)."""
    md = write_md(tmp_path, name, text)
    out = tmp_path / f"{name}.docx"
    proc = cli(md, "-o", out, *flags)
    assert proc.returncode == 0, proc.stderr[-2000:]
    return md, out, proc, V.Docx(out)


def warnings_of(proc) -> list:
    return [x for x in proc.stderr.splitlines() if x.startswith("Внимание:")]


# --------------------------------------------------------------------------
# 1. Все образцы: CLI + сверщик
# --------------------------------------------------------------------------


def test_samples_present():
    names = {p.stem for p in SAMPLE_FILES}
    assert set(EXPECT_ORIENT) <= names, f"не хватает образцов: {set(EXPECT_ORIENT) - names}"


@pytest.mark.parametrize("md", SAMPLE_FILES, ids=[p.stem for p in SAMPLE_FILES])
def test_sample_cli_and_verify(md, outdir):
    docx, proc = convert(outdir, md)
    lines = [x for x in proc.stdout.splitlines() if x.strip()]
    assert len(lines) == 1 and lines[0].startswith("Готово:"), \
        f"на stdout ожидалась одна строка «Готово: …», получено: {proc.stdout!r}"
    assert docx.name in lines[0], f"в строке итога нет имени файла: {lines[0]}"
    dx = V.Docx(docx)
    m = re.search(r"разделов\s+(\d+)\s*\(([^)]*)\)", lines[0])
    if m:  # формат строки в SPEC дан примером; если он соблюдён — числа обязаны совпасть
        assert int(m.group(1)) == len(dx.sections), lines[0]
        letters = [x.strip() for x in m.group(2).split(",")]
        assert letters == ["К" if o == "P" else "А" for o in orient(dx)], lines[0]
    m = re.search(r"таблиц\s+(\d+)", lines[0])
    if m:
        assert int(m.group(1)) == len(dx.tables), lines[0]
    args = [md, docx]
    if md.stem in EXPECT_ORIENT:
        args += ["--expect-orient", EXPECT_ORIENT[md.stem]]
    res = run_verify(*args)
    assert res.returncode == 0, f"сверщик нашёл расхождения:\n{res.stdout}\n{res.stderr}"


# --------------------------------------------------------------------------
# 2. Ориентация: маркеры, флаг, автоальбомная
# --------------------------------------------------------------------------


def test_orientation_markers(outdir):
    docx, proc = convert(outdir, sample("03-orientation-markers"))
    dx = V.Docx(docx)
    assert orient(dx) == ["P", "L", "P"], "маркеры: книжная → альбомная → книжная"
    assert dx.section_index(find_p(dx, "Первая страница книжная")) == 0
    assert dx.section_index(find_p(dx, "Альбомный участок")) == 1
    assert dx.tables[0].section == 1, "таблица после <!-- landscape --> — в альбомном разделе"
    assert dx.section_index(find_p(dx, "Повторный маркер той же ориентации")) == 1, \
        "повторный маркер той же ориентации не должен открывать раздел"
    assert dx.section_index(find_p(dx, "Снова книжная")) == 2
    code = find_p(dx, "это строка блока кода")
    assert "<!-- landscape -->" in V.para_text(code), "маркер внутри блока кода — это текст кода"
    assert dx.section_index(code) == 2
    assert dx.section_index(find_p(dx, "Маркер без пробелов")) == 2, \
        "маркер в конце документа должен игнорироваться"


# не помещается на книжной странице без разрыва слов — уходит на альбомную (SPEC §4.3)
WIDE_TABLE = ("| " + " | ".join(f"Колонка{i} информационно-телекоммуникационной" for i in range(9))
              + " |\n" + "|---" * 9 + "|\n| " + " | ".join(["значение"] * 9) + " |\n")


@pytest.mark.parametrize("text, flags, expected, warn", [
    ("<!-- landscape -->\n\n# Старт\n\nТекст первой страницы.\n", (), ["L"], False),
    ("Текст до маркеров.\n\n<!-- landscape -->\n<!-- portrait -->\n<!-- landscape -->\n\n"
     "Текст после трёх маркеров.\n", (), ["P", "L"], False),
    ("Текст документа.\n\n<!-- landscape -->\n\n<!-- pagebreak -->\n", (), ["P"], False),
    ("Текст на альбомной.\n\n<!-- portrait -->\n\nТекст на книжной.\n", ("--landscape",),
     ["L", "P"], False),
    ("Текст на альбомной.\n\n<!-- landscape -->\n\nВсё ещё альбомная.\n", ("--landscape",),
     ["L"], False),
    ("- пункт один\n\n  <!-- landscape -->\n\n- пункт два\n\nАбзац после списка.\n", (),
     ["P"], True),
    ("> цитата\n>\n> <!-- landscape -->\n\nАбзац после цитаты.\n", (), ["P"], True),
    ("> цитата\n>\n> <!-- pagebreak -->\n\nАбзац после цитаты.\n", (), ["P"],
     "маркер разрыва страницы"),
    # блок без видимого содержимого маркеры не разделяет и раздела не открывает (SPEC §4.1)
    ("Текст до.\n\n" + WIDE_TABLE + "\n> <!-- заметка для себя -->\n", (), ["P", "L"], False),
    ("<!-- landscape -->\n\nАльбомный текст.\n\n<!-- portrait -->\n\n>\n\n<!-- landscape -->\n\n"
     "Снова альбомный.\n", (), ["L"], False),
    ("Текст.\n\n<!-- landscape -->\n\n<br>\n\n<!-- portrait -->\n\nЕщё текст.\n", (), ["P"], False),
    ("Текст.\n\n<!-- landscape -->\n\n#\n\n<!-- portrait -->\n\nЕщё текст.\n", (), ["P"], False),
    ("Текст.\n\n<!-- landscape -->\n\n---\n\n<!-- portrait -->\n\nЕщё текст.\n", (), ["P", "L", "P"],
     False),  # линия — видимое содержимое: остаётся за маркером, на своей странице
    ("Текст.\n\n\\newpage\n\n>\n\n<!-- landscape -->\n\nАльбомный текст.\n", (), ["P", "L"], False),
    ("Текст.\n\n<!-- land scape -->\n\nЕщё текст.\n", (), ["P"], False),
    ("Текст.\n\n<!-- незакрытый комментарий\n\nОстаток файла.\n", (), ["P"],
     "незакрытый HTML-комментарий"),
    ("", (), ["P"], "нет видимого содержимого"),
], ids=["marker-at-start", "markers-in-a-row", "markers-at-end", "flag-landscape",
        "flag-landscape-same-marker", "marker-in-list", "marker-in-quote", "pagebreak-in-quote",
        "invisible-block-at-end", "empty-quote-between-markers", "br-only-section",
        "empty-heading-section", "rule-only-section", "pagebreak-before-section", "spaced-marker",
        "unclosed-comment", "empty-document"])
def test_marker_edge_cases(tmp_path, text, flags, expected, warn):
    md = write_md(tmp_path, "edge", text)
    out = tmp_path / "edge.docx"
    proc = cli(md, "-o", out, *flags)
    assert proc.returncode == 0, proc.stderr[-2000:]
    dx = V.Docx(out)
    assert orient(dx) == expected
    if warn:
        assert "Внимание:" in proc.stderr, "маркер внутри списка/цитаты — с предупреждением"
        assert warn is True or warn in proc.stderr, proc.stderr
    rep = V.run_checks(md, out, expected, landscape="--landscape" in flags)
    assert rep.ok, rep.render()


def test_marker_new_page_next_to_auto_landscape(tmp_path):
    """Маркер даёт новую страницу и тогда, когда автоальбомная таблица за ним сливается
    с альбомным участком перед ним (SPEC §4.1)."""
    for tail, expected in (("", ["L"]), ("\nКнижный текст.\n", ["L", "P"])):
        md = write_md(tmp_path, "glue", "<!-- landscape -->\n\nАльбомный текст.\n\n"
                                        "<!-- portrait -->\n\n" + WIDE_TABLE + tail)
        out = tmp_path / "glue.docx"
        proc = cli(md, "-o", out)
        assert proc.returncode == 0, proc.stderr[-2000:]
        dx = V.Docx(out)
        assert orient(dx) == expected
        seq = []  # содержимое первого раздела: «c» — абзац, «b» — разрыв страницы, «t» — таблица
        for el in dx.sections[0].elements:
            seq += ["t"] if V.local(el) == "tbl" else V.para_atoms(el) if V.local(el) == "p" else []
        assert seq == ["c", "b", "t"], f"таблица — с новой страницы после альбомного текста: {seq}"
        rep = V.run_checks(md, out, expected)
        assert rep.ok, rep.render()


def test_auto_landscape_pull_and_glue(outdir):
    """Подтягивание (SPEC §4.3): подпись с двоеточием и заголовок над ней уходят с таблицей,
    а через короткий вводный абзац — и заголовок старшего уровня (здесь H1 документа);
    вторая широкая таблица со своим заголовком приклеивается через линию; узкая — книжная."""
    docx, proc = convert(outdir, sample("04-auto-landscape"))
    dx = V.Docx(docx)
    assert orient(dx) == ["L", "P"], \
        "начало документа уходит к широким таблицам, хвост возвращается на книжную"
    sec = dx.section_index
    for text in ("Автоальбомная ориентация", "Вводный абзац на книжной странице",
                 "Сравнение направлений", "Показатели по девяти направлениям:",
                 "Продолжение сравнения"):
        assert sec(find_p(dx, text)) == 0, f"«{text}» — в альбомном разделе вместе с таблицей"
    wide1, wide2 = find_table(dx, "Аудитория"), find_table(dx, "Монетизация")
    assert wide1.section == wide2.section == 0, "склейка: обе широкие таблицы в одном разделе"
    assert sec(find_p(dx, "Абзац после широких таблиц")) == 1
    narrow = find_table(dx, "остаётся")
    assert narrow.section == 1, "узкая таблица остаётся книжной"
    assert sec(find_p(dx, "Финальный абзац документа")) == 1
    for t in (wide1, wide2):
        assert sum(t.grid) == V.TEXT_WIDTH["L"], "широкая таблица — во всю ширину альбомной"
    m = re.search(r"авто-альбомных\s+(\d+)", proc.stdout)
    if m:
        assert int(m.group(1)) == 2, proc.stdout


def test_no_auto_landscape_flag(outdir):
    md = sample("04-auto-landscape")
    docx, proc = convert(outdir, md, "--no-auto-landscape")
    dx = V.Docx(docx)
    assert orient(dx) == ["P"], "с --no-auto-landscape всё остаётся книжным"
    assert "Внимание:" in proc.stderr, "таблица не влезла без разрыва слов — нужно предупреждение"
    for t in dx.tables:
        assert sum(t.grid) == V.TEXT_WIDTH["P"]
    res = run_verify(md, docx, "--no-auto-landscape", "--expect-orient", "P")
    assert res.returncode == 0, res.stdout


def test_landscape_flag(outdir):
    md = sample("01-headings")
    docx, _ = convert(outdir, md, "--landscape")
    dx = V.Docx(docx)
    assert orient(dx) == ["L"]
    res = run_verify(md, docx, "--landscape", "--expect-orient", "L")
    assert res.returncode == 0, res.stdout


def test_pagebreaks(outdir):
    docx, _ = convert(outdir, sample("07-pagebreak"))
    dx = V.Docx(docx)
    assert orient(dx) == ["P", "L"]
    breaks = []
    for p in paragraphs(dx):
        n = sum(1 for br in p.iter(q("br")) if wa(br, "type") == "page")
        n += 1 if V.is_on(p.find("w:pPr/w:pageBreakBefore", NS)) else 0
        breaks += [dx.section_index(p)] * n
    assert breaks == [0, 0, 0, 0], (f"ожидалось 4 разрыва страницы в первом разделе (маркер в "
                                    f"начале и в конце игнорируются, два подряд — один), "
                                    f"найдено по разделам: {breaks}")
    assert dx.section_index(find_p(dx, "Шестая страница альбомная")) == 1


def test_pagebreak_line_inside_paragraph(tmp_path):
    """Строка \\newpage / \\pagebreak посреди абзаца — тоже маркер (как <!-- pagebreak -->,
    который markdown-it выделяет из абзаца); регистр не важен; внутри списка — игнорируется
    с предупреждением; «… \\newpage …» внутри строки — обычный текст."""
    md = write_md(tmp_path, "inline-break",
                  "Текст до разрыва\n\\pagebreak\nТекст после разрыва.\n\n"
                  "Слово \\newpage в строке — не маркер.\n\n"
                  "- пункт списка\n  \\NewPage\n- второй пункт\n")
    out = tmp_path / "inline-break.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert "Внимание:" in proc.stderr, "маркер внутри списка — с предупреждением"
    seq = []  # тексты абзацев и разрывы страниц («|») в порядке документа
    for p in paragraphs(V.Docx(out)):
        for a in V.para_atoms(p):
            if a == "b":
                seq.append("|")
            elif not seq or seq[-1] != V.para_text(p):
                seq.append(V.para_text(p))
    assert seq == ["Текст до разрыва", "|", "Текст после разрыва.",
                   "Слово \\newpage в строке — не маркер.", "пункт списка", "второй пункт"], seq
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


# Правило подтягивания к автоальбомной таблице (SPEC §4.3): (текст, разделы, {абзац: раздел}).
# Невидимые блоки прозрачны; подводка (до двух коротких абзацев и одной линии) уходит только
# вместе с заголовком за ней; подъём — к заголовку строго старшего уровня; список, разрыв,
# третий абзац, абзац длиннее 600 знаков останавливают; склейка считает только видимые блоки
# и через разрыв страницы не идёт.
PULL_CASES = [
    ("Вступление.\n\n## Раздел\n\n<br>\n\n" + WIDE_TABLE, ["P", "L"],
     {"Вступление.": 0, "Раздел": 1}),
    ("Вступление.\n\n## Раздел\n\n> <!-- заметка -->\n\n" + WIDE_TABLE, ["P", "L"],
     {"Вступление.": 0, "Раздел": 1}),
    ("Текст до.\n\n## Раздел\n\nКороткое вступление без двоеточия.\n\n" + WIDE_TABLE, ["P", "L"],
     {"Текст до.": 0, "Раздел": 1, "Короткое вступление": 1}),
    ("Текст до.\n\n## Раздел\n\nПервый абзац.\n\n---\n\nВторой абзац.\n\n" + WIDE_TABLE,
     ["P", "L"], {"Текст до.": 0, "Раздел": 1, "Первый абзац.": 1, "Второй абзац.": 1}),
    ("## Раздел\n\nПервый.\n\nВторой.\n\nТретий.\n\n" + WIDE_TABLE, ["P", "L"],
     {"Раздел": 0, "Первый.": 0, "Третий.": 0}),
    ("## Раздел\n\n" + "длинное вступление " * 40 + "\n\n" + WIDE_TABLE, ["P", "L"],
     {"Раздел": 0, "длинное вступление": 0}),
    ("# Глава\n\nВступление главы.\n\n---\n\n## Раздел\n\n" + WIDE_TABLE, ["L"],
     {"Глава": 0, "Вступление главы.": 0, "Раздел": 0}),
    ("## Раздел А\n\nТекст А.\n\n## Раздел Б\n\n" + WIDE_TABLE, ["P", "L"],
     {"Раздел А": 0, "Текст А.": 0, "Раздел Б": 1}),
    ("# Глава\n\n- пункт списка\n\n## Раздел\n\n" + WIDE_TABLE, ["P", "L"],
     {"Глава": 0, "пункт списка": 0, "Раздел": 1}),
    ("# Глава\n\nВступление.\n\n\\newpage\n\n## Раздел\n\n" + WIDE_TABLE, ["P", "L"],
     {"Глава": 0, "Вступление.": 0, "Раздел": 1}),
    ("Текст до.\n\nТаблица показателей:\n\n" + WIDE_TABLE, ["P", "L"],
     {"Текст до.": 0, "Таблица показателей:": 1}),
    (WIDE_TABLE + "\nАбзац один.\n\n<br>\n\n> <!-- заметка -->\n\nАбзац два.\n\n" + WIDE_TABLE,
     ["L"], {"Абзац один.": 0, "Абзац два.": 0}),
    (WIDE_TABLE + "\nПервый.\n\nВторой.\n\nТретий.\n\n" + WIDE_TABLE, ["L", "P", "L"],
     {"Второй.": 1}),
    (WIDE_TABLE + "\nАбзац.\n\n\\newpage\n\n" + WIDE_TABLE, ["L", "P", "L"], {"Абзац.": 1}),
]
PULL_IDS = ["br-between-heading-and-table", "invisible-quote-between", "lead-in-without-colon",
            "two-paragraphs-and-rule", "third-paragraph-stops", "long-paragraph-stops",
            "chapter-intro-and-rule", "same-level-heading-stops", "list-stops", "pagebreak-stops",
            "caption-without-heading", "glue-counts-visible-only", "three-visible-no-glue",
            "pagebreak-stops-glue"]


@pytest.mark.parametrize("text, expected, where", PULL_CASES, ids=PULL_IDS)
def test_auto_landscape_pull_rule(tmp_path, text, expected, where):
    md, out, proc, dx = run_md(tmp_path, "pull", text)
    assert orient(dx) == expected
    for needle, k in where.items():
        assert dx.section_index(find_p(dx, needle)) == k, f"«{needle}» — в разделе {k + 1} из {expected}"
    rep = V.run_checks(md, out, expected)
    assert rep.ok, rep.render()


@pytest.mark.parametrize("line", ["\\newpage", "\\pagebreak"])
def test_pagebreak_right_under_table(tmp_path, line):
    """Строка-маркер разрыва сразу под строкой таблицы (без пустой строки) — разрыв страницы,
    а не строка таблицы, хотя по GFM она стала бы строкой (SPEC §4.1)."""
    md, out, proc, dx = run_md(tmp_path, "under-table",
                               f"| A | B |\n|---|---|\n| 1 | 2 |\n{line}\nПосле маркера.\n")
    rows = [len(t.rows) for t in dx.tables]
    assert rows == [2], f"строк в таблицах {rows} — маркер или текст за ним стал строкой таблицы"
    assert line[1:] not in V.docx_text(dx), "маркер не выводится текстом"
    seq = []  # атомы абзацев после таблицы до «После маркера.» включительно
    for el in dx.tables[0].el.itersiblings():
        if V.local(el) == "p":
            seq += V.para_atoms(el)
            if "После маркера." in V.para_text(el):
                break
    assert seq[:2] == ["b", "c"], f"абзац после маркера — с новой страницы: {seq}"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_pagebreaks_between_table_pieces(tmp_path):
    """Каждый \\newpage под строкой таблицы — разрыв, а не только первый: три куска таблицы
    через два \\newpage — три таблицы и два разрыва (SPEC §4.1)."""
    piece = "| a | b |\n|---|---|\n| {0} | {0} |\n"
    md, out, proc, dx = run_md(tmp_path, "pieces", piece.format(1) + "\\newpage\n" + piece.format(2)
                               + "\\newpage\n" + piece.format(3))
    rows = [len(t.rows) for t in dx.tables]
    assert rows == [2, 2, 2], f"строк в таблицах {rows} — второй \\newpage стал строкой таблицы"
    found = sum(V.para_atoms(p).count("b") for p in paragraphs(dx))
    assert found == 2, f"разрывов страницы {found}, ожидалось 2"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


SPLIT_TABLE = ("Абзац до таблицы.\n\n| Колонка | Значение |\n|---|---|\n| первая | 1 |\n| вторая | 2 |\n"
               "{brk}\n| третья | 3 |\n| четвёртая | 4 |\n\nАбзац после таблицы.\n")


def _first_col(t) -> list:
    return [V.cell_text(r[0]).strip() for r in t.rows]


def _first_row_flag(t):
    return wa(t.el.find("w:tblPr/w:tblLook", NS), "firstRow")


def _no_tbl_header(dx) -> bool:
    return not any(V.is_on(tr.find("w:trPr/w:tblHeader", NS)) for t in dx.tables
                   for tr in t.el.findall(q("tr")))


@pytest.mark.parametrize("line", ["\\newpage", "\\pagebreak"])
def test_newpage_inside_table_cuts_it(tmp_path, line):
    """\\newpage между строками тела режет таблицу по целой строке: вторая часть — отдельная
    таблица с новой страницы, без шапки, с той же сеткой; маркер текстом не выводится (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "cut", SPLIT_TABLE.format(brk=line))
    assert not warnings_of(proc), proc.stderr
    assert [_first_col(t) for t in dx.tables] == [["Колонка", "первая", "вторая"], ["третья", "четвёртая"]]
    t1, t2 = dx.tables
    assert t1.grid == t2.grid, f"сетка частей: {t1.grid} и {t2.grid}"
    assert wa(t1.el.find("w:tblPr/w:tblW", NS), "w") == wa(t2.el.find("w:tblPr/w:tblW", NS), "w")
    assert (_first_row_flag(t1), _first_row_flag(t2)) == ("1", "0"), "шапка — только у первой части"
    fills = {(wa(tc.find("w:tcPr/w:shd", NS), "fill") or "").upper() for tc in t2.el.iter(q("tc"))}
    assert "F2F2F2" not in fills, "у второй части нет заливки шапки"
    first = t2.rows[0][0].find(q("p"))
    assert not V.is_on(dx.eff_rpr(run_with_text(first, "третья"), first, "b")), "не полужирная"
    assert _no_tbl_header(dx), "шапка на страницах не повторяется"
    gap = t1.el.getnext()
    assert V.local(gap) == "p" and V.para_atoms(gap) == ["b"] and gap.getnext() == t2.el, \
        "между частями — только абзац с разрывом страницы"
    assert line[1:] not in V.docx_text(dx), "маркер не выводится текстом"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("text, parts, breaks", [
    ("| A | B |\n|---|---|\n| 1 | 1 |\n\\newpage\n| 2 | 2 |\n\\pagebreak\n| 3 | 3 |\n| 4 | 4 |\n",
     [["A", "1"], ["2"], ["3", "4"]], 2),
    ("| A | B |\n|---|---|\n| 1 | 1 |\n\\newpage\n\\pagebreak\n\\newpage\n| 2 | 2 |\n",
     [["A", "1"], ["2"]], 1),
    ("Абзац.\n\n| A | B |\n|---|---|\n\\newpage\n| 1 | 1 |\n| 2 | 2 |\n", [["A", "1", "2"]], 1),
    ("| A | B |\n|---|---|\n| 1 | 1 |\n| 2 | 2 |\n\\newpage\n\nАбзац.\n", [["A", "1", "2"]], 1),
    ("|  |  |\n|---|---|\n| 1 | 1 |\n\\newpage\n| 2 | 2 |\n", [["1"], ["2"]], 1),
], ids=["several", "series", "under-delimiter", "after-last-row", "empty-header"])
def test_newpage_inside_table_parts(tmp_path, text, parts, breaks):
    """Несколько разрывов — несколько частей; серия подряд — один разрез; разрыв сразу под
    строкой-разделителем — разрыв перед таблицей, таблица целиком; за последней строкой — как в
    0.2.0; при пустой шапке шапки нет ни у одной части. Сетка у частей одна (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "parts", text)
    assert not warnings_of(proc), proc.stderr
    assert [_first_col(t) for t in dx.tables] == parts
    assert len({tuple(t.grid) for t in dx.tables}) == 1, [t.grid for t in dx.tables]
    assert [_first_row_flag(t) for t in dx.tables] == \
        ["1" if k == 0 and parts[0][0] == "A" else "0" for k in range(len(parts))]
    assert _no_tbl_header(dx)
    found = sum(V.para_atoms(p).count("b") for p in paragraphs(dx))
    assert found == breaks, f"разрывов страницы {found}, ожидалось {breaks}"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("text", [
    "- Пункт\n\n  | A | B |\n  |---|---|\n  | 1 | 2 |\n  \\newpage\n  | 3 | 4 |\n",
    "> Цитата\n>\n> | A | B |\n> |---|---|\n> | 1 | 2 |\n> \\newpage\n> | 3 | 4 |\n",
], ids=["list", "quote"])
def test_newpage_inside_nested_table_ignored(tmp_path, text):
    """В таблице внутри списка или цитаты \\newpage игнорируется с предупреждением, как маркер:
    таблица целая, без строки-разрыва и без разрыва страницы (SPEC §3.4, §4.1)."""
    md, out, proc, dx = run_md(tmp_path, "nested", text)
    assert [_first_col(t) for t in dx.tables] == [["A", "1", "3"]]
    assert "newpage" not in V.docx_text(dx), "маркер не выводится текстом"
    assert not any("b" in V.para_atoms(p) for p in paragraphs(dx))
    warns = warnings_of(proc)
    assert len(warns) == 1 and "строка 6" in warns[0] and "игнорируется" in warns[0], warns
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_newpage_split_table_landscape_as_whole(tmp_path):
    """Широкой таблицу делает только вторая часть — в альбомный раздел уходит вся таблица, сетка у
    частей одна, по всем строкам (SPEC §3.4, §4.3)."""
    wide = " | ".join(f"Длинноесловобезпроб{i}" for i in range(5))
    md, out, proc, dx = run_md(tmp_path, "wide-part", "Абзац.\n\n| К0 | К1 | К2 | К3 | К4 |\n"
                               f"|---|---|---|---|---|\n| 0 | 1 | 2 | 3 | 4 |\n\\newpage\n| {wide} |\n\n"
                               "Абзац после.\n")
    assert not warnings_of(proc), proc.stderr
    assert orient(dx) == ["P", "L", "P"]
    t1, t2 = dx.tables
    assert (t1.section, t2.section) == (1, 1), "обе части — в альбомном разделе"
    assert t1.grid == t2.grid and sum(t1.grid) == V.TEXT_WIDTH["L"], [t1.grid, t2.grid]
    rep = V.run_checks(md, out, ["P", "L", "P"])
    assert rep.ok, rep.render()


def test_verify_catches_split_table_defects(tmp_path):
    """Сверщик ловит порчу частей таблицы: другая сетка у второй части, шапка у второй части,
    нет разрыва между частями, повтор шапки (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "split", SPLIT_TABLE.format(brk="\\newpage"))
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()

    def tbl(root, k):
        return root.findall(".//" + q("tbl"))[k]

    def regrid(root):  # вторая часть: колонки другой ширины, сумма та же
        t2 = tbl(root, 1)
        cols = t2.findall(f"{q('tblGrid')}/{q('gridCol')}")
        cols[0].set(q("w"), str(int(cols[0].get(q("w"))) + 300))
        cols[1].set(q("w"), str(int(cols[1].get(q("w"))) - 300))

    def head_on_part2(root):  # вторая часть размечена таблицей с шапкой
        tbl(root, 1).find(f"{q('tblPr')}/{q('tblLook')}").set(q("firstRow"), "1")

    def drop_break(root):  # между частями нет разрыва страницы
        pb = tbl(root, 0).getnext().find(f"{q('pPr')}/{q('pageBreakBefore')}")
        pb.getparent().remove(pb)

    def repeat_head(root):  # повтор шапки на страницах
        etree.SubElement(tbl(root, 0).find(f"{q('tr')}/{q('trPr')}"), q("tblHeader"))

    for name, fn, code, needle in [("regrid", regrid, "table-format", "сетке первой части"),
                                   ("head-on-part2", head_on_part2, "table-format", "tblLook"),
                                   ("drop-break", drop_break, "plan", "нет разрыва страницы"),
                                   ("repeat-head", repeat_head, "table-format", "tblHeader")]:
        bad = mutate(out, tmp_path / f"split-{name}.docx", "word/document.xml", fn)
        rep = V.run_checks(md, bad, ["P"])
        assert rep.has(code, needle), f"порча «{name}» не замечена:\n{rep.render()}"


@pytest.mark.parametrize("text, expected, breaks", [
    ("Абзац A.\n\n<!-- landscape -->\n\n\\newpage\n\nАбзац B.\n", ["P", "L"], 0),
    ("Абзац A.\n\n\\newpage\n\n<!-- portrait -->\n\nАбзац B.\n", ["P"], 1),
    ("Абзац A.\n\n\\newpage\n\n<!-- landscape -->\n\nАбзац B.\n", ["P", "L"], 0),
    ("Абзац A.\n\n<!-- pagebreak -->\n<!-- landscape -->\n<!-- portrait -->\n\nАбзац B.\n",
     ["P"], 1),
], ids=["orient-then-newpage", "newpage-then-same-orient", "newpage-then-new-orient",
        "break-and-two-orients"])
def test_marker_series(tmp_path, text, expected, breaks):
    """Маркеры подряд: ориентация — последняя из маркеров ориентации, разрыв — если в серии
    был хоть один; новый раздел и так с новой страницы — второго разрыва нет (SPEC §4.1)."""
    md, out, proc, dx = run_md(tmp_path, "series", text)
    assert orient(dx) == expected
    found = sum(V.para_atoms(p).count("b") for p in paragraphs(dx))
    assert found == breaks, f"разрывов страницы {found}, ожидалось {breaks}"
    assert dx.section_index(find_p(dx, "Абзац B.")) == len(expected) - 1
    rep = V.run_checks(md, out, expected)
    assert rep.ok, rep.render()


def test_escaped_newpage_is_text(tmp_path):
    """Маркер, экранированный второй обратной косой чертой, — текст «\\newpage», а не разрыв
    (SPEC §4.1)."""
    md, out, proc, dx = run_md(tmp_path, "escaped", "Текст до.\n\n\\\\newpage\n\nТекст после.\n")
    assert V.para_text(find_p(dx, "newpage")).strip() == "\\newpage"
    assert not any("b" in V.para_atoms(p) for p in paragraphs(dx)), "экранированный маркер — не разрыв"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("text, expected", [
    ("текст\n&#92;newpage\nещё\n", ["текст \\newpage ещё"]),
    ("текст\n&#92;newpage\n\\newpage\nещё\n", ["текст \\newpage", "ещё"]),
    ("текст\n&bsol;newpage\n\\newpage\nещё\n", ["текст \\newpage", "ещё"]),
], ids=["entity", "entity-then-break", "named-entity-then-break"])
def test_entity_newpage_is_text(tmp_path, text, expected):
    """Разрыв — только строка, которая в исходнике дословно \\newpage: «&#92;newpage» — текст,
    а настоящий \\newpage ниже в том же абзаце режет по своей строке (SPEC §4.1)."""
    md, out, proc, dx = run_md(tmp_path, "entity", text)
    texts = [V.para_text(p) for p in paragraphs(dx) if V.para_text(p).strip()]
    assert texts == expected
    assert [V.para_atoms(p)[0] for p in paragraphs(dx) if V.para_text(p).strip()][1:] == \
        ["b"] * (len(expected) - 1), "разрыв — перед «ещё»"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


# --------------------------------------------------------------------------
# 3. Таблицы
# --------------------------------------------------------------------------


def test_two_tables_separator_and_edges(outdir):
    docx, _ = convert(outdir, sample("02-tables"))
    dx = V.Docx(docx)
    kids = [c for c in dx.body if V.local(c)]
    assert V.local(kids[0]) == "tbl", "таблица первым блоком документа"
    for i, el in enumerate(kids):
        if V.local(el) == "tbl":
            assert V.local(kids[i + 1]) == "p", "за таблицей всегда абзац (иначе Word склеит)"
    first, second = find_table(dx, "x1"), find_table(dx, "x2")
    sep = first.el.getnext()
    assert sep is not None and V.local(sep) == "p" and sep.getnext() is second.el, \
        "между двумя таблицами подряд ровно один абзац-разделитель"
    assert not V.has_content(sep)
    assert spacing(dx, sep) == {"before": 0, "after": 0, "line": 160, "lineRule": "exact"}
    last = dx.tables[-1]
    assert "конец" in V.cell_text(last.rows[1][0]), "последний блок — таблица"
    tail = kids[kids.index(last.el) + 1:]
    assert [V.local(c) for c in tail] == ["p", "sectPr"], \
        f"после последней таблицы — один абзац и конец, а не {[V.local(c) for c in tail]}"
    assert not V.has_content(tail[0])
    assert spacing(dx, tail[0]) == {"before": 0, "after": 0, "line": 20, "lineRule": "exact"}, \
        "документ кончается таблицей — последний абзац в 1 pt, а не разделитель 8 pt (SPEC §3.4)"


def test_table_cells(outdir):
    docx, _ = convert(outdir, sample("02-tables"))
    dx = V.Docx(docx)
    al = find_table(dx, "По центру")
    want = [None, "center", "right", None]
    for ri, row in enumerate(al.rows):
        for ci, tc in enumerate(row):
            for p in tc.findall(q("p")):
                jc = dx.eff_ppr(p, "jc", "val")
                if want[ci] is None:
                    assert jc in (None, "left", "start"), (ri, ci, jc)
                else:
                    assert jc in ((want[ci],) if want[ci] == "center" else ("right", "end")), \
                        (ri, ci, jc)
    br = find_table(dx, "перенос")
    cell = br.rows[1][1]
    assert V.norm_lines(V.cell_text(cell)) == ("первая строка", "вторая строка"), \
        "<br> в ячейке — перенос строки"
    assert V.norm_lines(V.cell_text(br.rows[2][1])) == ("строка один", "строка два", "строка три")
    assert V.norm_ws(V.cell_text(br.rows[3][1])) == "значение | с вертикальной чертой"
    assert V.cell_text(br.rows[4][1]).strip() == ""
    assert [V.cell_text(tc).strip() for tc in br.rows[5]] == ["", "без ключа", ""]
    assert [V.cell_text(tc).strip() for tc in br.rows[6]] == ["короткая строка", "", ""], \
        "короткая строка дополняется пустыми ячейками"
    assert len(br.rows[7]) == 3 and "отброшенная" not in V.docx_text(dx), \
        "лишняя ячейка отбрасывается"
    mk = find_table(dx, "полужирный текст")
    cells = {V.cell_text(r[0]).strip(): r[1] for r in mk.rows[1:]}

    def run_in(cellname, text):
        tc = cells[cellname]
        return tc, run_with_text(tc.find(q("p")), text)

    tc, r = run_in("жирный", "полужирный текст")
    assert V.is_on(dx.eff_rpr(r, tc.find(q("p")), "b"))
    tc, r = run_in("курсив", "наклонный текст")
    assert V.is_on(dx.eff_rpr(r, tc.find(q("p")), "i"))
    tc, r = run_in("зачёркнутый", "старое значение")
    assert V.is_on(dx.eff_rpr(r, tc.find(q("p")), "strike"))
    tc, r = run_in("код", "inline-код")
    p = tc.find(q("p"))
    assert dx.eff_font(r, p) == "Menlo"
    assert V.to_int(wa(dx.eff_rpr(r, p, "sz"), "val")) == 20
    assert (wa(dx.eff_rpr(r, p, "shd"), "fill") or "").upper() == "F2F2F2"
    hl = [h for h in cells["ссылка"].iter(q("hyperlink"))]
    assert hl, "ссылка в ячейке — настоящая гиперссылка"
    rel = dx.rel_by_id[hl[0].get("{%s}id" % V.R_NS)]
    assert rel.get("Target") == "https://example.com/cell"
    for t in dx.tables:
        assert sum(t.grid) == V.TEXT_WIDTH["P"], f"таблица {t.index}: ширина = ширине текста"
        assert not any(V.is_on(tr.find("w:trPr/w:tblHeader", NS)) for tr in t.el.findall(q("tr"))), \
            f"таблица {t.index}: tblHeader — шапка на страницах не повторяется"


LONG_CELL = " ".join(["стороны обязуются соблюдать условия настоящего договора"] * 12)


@pytest.mark.parametrize("text", [
    f"| Описание | Срок | Действие | Примечание |\n|---|---|---|---|\n"
    f"| {LONG_CELL} | *отметить* | ***отметить*** | (договора) |\n",
    "| " + " | ".join(f"Дата{i}" for i in range(9)) + " |\n" + "|---" * 9 + "|\n| "
    + " | ".join(["2026-10-22"] * 9) + " |\n| " + " | ".join(["+7-916-123-45-67"] * 9) + " |\n",
    "| " + " | ".join(f"Ист{i}" for i in range(7)) + " |\n" + "|---" * 7 + "|\n| "
    + " | ".join(["см. https://buh.ru/news/peredacha-dannyh"] * 7) + " |\n",
    "| " + " | ".join(["Слово\u200bслово\u200bслово\u200bслово\u200bслово"] * 6) + " |\n"
    + "|---" * 6 + "|\n| " + " | ".join(["a"] * 6) + " |\n",
    f"| Логотип | Описание |\n|---|---|\n| ![логотип](logo.png) | {LONG_CELL} |\n"
    f"| ![схема](нет.png) | коротко |\n",
], ids=["arial-glyphs", "dates-phones", "urls", "zwsp", "images"])
def test_table_column_widths(tmp_path, text):
    """Ширины колонок (SPEC §3.4 п. 1–2): по бо́льшему из замеров Helvetica Neue и Arial;
    перенос после дефиса разрешён и перед цифрой; URL и ZWSP колонку не расширяют;
    картинка и заглушка в ячейке — содержимое колонки. Все таблицы остаются книжными."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "logo.png", 200, 100)
    md = write_md(tmp_path, "widths", text)
    out = tmp_path / "widths.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert orient(V.Docx(out)) == ["P"]
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_overflow_keeps_narrow_columns(tmp_path):
    """Таблица не влезает и в альбомную: рвутся слова только самых широких колонок,
    номера пунктов остаются целыми (SPEC §3.4 п. 3)."""
    word = "конфиденциальностьинформации"  # 28 знаков без места для переноса
    md = write_md(tmp_path, "overflow", "| № | " + " | ".join(f"{word}{i}" for i in range(6))
                  + " |\n" + "|---" * 7 + "|\n| 11.12 | " + " | ".join(["да"] * 6) + " |\n")
    out = tmp_path / "overflow.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert "не влезает" in proc.stderr, proc.stderr
    dx = V.Docx(out)
    assert orient(dx) == ["L"]
    t = dx.tables[0]
    mins, words = V.table_min_widths(dx, t, V.measurer())
    assert t.grid[0] >= mins[0] - 2, \
        f"колонка «№»: gridCol {t.grid[0]} уже слова «{words[0]}» ({round(mins[0])} с полями)"


def test_tall_row_may_split(tmp_path):
    """Строке выше страницы cantSplit не ставится: Word обрезал бы её текст, LibreOffice
    вывел бы его за нижнее поле (SPEC §3.4, «Строки»). Сверщик это различает."""
    long = " ".join(["условие договора о порядке расчётов"] * 400)
    md = write_md(tmp_path, "tall", f"| Пункт | Текст |\n|---|---|\n| 1 | коротко |\n| 2 | {long} |\n")
    out = tmp_path / "tall.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in V.Docx(out).body.iter(q("tr"))]
    assert cant == [True, True, False], f"cantSplit по строкам: {cant}"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()

    def add_cant_split(root):
        for pr in root.iter(q("trPr")):
            if pr.find(q("cantSplit")) is None:
                pr.insert(0, etree.Element(q("cantSplit")))

    bad = mutate(out, tmp_path / "tall-cant.docx", "word/document.xml", add_cant_split)
    assert V.run_checks(md, bad, ["P"]).has("table-format", "выше страницы")


def _need(word: str, bold: bool = False) -> float:
    """Ширина колонки, в которой слово стоит целиком: замер + 3 % запаса + поля ячейки 216
    (SPEC §3.4 п. 2); замер — свой, сверщика."""
    return V.measurer().twips(word, 11, bold=bold) * 1.03 + 216


def test_column_margin_three_percent(tmp_path):
    """Запас 3 % к замеру слова сторожится без LibreOffice: колонка из одного длинного слова
    не уже замера + 3 % + полей."""
    long = " ".join(["стороны обязуются соблюдать условия настоящего договора"] * 8)
    words = ("Электроэнцефалографический", "Высокотехнологичный")
    md, out, proc, dx = run_md(tmp_path, "margin", "| А | Б | Текст |\n|---|---|---|\n"
                               + f"| {words[0]} | {words[1]} | {long} |\n" * 3)
    assert orient(dx) == ["P"]
    grid = dx.tables[0].grid
    for j, w in enumerate(words):
        assert grid[j] >= _need(w) - 2, \
            f"колонка «{w}»: gridCol {grid[j]}, нужно не меньше {round(_need(w))} (замер + 3 % + поля)"


# Даты и коды в узких колонках книжной таблицы, где целые токены влезают (Σ мягких
# минимумов ≤ ширины) — по дефисам не делятся (SPEC §3.4, два минимума).
DATE_TOKENS = ("2026-10-22", "2027-01-31", "Ф-123456", "АБ-2026-001")
DATES_TABLE = ("| Дата начала | Дата конца | Ф-код | Номер | Комментарий |\n|---|---|---|---|---|\n"
               + ("| " + " | ".join(DATE_TOKENS) + " | "
                  + " ".join(["стороны обязуются соблюдать условия настоящего договора"] * 3)
                  + " |\n") * 3)


def test_codes_and_dates_not_split(tmp_path):
    words = "стороны обязуются соблюдать условия настоящего договора Комментарий".split()
    assert sum(map(_need, DATE_TOKENS)) + max(_need(w, True) for w in words) <= V.TEXT_WIDTH["P"], \
        "условие теста: целые токены влезают в книжную"
    md, out, proc, dx = run_md(tmp_path, "codes", "# Коды и даты\n\n" + DATES_TABLE)
    assert orient(dx) == ["P"]
    grid = dx.tables[0].grid
    for j, tok in enumerate(DATE_TOKENS):
        assert grid[j] >= _need(tok) - 2, \
            f"колонка {j + 1}: gridCol {grid[j]} уже «{tok}» ({round(_need(tok))} с запасом и полями)"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_soft_overflow_stays_portrait(tmp_path):
    """Целиком даты и телефоны не влезают (Σ мягких > ширины), с переносом после дефисов
    влезают — таблица книжная и без предупреждения: ориентацию решает жёсткий минимум."""
    toks = ("2026-10-22", "2027-01-31", "+7-916-123-45-67", "Ф-123456", "2026-12-01",
            "8-800-555-35-35", "АБ-2026-001", "77-88-99")
    assert sum(map(_need, toks)) > V.TEXT_WIDTH["P"], "условие теста: Σ мягких минимумов > ширины"
    md, out, proc, dx = run_md(
        tmp_path, "soft", "| Дата начала | Дата конца | Телефон | Код | Дата отчёта | Телефон 2 | "
        "Ф-код | Номер |\n" + "|---" * 8 + "|\n" + ("| " + " | ".join(toks) + " |\n") * 3)
    assert orient(dx) == ["P"]
    assert "не влезает" not in proc.stderr, proc.stderr
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_letter_dash_codes_decide_orientation(tmp_path):
    """«Ф-123456» после дефиса не переносится (буква, дефис, цифра — как в LibreOffice), значит,
    код целиком — жёсткий минимум: девять таких колонок не влезают в книжную — альбомная."""
    codes = [f"Ф-12345{i}" for i in range(9)]
    assert 9 * _need("123456") <= V.TEXT_WIDTH["P"] < sum(map(_need, codes)), \
        "условие теста: целые коды не влезают, а хвосты после дефиса влезли бы"
    md, out, proc, dx = run_md(tmp_path, "letter-dash", "| " + " | ".join(f"К{i}" for i in range(9))
                               + " |\n" + "|---" * 9 + "|\n| " + " | ".join(codes) + " |\n")
    assert orient(dx) == ["L"]
    rep = V.run_checks(md, out, ["L"])
    assert rep.ok, rep.render()


def test_widths_follow_text_volume(tmp_path):
    """Ширины по объёму текста тела (SPEC §3.4): одна длинная реплика в первой колонке не
    отнимает ширину у второй, где длинный текст в каждой строке."""
    rare = " ".join(["редкий длинный комментарий в одной ячейке"] * 6)
    dense = " ".join(["стороны обязуются соблюдать условия"] * 4)
    md, out, proc, dx = run_md(tmp_path, "volume", f"| Пояснение | Обязательства сторон |\n|---|---|\n"
                               f"| {rare} | {dense} |\n" + f"| — | {dense} |\n" * 5)
    assert orient(dx) == ["P"]
    grid = dx.tables[0].grid
    assert sum(grid) == V.TEXT_WIDTH["P"] and grid[1] > grid[0], \
        f"колонка с бо́льшим объёмом текста шире: {grid}"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


SMALL_TABLE = "| A | B |\n|---|---|\n| 1 | 2 |\n"


@pytest.mark.parametrize("text, flags, expected", [
    ("## Раздел\n\n" + SMALL_TABLE, (), ["P"]),
    ("## Раздел\n\n<!-- заметка -->\n\n" + SMALL_TABLE, (), ["P"]),
    ("## Раздел\n\n" + SMALL_TABLE, ("--landscape",), ["L"]),
    ("Текст.\n\n## Раздел\n\n" + WIDE_TABLE, (), ["P", "L"]),
], ids=["portrait", "invisible-block-between", "landscape-flag", "auto-landscape"])
def test_spacer_between_heading_and_table(tmp_path, text, flags, expected):
    """Таблица сразу за заголовком — перед ней пустой абзац 6 pt (before 0, after 0, line 120
    exact, keepNext), как в эталоне; в любой ориентации и в разделе таблицы (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "spacer", text, *flags)
    assert orient(dx) == expected
    t = dx.tables[0]
    sp = t.el.getprevious()
    assert sp is not None and V.local(sp) == "p" and not V.has_content(sp), \
        "перед таблицей за заголовком — пустой абзац-отбивка"
    got = spacing(dx, sp)
    assert (got["before"] or 0, got["after"], got["line"], got["lineRule"]) == (0, 0, 120, "exact"), got
    assert V.is_on(dx.eff_ppr(sp, "keepNext")), "отбивка не отрывается от таблицы"
    head = sp.getprevious()
    assert head is not None and dx.heading_level(head) == 2, "отбивка стоит между заголовком и таблицей"
    assert dx.section_index(sp) == t.section
    rep = V.run_checks(md, out, expected, landscape="--landscape" in flags)
    assert rep.ok, rep.render()


@pytest.mark.parametrize("text", [
    "## Раздел\n\nАбзац перед таблицей.\n\n" + SMALL_TABLE,
    "> ## Раздел\n>\n" + "".join(f"> {x}\n" for x in SMALL_TABLE.splitlines()),
    "- ## Раздел\n\n" + "".join(f"  {x}\n" for x in SMALL_TABLE.splitlines()),
], ids=["paragraph-between", "inside-quote", "inside-list"])
def test_no_spacer_without_heading_right_before(tmp_path, text):
    """Отбивка — только для пары «заголовок → таблица» на верхнем уровне."""
    md, out, proc, dx = run_md(tmp_path, "no-spacer", text)
    prev = dx.tables[0].el.getprevious()
    assert prev is not None and V.has_content(prev), "перед таблицей — соседний блок, без отбивки"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("after", ["Абзац после.\n", SMALL_TABLE], ids=["paragraph", "table"])
def test_no_separator_before_pagebreak(tmp_path, after):
    """Таблица, за ней \\newpage: разделителя 8 pt нет — таблицу отделяет абзац с разрывом
    страницы (следующий блок или носитель 1 pt перед таблицей). Иначе таблица, дошедшая
    до низа страницы, уносит разделитель на новую страницу, и та остаётся пустой (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "tbl-break", SMALL_TABLE + "\n\\newpage\n\n" + after)
    nxt = dx.tables[0].el.getnext()
    assert V.local(nxt) == "p" and V.para_atoms(nxt)[:1] == ["b"], \
        f"сразу за таблицей — абзац с разрывом страницы, а не разделитель: {V.para_atoms(nxt)}"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def _keep_rows(dx, t) -> list:
    """Номера строк таблицы, у которых keepNext стоит на всех абзацах."""
    return [i for i, tr in enumerate(t.el.findall(q("tr")))
            if all(V.is_on(dx.eff_ppr(p, "keepNext")) for p in tr.iter(q("p")))]


def test_table_rows_keep_together(tmp_path):
    """Таблица не начинается и не кончается одной строкой на странице: keepNext у шапки,
    первой строки тела и предпоследней строки; строке выше страницы — нет (SPEC §3.4)."""
    long = " ".join(["условие договора о порядке расчётов"] * 400)
    md, out, proc, dx = run_md(
        tmp_path, "keep", "| № | Пункт |\n|---|---|\n"
        + "".join(f"| {i} | пункт {i} |\n" for i in range(1, 5)) + "\nАбзац.\n\n"
        + f"| № | Пункт |\n|---|---|\n| 1 | коротко |\n| 2 | {long} |\n| 3 | коротко |\n")
    t1, t2 = dx.tables
    assert _keep_rows(dx, t1) == [0, 1, 3], "шапка + первая строка тела + предпоследняя"
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in t2.el.findall(q("tr"))]
    assert cant == [True, True, False, True], cant
    assert _keep_rows(dx, t2) == [0, 1], "строке без cantSplit (выше страницы) keepNext не ставится"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_keep_chain_must_fit(tmp_path):
    """Цепочка keepNext держится, только если влезает: шапка с первой строкой тела выше ¾ страницы —
    у шапки keepNext нет (иначе заголовок над таблицей остался бы внизу страницы один), хвост
    от первой строки тела держится; снятый keepNext влезающего хвоста сверщик ловит. Предпоследняя
    строка выше 90 % страницы за шапкой (без cantSplit) с последней не держится, даже если хвост
    влезает в страницу (SPEC §3.4)."""
    long = " ".join(["условие договора о порядке расчётов"] * 88)
    md, out, proc, dx = run_md(tmp_path, "chain", f"## Раздел\n\n| № | Пункт |\n|---|---|\n| 1 | {long} |\n"
                               "| 2 | коротко |\n| 3 | коротко |\n")
    t = dx.tables[0]
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in t.el.findall(q("tr"))]
    assert cant == [True] * 4, cant
    assert _keep_rows(dx, t) == [1, 2], "шапка с первой строкой тела выше ¾ страницы — без keepNext"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()

    def drop_tail_keep(root):
        for k in list(root.find(".//" + q("tbl")).findall(q("tr"))[2].iter(q("keepNext"))):
            k.getparent().remove(k)

    bad = mutate(out, tmp_path / "chain-tail.docx", "word/document.xml", drop_tail_keep)
    assert V.run_checks(md, bad, ["P"]).has("table-format", "нет keepNext")
    tall = "<br>".join(f"строка {k}" for k in range(1, 51))  # 50 строк: 0,9–1,0 страницы за шапкой
    md, out, proc, dx = run_md(tmp_path, "uncut", "| A | B |\n|---|---|\n| r1 | x |\n| r2 | x |\n"
                               f"| r3 | x |\n| {tall} | x |\n| r5 | z |\n")
    t = dx.tables[0]
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in t.el.findall(q("tr"))]
    assert cant == [True] * 4 + [False, True], cant
    assert _keep_rows(dx, t) == [0, 1], "строку без cantSplit с последней не держат, хоть хвост и влезает"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_keep_tail_fits_full_page(tmp_path):
    """Шапка на страницах не повторяется — хвост (предпоследняя строка с последней) держится, если
    влезает во всю высоту текста, а не в страницу за вычетом шапки: шапка в 10 строк, хвост из двух
    строк по 24 (SPEC §3.4)."""
    head = "<br>".join(f"шапка {k}" for k in range(1, 11))
    big = "<br>".join(f"строка {k}" for k in range(1, 25))
    md, out, proc, dx = run_md(tmp_path, "tail", f"| {head} | B |\n|---|---|\n| 1 | x |\n| 2 | x |\n"
                               f"| 3 | x |\n| {big} | x |\n| {big} | y |\n")
    t = dx.tables[0]
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in t.el.findall(q("tr"))]
    assert cant == [True] * 6, cant
    assert _keep_rows(dx, t) == [0, 1, 4], "хвост выше страницы за шапкой, но в страницу влезает"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_row_height_counts_images(tmp_path):
    """Высота строки — с картинками ячейки: строка с двумя снимками экрана (каждый до 0,8 высоты
    текста) выше страницы — без cantSplit и keepNext; шапка с тремя строками по картинке в 6000
    twips выше ¾ страницы — держится только шапка с первой строкой (SPEC §3.4)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "shot.png", 1170, 2532)
    _png(tmp_path / "photo.png", 200, 400)
    md, out, proc, dx = run_md(
        tmp_path, "img-rows", "| Экран | Описание |\n|---|---|\n"
        "| ![a](shot.png)<br>![b](shot.png) | два состояния |\n| следующий | экран |\n\nАбзац.\n\n"
        "| Фото | Описание |\n|---|---|\n" + "".join(f"| ![ф{i}](photo.png) | снимок {i} |\n" for i in (1, 2, 3)))
    t1, t2 = dx.tables
    cant = [V.is_on(tr.find("w:trPr/w:cantSplit", NS)) for tr in t1.el.findall(q("tr"))]
    assert cant == [True, False, True], f"строке выше страницы cantSplit не ставится: {cant}"
    assert 1 not in _keep_rows(dx, t1), "строке без cantSplit keepNext не ставится"
    assert _keep_rows(dx, t2) == [0, 2], "цепочка от шапки — только если влезает в ¾ страницы"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("head", ["|  |  |", "| &nbsp; | &nbsp; |"], ids=["empty", "nbsp"])
def test_empty_header_row_dropped(tmp_path, head):
    """Шапка из пустых ячеек не выводится: ни заливки, ни tblHeader, первая строка тела —
    обычная; keepNext — у первой и предпоследней строки тела (SPEC §3.4)."""
    md, out, proc, dx = run_md(tmp_path, "no-head", head + "\n|---|---|\n"
                               + "".join(f"| ячейка {i} | значение {i} |\n" for i in range(1, 6)))
    t = dx.tables[0]
    assert [V.cell_text(r[0]).strip() for r in t.rows] == [f"ячейка {i}" for i in range(1, 6)]
    trs = t.el.findall(q("tr"))
    assert not any(V.is_on(tr.find("w:trPr/w:tblHeader", NS)) for tr in trs), "tblHeader без шапки"
    fills = {(wa(tc.find("w:tcPr/w:shd", NS), "fill") or "").upper() for tc in t.el.iter(q("tc"))}
    assert "F2F2F2" not in fills, "строка тела — без заливки шапки"
    first = t.rows[0][0].find(q("p"))
    assert not V.is_on(dx.eff_rpr(run_with_text(first, "ячейка 1"), first, "b"))
    look = t.el.find("w:tblPr/w:tblLook", NS)
    assert look is None or wa(look, "firstRow") in ("0", "false", "off"), etree.tostring(look)
    assert _keep_rows(dx, t) == [0, 3]
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("depth", [30, 45])
def test_deep_quote_table_width(tmp_path, depth):
    """Таблица в глубокой цитате: отступ не глубже 8 уровней, ширина положительна и таблица
    не уходит за правое поле (SPEC §3.8)."""
    rows = ("| № | Имя | Дата |", "|---|---|---|", "| 1 | Иванов | 01.02.2026 |")
    md, out, proc, dx = run_md(tmp_path, "deep-quote", "".join(">" * depth + f" {x}\n" for x in rows))
    t = find_table(dx, "Иванов")
    width = dx.sections[t.section].text_width
    ind = V.to_int(wa(t.el.find("w:tblPr/w:tblInd", NS), "w")) or 0
    tblw = V.to_int(wa(t.el.find("w:tblPr/w:tblW", NS), "w")) or 0
    assert 0 <= ind <= 8 * 360, f"отступ таблицы {ind} — глубже 8 уровней"
    assert tblw > 0 and all(g > 0 for g in t.grid) and sum(t.grid) == tblw, (tblw, t.grid)
    assert ind + tblw <= width, f"таблица уходит за правое поле: {ind} + {tblw} > {width}"
    assert tblw >= width - 8 * 360, f"ширина {tblw} меньше ширины текста за вычетом 8 уровней"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


# --------------------------------------------------------------------------
# 4. Заголовки
# --------------------------------------------------------------------------


def test_headings(outdir):
    docx, _ = convert(outdir, sample("01-headings"))
    dx = V.Docx(docx)
    levels = [(dx.heading_level(p), V.norm_lines(V.para_text(p))) for p in paragraphs(dx)
              if dx.heading_level(p)]
    assert [lv for lv, _ in levels] == [1, 2, 3, 4, 5, 6, 1, 2, 2, 3, 2, 2]
    for p in paragraphs(dx):
        lv = dx.heading_level(p)
        if lv:
            assert dx.pstyle_id(p) == f"Heading{lv}", "встроенные стили Heading1…6"
    h = find_p(dx, "Заголовок")
    assert dx.heading_level(h) == 1
    h2 = find_p(dx, "Второй уровень с ")
    assert V.is_on(dx.eff_rpr(run_with_text(h2, "курсивом"), h2, "i"))
    assert V.is_on(dx.eff_rpr(run_with_text(h2, "зачёркнутым"), h2, "strike"))
    h3 = find_p(dx, "Третий уровень с ")
    assert dx.eff_font(run_with_text(h3, "кодом"), h3) == "Menlo"
    assert any(True for _ in h3.iter(q("hyperlink")))
    brh = find_p(dx, "с переносом строки")
    assert V.norm_lines(V.para_text(brh)) == ("Заголовок", "с переносом строки")
    assert sum(1 for _ in brh.iter(q("br"))) == 1
    br3 = find_p(dx, "через слэш")
    assert V.norm_lines(V.para_text(br3)) == ("Ещё перенос", "через слэш", "и с пробелом")
    assert V.norm_ws(V.para_text(find_p(dx, "Номер берётся"))) == "2.1 Номер берётся из текста"
    assert V.norm_ws(V.para_text(find_p(dx, "закрывающими"))) == \
        "Заголовок с закрывающими решётками"
    assert V.core_props(dx)["title"] == "Заголовок первого уровня"


def test_title_without_h1(outdir):
    docx, _ = convert(outdir, sample("02-tables"))
    assert V.core_props(V.Docx(docx))["title"] == "02-tables", "нет H1 — имя файла без расширения"


def test_title_empty_h1_and_control_chars(tmp_path):
    """Пустой H1 — заголовок документа из имени файла (SPEC §3.2); управляющие символы
    и U+FFFE в H1, тексте и alt вырезаются, а не роняют конвертер (SPEC §2)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "red.png", 20, 10)
    for name, text, title in (
            ("пустой-h1", "#\n\nТекст без заголовка.\n", "пустой-h1"),
            ("служебные", "# Титул\x01 с управляющим\ufffe\n\nабзац\x0bс табом\n\n"
                          "![alt\x02x](red.png)\n", "Титул с управляющим")):
        md = write_md(tmp_path, name, text)
        out = tmp_path / f"{name}.docx"
        proc = cli(md, "-o", out)
        assert proc.returncode == 0 and "Traceback" not in proc.stderr, proc.stderr[-2000:]
        assert V.core_props(V.Docx(out))["title"] == title
        rep = V.run_checks(md, out, ["P"])
        assert rep.ok, rep.render()


def test_long_title_truncated(tmp_path):
    """H1 длиннее 255 знаков: заголовок документа обрезан до 255, конвертация не падает,
    в тексте заголовок целиком (SPEC §3.2)."""
    h1 = "Договор " + "оказания услуг " * 20 + "конец"
    md, out, proc, dx = run_md(tmp_path, "long-title", f"# {h1}\n\nТекст.\n")
    title = V.core_props(dx)["title"]
    assert 200 <= len(title) <= 255 and h1.startswith(title.rstrip("… ")), \
        f"заголовок документа — начало H1 не длиннее 255 знаков, а не {title!r}"
    assert V.norm_ws(V.para_text(find_p(dx, "конец"))) == h1
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_inline_code_in_heading_size(tmp_path):
    """Inline-код в заголовке — кеглем заголовка (Menlo и заливка остаются), вне заголовков —
    10 pt (SPEC §3.6)."""
    md, out, proc, dx = run_md(tmp_path, "code-heading", "# Первый с `кодом`\n\n## Второй с `кодом`\n\n"
                               "### Третий с `кодом`\n\nАбзац с `кодом`.\n")
    for needle, sz in (("Первый", V.HEADING_SPEC[1][0]), ("Второй", V.HEADING_SPEC[2][0]),
                       ("Третий", V.HEADING_SPEC[3][0]), ("Абзац", 20)):
        p = find_p(dx, needle)
        r = run_with_text(p, "кодом")
        assert V.to_int(wa(dx.eff_rpr(r, p, "sz"), "val")) == sz, f"«{needle}»: кегль кода не {sz}"
        assert dx.eff_font(r, p) == "Menlo"
        assert (wa(dx.eff_rpr(r, p, "shd"), "fill") or "").upper() == "F2F2F2"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


# --------------------------------------------------------------------------
# 5. Списки
# --------------------------------------------------------------------------


def _num(dx, p):
    np_ = p.find("w:pPr/w:numPr", NS)
    assert np_ is not None, f"у пункта «{V.para_text(p)}» нет numPr"
    return wa(np_.find(q("numId")), "val"), V.to_int(wa(np_.find(q("ilvl")), "val")) or 0


def _lvl(dx, num_id, ilvl):
    info = V.numbering_info(dx)
    aid, so = info.num_ids[num_id]
    return info.abstracts[aid][ilvl], so


def test_numbering_restart(outdir):
    docx, _ = convert(outdir, sample("05-lists"))
    dx = V.Docx(docx)
    n = {t: _num(dx, find_p(dx, t)) for t in (
        "Первый пункт", "Второй пункт", "Вложенный пункт второго уровня", "Пункт третьего уровня",
        "Ещё пункт второго уровня", "Последний пункт маркированного списка",
        "Третий пункт", "Четвёртый пункт", "Вложенный нумерованный", "Второй вложенный",
        "Пятый пункт", "Альфа", "Бета", "Гамма", "Дельта", "Эпсилон со скобкой",
        "Дзета со скобкой", "Первый абзац пункта.", "Следующий пункт.")}
    assert n["Альфа"][0] == n["Бета"][0]
    assert n["Гамма"][0] == n["Дельта"][0]
    assert n["Эпсилон со скобкой"][0] == n["Дзета со скобкой"][0]
    assert len({n["Альфа"][0], n["Гамма"][0], n["Эпсилон со скобкой"][0],
                n["Третий пункт"][0], n["Первый абзац пункта."][0]}) == 5, \
        "соседние списки — разные w:num, иначе нумерация продолжится"
    assert n["Вложенный нумерованный"][0] != n["Третий пункт"][0], \
        "вложенный список — свой w:num"
    for text, ilvl, start, txt in (("Альфа", 0, 1, "%1."), ("Гамма", 0, 1, "%1."),
                                   ("Эпсилон со скобкой", 0, 1, "%1)"), ("Третий пункт", 0, 3, "%1."),
                                   ("Вложенный нумерованный", 1, 1, "%2."),
                                   ("Первый абзац пункта.", 0, 1, "%1.")):
        nid, lv = n[text]
        assert lv == ilvl, (text, lv)
        lvl, so = _lvl(dx, nid, ilvl)
        assert wa(lvl.find(q("numFmt")), "val") == "decimal", text
        assert wa(lvl.find(q("lvlText")), "val") == txt, text
        assert so.get(ilvl) == start, f"«{text}»: startOverride {so.get(ilvl)}, нужно {start}"
    for text, ilvl in (("Первый пункт", 0), ("Вложенный пункт второго уровня", 1),
                       ("Пункт третьего уровня", 2)):
        nid, lv = n[text]
        assert lv == ilvl
        lvl, _ = _lvl(dx, nid, ilvl)
        assert wa(lvl.find(q("numFmt")), "val") == "bullet"
        assert wa(lvl.find(q("lvlText")), "val") == V.BULLETS[ilvl]
        ind = lvl.find("w:pPr/w:ind", NS)
        assert V.to_int(wa(ind, "left") or wa(ind, "start")) == 360 * (ilvl + 1)
        assert V.to_int(wa(ind, "hanging")) == 360
    last_of_outer = {"Последний пункт маркированного списка", "Пятый пункт", "Бета", "Дельта",
                     "Дзета со скобкой", "Следующий пункт."}
    for text in n:
        p = find_p(dx, text)
        want = 160 if text in last_of_outer else 80
        assert spacing(dx, p)["after"] == want, f"«{text}»: after={spacing(dx, p)['after']}, нужно {want}"
        assert dx.pstyle_id(p) in ("Normal", dx.default_pstyle), "пункт — стиль Normal + numPr"
    second = find_p(dx, "Второй абзац того же пункта.")
    assert second.find("w:pPr/w:numPr", NS) is None, "второй абзац пункта — без numPr"
    assert V.to_int(dx.eff_ppr(second, "ind", "left") or dx.eff_ppr(second, "ind", "start")) == 360


def test_list_items_without_leading_paragraph(tmp_path):
    """Пустой пункт и пункт, который начинается с кода, тоже занимают номер — пустой абзац
    с numPr; иначе нумерация Word разойдётся с md (третий пункт получит номер 2)."""
    md = write_md(tmp_path, "items", "1. первый\n2. ```\n   код во втором пункте\n   ```\n3. третий\n\n"
                                     "- маркер\n-\n- после пустого\n")
    out = tmp_path / "items.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    dx = V.Docx(out)
    numbered = [V.para_text(p) for p in paragraphs(dx) if p.find("w:pPr/w:numPr", NS) is not None]
    assert numbered == ["первый", "", "третий", "маркер", "", "после пустого"], numbered
    assert "код во втором пункте" in [V.para_text(p) for p in paragraphs(dx)]
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_deep_nesting_keeps_tail(tmp_path):
    """Список глубже девяти уровней и цитата глубже двадцати не обрывают документ;
    пункты глубже девятого — на уровне Word 9 (SPEC §2, §3.5)."""
    items = "".join("  " * i + f"- уровень {i + 1}\n" for i in range(12))
    md = write_md(tmp_path, "deep", f"{items}\n{'>' * 22} глубокая цитата\n\n"
                                    "Абзац после вложенности.\n")
    out = tmp_path / "deep.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    dx = V.Docx(out)
    levels = [_num(dx, p)[1] for p in paragraphs(dx) if p.find("w:pPr/w:numPr", NS) is not None]
    assert levels == list(range(9)) + [8] * 3, levels
    find_p(dx, "глубокая цитата")
    find_p(dx, "Абзац после вложенности.")
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_bullets_in_arial(outdir):
    """Глифов ◦ и ▪ нет в Helvetica Neue — знаки маркированных списков набраны Arial (SPEC §3.5)."""
    dx = V.Docx(convert(outdir, sample("05-lists"))[0])
    used = [_lvl(dx, nid, ilvl)[0] for _, nid, ilvl in V.numbered_paragraphs(dx)]
    bullets = [lv for lv in used if wa(lv.find(q("numFmt")), "val") == "bullet"]
    assert bullets and all(wa(lv.find("w:rPr/w:rFonts", NS), "ascii") == "Arial" for lv in bullets)


def _effective_left(dx, p) -> int:
    """Левый отступ абзаца: свой ind → уровень нумерации → стиль."""
    own = dx.eff_ppr(p, "ind", "left") or dx.eff_ppr(p, "ind", "start")
    if own is None and p.find("w:pPr/w:numPr", NS) is not None:
        lvl, _ = _lvl(dx, *_num(dx, p))
        ind = lvl.find("w:pPr/w:ind", NS)
        own = wa(ind, "left") or wa(ind, "start")
    return V.to_int(own) or 0


@pytest.mark.parametrize("text, count", [
    ("".join("  " * i + f"- уровень{i + 1} слово\n" for i in range(30)), 30),
    ("".join(">" * (i + 1) + f" уровень{i + 1} слово\n\n" for i in range(45)), 45),
], ids=["list-30", "quote-45"])
def test_deep_indent_capped(tmp_path, text, count):
    """Отступ уровней списков и цитат — не глубже 8 уровней (8 × 360 twips); глубже — на
    отступе 8-го, текст не уходит за правое поле (SPEC §3.5, §3.8)."""
    md, out, proc, dx = run_md(tmp_path, "deep-indent", text + "\nАбзац после вложенности.\n")
    deep = [p for p in paragraphs(dx) if "уровень" in V.para_text(p)]
    assert len(deep) == count
    lefts = [_effective_left(dx, p) for p in deep]
    assert max(lefts) <= 8 * 360, f"отступы глубже 8 уровней: {sorted(set(lefts))}"
    find_p(dx, "Абзац после вложенности.")
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_max_nesting_warns(tmp_path):
    """Разбор оборвался по пределу вложенности — предупреждение, что часть документа
    пропущена (SPEC §2); на допустимой глубине — без него."""
    for depth, warn in ((50, True), (30, False)):
        items = "".join("  " * i + f"- пункт {i + 1}\n" for i in range(depth))
        md, out, proc, dx = run_md(tmp_path, f"nest-{depth}", items + "\nХвост документа.\n")
        hit = [x for x in warnings_of(proc) if "вложенности" in x and "пропущен" in x]
        assert bool(hit) == warn, f"глубина {depth}: {warnings_of(proc)}"
        if not warn:
            find_p(dx, "Хвост документа.")


# --------------------------------------------------------------------------
# 6. Цитаты, код, ссылки, картинки, линия, HTML, сущности, переносы
# --------------------------------------------------------------------------


def test_blocks(outdir):
    docx, proc = convert(outdir, sample("06-blocks"))
    dx = V.Docx(docx)
    text = V.docx_text(dx)
    for gone in ("Пример блоков", "Тестовый автор", "title:", "скрытый текст",
                 "отдельный комментарий"):
        assert gone not in text, f"«{gone}» не должно попасть в docx"
    q1, q2 = find_p(dx, "Цитата первого уровня"), find_p(dx, "Вложенная цитата")
    quote = dx.styles_by_id.get(dx.pstyle_id(q1))
    assert quote is not None and (wa(quote.find(q("name")), "val") or "") == "Quote"
    assert dx.pstyle_id(q2) == dx.pstyle_id(q1)
    assert V.to_int(dx.eff_ppr(q1, "ind", "left")) == 360
    assert V.to_int(dx.eff_ppr(q2, "ind", "left")) == 720, "360 на каждый уровень вложенности"
    left = quote.find("w:pPr/w:pBdr/w:left", NS)
    assert (wa(left, "val"), wa(left, "sz"), wa(left, "space"), (wa(left, "color") or "").upper()) \
        == ("single", "18", "8", "DDDDDD")
    assert spacing(dx, q1) == {"before": 160, "after": 160, "line": 276, "lineRule": "auto"}
    color = quote.find("w:rPr/w:color", NS)
    assert (wa(color, "val") or "").upper() == "555555" and wa(color, "themeColor") is None, \
        "цвет цитаты 555555 без themeColor (themeColor перебивает val)"
    assert V.is_on(quote.find("w:rPr/w:i", NS))
    item = find_p(dx, "пункт внутри цитаты")
    nid, ilvl = _num(dx, item)
    lvl, _ = _lvl(dx, nid, ilvl)
    own = V.to_int(dx.eff_ppr(item, "ind", "left"))
    lvl_left = V.to_int(wa(lvl.find("w:pPr/w:ind", NS), "left"))
    assert (own if own is not None else lvl_left) == 720, "список в цитате — с отступом цитаты"
    para = find_p(dx, "Абзац с ")
    r = run_with_text(para, "inline-кодом")
    assert dx.eff_font(r, para) == "Menlo" and V.to_int(wa(dx.eff_rpr(r, para, "sz"), "val")) == 20
    assert (wa(dx.eff_rpr(r, para, "shd"), "fill") or "").upper() == "F2F2F2"
    links = list(para.iter(q("hyperlink")))
    ext = [h for h in links if h.get("{%s}id" % V.R_NS)]
    assert len(ext) == 1, "одна внешняя гиперссылка, якорная — просто текст"
    rel = dx.rel_by_id[ext[0].get("{%s}id" % V.R_NS)]
    assert rel.get("Target") == "https://example.com/blocks" and rel.get("TargetMode") == "External"
    assert "внешней ссылкой" in V.para_text(ext[0])
    assert all("якорной" not in V.para_text(h) for h in links)
    hr_run = ext[0].find(q("r"))
    assert wa(hr_run.find("w:rPr/w:rStyle", NS), "val") is not None
    hstyle = dx.styles_by_id[wa(hr_run.find("w:rPr/w:rStyle", NS), "val")]
    assert (wa(hstyle.find("w:rPr/w:color", NS), "val") or "").upper() == "0563C1"
    assert wa(hstyle.find("w:rPr/w:u", NS), "val") == "single"
    code = find_p(dx, "строка без отступа")
    assert V.para_text(code) == \
        "строка без отступа\n    четыре пробела\n\tтабуляция\n        восемь пробелов", \
        f"блок кода — один абзац, отступы и табы сохранены: {V.para_text(code)!r}"
    for t in code.iter(q("t")):
        if t.text and (t.text != t.text.strip()):
            assert t.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    rc = run_with_text(code, "четыре")
    assert dx.eff_font(rc, code) == "Menlo" and V.to_int(wa(dx.eff_rpr(rc, code, "sz"), "val")) == 20
    assert (dx.eff_ppr(code, "shd", "fill") or "").upper() == "F2F2F2"
    assert spacing(dx, code) == {"before": 0, "after": 160, "line": 240, "lineRule": "auto"}
    assert dx.eff_ppr(code, "jc", "val") in (None, "left", "start")
    ind_code = find_p(dx, "отступной блок кода")
    assert V.para_text(ind_code) == "отступной блок кода\n    с вложенным отступом"
    ph = find_p(dx, "[изображение: Схема процесса]")
    pr = run_with_text(ph, "изображение")
    assert V.is_on(dx.eff_rpr(pr, ph, "i")), "заглушка картинки курсивом"
    col = (wa(dx.eff_rpr(pr, ph, "color"), "val") or "").upper()
    assert re.fullmatch(r"([0-9A-F]{2})\1\1", col) and col not in ("000000", "FFFFFF"), \
        f"заглушка серым цветом, а не {col!r}"
    assert any(x.startswith("Внимание:") and "missing-diagram.png" in x
               for x in proc.stderr.splitlines()), \
        f"пропущенная картинка — предупреждение с её путём: {proc.stderr!r}"
    hrs = [p for p in paragraphs(dx) if p.find("w:pPr/w:pBdr/w:bottom", NS) is not None
           and (dx.eff_ppr(p, "shd", "fill") or "").upper() != "F2F2F2"]  # рамки плашки кода — не линия
    assert len(hrs) == 1
    b = hrs[0].find("w:pPr/w:pBdr/w:bottom", NS)
    assert (wa(b, "val"), wa(b, "sz"), wa(b, "space"), (wa(b, "color") or "").upper()) == \
        ("single", "6", "1", "BBBBBB") and not V.para_text(hrs[0]).strip()
    assert V.norm_ws(V.para_text(find_p(dx, "Текст из"))) == "Текст из HTML-блока без тегов"
    assert any(x.startswith("Внимание:") and "HTML" in x for x in proc.stderr.splitlines()), \
        f"HTML-блок — с предупреждением: {proc.stderr!r}"
    assert V.norm_ws(V.para_text(find_p(dx, "Сущности:"))) == \
        "Сущности: «кавычки» & амперсанд — тире, © знак."
    hb = find_p(dx, "Жёсткий перенос")
    assert V.norm_lines(V.para_text(hb)) == ("Жёсткий перенос двумя пробелами",
                                             "и обратной косой чертой", "конец абзаца.")
    assert V.norm_ws(V.para_text(find_p(dx, "Комментарий"))) == \
        "Комментарий внутри абзаца не выводится."
    cp = V.core_props(dx)
    assert cp["title"] == "Блоки и inline-разметка" and not cp["author"]


WP_INLINE = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
WP_ANCHOR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor"
WP_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
EMU_PER_PX = 914400 // 96  # 9525 EMU — пиксель при 96 dpi
EMU_PER_TWIP = 635


def _png(path: Path, w: int, h: int):
    from PIL import Image
    Image.new("RGB", (w, h), (200, 60, 60)).save(path)  # без pHYs: плотность не задана


def test_images(tmp_path):
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовые картинки")
    (tmp_path / "img").mkdir()
    _png(tmp_path / "small.png", 200, 100)
    _png(tmp_path / "img" / "wide.png", 2000, 500)
    _png(tmp_path / "img" / "схема.png", 300, 150)
    md = write_md(tmp_path, "images",
                  "# Картинки\n\nМалая ![Малая](small.png) в строке текста.\n\n"
                  "![Широкая](img/wide.png)\n\n![Кириллица в имени](img/схема.png)\n\n"
                  "![Удалённая](https://example.com/remote.png)\n\nАбзац после картинок.\n")
    out = tmp_path / "images.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    dx = V.Docx(out)
    assert not list(dx.body.iter(WP_ANCHOR)), "картинки — в строку, не плавающие"
    small = list(find_p(dx, "Малая").iter(WP_INLINE))
    assert len(small) == 1, "картинка внутри абзаца остаётся в его строке"
    inl = list(dx.body.iter(WP_INLINE))
    assert len(inl) == 3, (f"три локальные картинки вставлены, а не заменены заглушкой "
                           f"(в т. ч. файл с кириллицей в имени); рисунков {len(inl)}, "
                           f"stderr: {proc.stderr!r}")
    ext = [(int(x.get("cx")), int(x.get("cy"))) for x in (e.find(WP_EXTENT) for e in inl)]
    max_cx = V.TEXT_WIDTH["P"] * EMU_PER_TWIP
    (cx1, cy1), (cx2, cy2), (cx3, cy3) = ext
    assert abs(cx1 - 200 * EMU_PER_PX) <= 200 * EMU_PER_PX // 100 and \
        abs(cy1 - 100 * EMU_PER_PX) <= 100 * EMU_PER_PX // 100, \
        f"200×100 px при 96 dpi = {200 * EMU_PER_PX}×{100 * EMU_PER_PX} EMU, а вставлено {cx1}×{cy1}"
    assert cx2 <= max_cx and abs(cx2 - max_cx) <= max_cx // 100, \
        f"широкая картинка ужимается до ширины текста {max_cx} EMU, а не {cx2}"
    assert abs(cy2 / cx2 - 500 / 2000) <= 0.01 * 500 / 2000, f"пропорции не сохранены: {cx2}×{cy2}"
    assert abs(cx3 - 300 * EMU_PER_PX) <= 300 * EMU_PER_PX // 100 and \
        abs(cy3 - 150 * EMU_PER_PX) <= 150 * EMU_PER_PX // 100
    assert V.norm_ws(V.para_text(find_p(dx, "[изображение:"))) == "[изображение: Удалённая]"
    assert any(x.startswith("Внимание:") and "remote.png" in x for x in proc.stderr.splitlines())
    rep = V.run_checks(md, out, "P")
    assert rep.ok, rep.render()


def test_html_blocks_and_comments(tmp_path):
    """HTML-блок: слова соседних элементов не склеиваются, содержимое <style> и <script>
    не выводится; </br> — перенос строки; на месте вырезанного комментария один пробел (SPEC §2)."""
    md = write_md(tmp_path, "html",
                  "<p>альфа</p><p>бета</p>\n\n<details><summary>Сводка</summary>Подробности"
                  "</details>\n\n<style>body { color: red }</style>\n\n<script>var секрет = 1;"
                  "</script>\n\nслово</br>слово\n\nТекст <!-- комментарий --> продолжение.\n")
    out = tmp_path / "html.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == 0, proc.stderr[-2000:]
    dx = V.Docx(out)
    text = V.docx_text(dx)
    assert "альфабета" not in text and "СводкаПодробности" not in text, text
    assert "color" not in text and "секрет" not in text, text
    assert V.norm_lines(V.para_text(find_p(dx, "слово"))) == ("слово", "слово")
    assert V.para_text(find_p(dx, "продолжение")) == "Текст продолжение."
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_front_matter(tmp_path):
    """Front matter вырезается и со списком YAML; «---» и пустая строка за ней — линия,
    а не метаданные: текст после неё не теряется (SPEC §2)."""
    for name, text, kept, gone in (
            ("yaml", "---\ntitle: Док\ntags:\n  - альфа\n  - бета\n---\n\n# Заголовок\n\nТекст.\n",
             "Текст.", "альфа"),
            ("hr", "---\n\nПримечание: это важно\n\n---\n\nТекст.\n", "Примечание: это важно", None)):
        md = write_md(tmp_path, name, text)
        out = tmp_path / f"{name}.docx"
        proc = cli(md, "-o", out)
        assert proc.returncode == 0, proc.stderr[-2000:]
        dx = V.Docx(out)
        find_p(dx, kept)
        assert gone is None or gone not in V.docx_text(dx)
        rep = V.run_checks(md, out, ["P"])
        assert rep.ok, rep.render()


def test_front_matter_dots_closer(tmp_path):
    """Front matter, закрытый «...» (как в pandoc и YAML), — тоже метаданные (SPEC §2)."""
    md, out, proc, dx = run_md(tmp_path, "fm-dots",
                               "---\ntitle: Отчёт\nauthor: Я\n...\n\n# Заголовок\n\nТекст.\n")
    text = V.docx_text(dx)
    assert "title:" not in text and "author:" not in text, text
    assert not any(p.find("w:pPr/w:pBdr", NS) is not None for p in paragraphs(dx)), "линии нет"
    assert V.core_props(dx)["title"] == "Заголовок"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_footnotes_as_text(tmp_path):
    """Сноски — обычным текстом: ссылка «[^1]» на месте, строка определения — абзац с меткой;
    ни битых гиперссылок, ни потери текста; в блоке кода — код (SPEC §3.6)."""
    md, out, proc, dx = run_md(tmp_path, "footnotes",
                               "Текст со сноской[^1] и ещё[^note].\n\n[^1]: Примечание.\n\n"
                               "[^note]: Длинное примечание из слов.\n\n```\n[^2]: в коде\n```\n")
    text = V.docx_text(dx)
    assert "сноской[^1]" in text and "ещё[^note]" in text, text
    assert V.para_text(find_p(dx, "Примечание.")).startswith("[^1]")
    assert V.para_text(find_p(dx, "Длинное примечание из слов.")).startswith("[^note]")
    assert not list(dx.body.iter(q("hyperlink"))), "сноска — не гиперссылка"
    code = find_p(dx, "в коде")
    assert V.para_text(code) == "[^2]: в коде"
    assert dx.eff_font(run_with_text(code, "в коде"), code) == "Menlo", "сноска в коде — код"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_link_schemes(tmp_path):
    """Гиперссылка — только http, https, mailto, tel (регистр не важен). Остальные схемы —
    текст ссылки с предупреждением (адрес обрезан), относительные пути и якоря — текст молча;
    форматирование текста ссылки сохраняется (SPEC §3.6)."""
    far = "smb://server.example.com/" + "/".join(["very-long-folder-name"] * 8) + "/file.docx"
    md, out, proc, dx = run_md(
        tmp_path, "links",
        "Можно: [сайт](https://example.com/a), [КРУПНО](HTTP://EXAMPLE.COM/B), "
        "[почта](mailto:a@example.com), [звонок](tel:+79161234567).\n\n"
        "Нельзя: [данные](data:image/png;base64,iVBORw0KGgo=), [шара](smb://127.0.0.1/share), "
        "[диагностика](ms-msdt:/id%20PCWDiagnostic), [**жирная**](ftp://example.com/x), "
        f"[далёкая]({far}), [хосты](file:///etc/hosts), [скрипт](javascript:alert(1)).\n\n"
        "Молча: [файл](docs/readme.md), [якорь](#razdel).\n")
    targets = set()
    for h in dx.body.iter(q("hyperlink")):
        rid = h.get("{%s}id" % V.R_NS)
        targets.add(dx.rel_by_id[rid].get("Target").lower() if rid else "#" + (wa(h, "anchor") or ""))
    assert targets == {"https://example.com/a", "http://example.com/b", "mailto:a@example.com",
                       "tel:+79161234567"}, targets
    text = V.docx_text(dx)
    for word in ("сайт", "КРУПНО", "почта", "звонок", "данные", "шара", "диагностика", "жирная",
                 "далёкая", "хосты", "скрипт", "файл", "якорь"):
        assert word in text, f"текст ссылки «{word}» потерян"
    p = find_p(dx, "жирная")
    assert V.is_on(dx.eff_rpr(run_with_text(p, "жирная"), p, "b")), "форматирование текста ссылки"
    warns = warnings_of(proc)
    for scheme in ("data:", "smb://127", "ms-msdt:", "ftp:", "smb://server.example.com", "file:",
                   "javascript:"):
        assert any(scheme in x for x in warns), f"нет предупреждения про {scheme}: {warns}"
    assert far not in proc.stderr, "длинный адрес в предупреждении обрезан"
    assert not any("readme.md" in x or "#razdel" in x for x in warns), "относительные и якоря — молча"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("text, flags, o", [
    ("![экран](tall.png)\n", (), "P"),
    ("![экран](tall.png)\n", ("--landscape",), "L"),
    ("| Снимок |\n|---|\n| ![экран](tall.png) |\n", (), "P"),
], ids=["portrait", "landscape", "table-cell"])
def test_tall_image_fits_page(tmp_path, text, flags, o):
    """Высокая картинка (снимок экрана телефона) — не выше 0,8 высоты текста страницы своей
    ориентации, пропорции сохранены; так же в ячейке (SPEC §3.6)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "tall.png", 1170, 2532)
    md, out, proc, dx = run_md(tmp_path, "tall-image", text, *flags)
    assert orient(dx) == [o]
    (inl,) = list(dx.body.iter(WP_INLINE))
    ext = inl.find(WP_EXTENT)
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    limit = 0.8 * V.TEXT_HEIGHT[o] * EMU_PER_TWIP
    assert cy <= limit + EMU_PER_TWIP, f"высота {cy} EMU больше 0,8 высоты текста ({limit:.0f} EMU)"
    assert abs(cy / cx - 2532 / 1170) <= 0.01 * 2532 / 1170, f"пропорции не сохранены: {cx}×{cy}"
    rep = V.run_checks(md, out, [o], landscape="--landscape" in flags)
    assert rep.ok, rep.render()


def test_narrow_tall_images_keep_portrait(tmp_path):
    """Ширина картинки для расчёта колонок — с учётом предела по высоте: шесть узких высоких
    картинок в строке таблицы не уводят её на альбомную (SPEC §3.4, §3.6)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "narrow.png", 300, 4000)
    md, out, proc, dx = run_md(tmp_path, "narrow-images", "| " + " | ".join(f"Рис{i}" for i in range(6))
                               + " |\n" + "|---" * 6 + "|\n| "
                               + " | ".join(f"![р{i}](narrow.png)" for i in range(6)) + " |\n")
    assert orient(dx) == ["P"]
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_huge_image_file_skipped(tmp_path):
    """Файл картинки больше 50 МБ не читается: заглушка и предупреждение (SPEC §3.6)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    big = tmp_path / "big.png"
    _png(big, 20, 10)
    os.truncate(big, 51 * 1024 * 1024)  # разреженный хвост нулями: места на диске почти не занимает
    md, out, proc, dx = run_md(tmp_path, "huge", "![огромная](big.png)\n")
    assert not list(dx.body.iter(WP_INLINE)), "файл больше 50 МБ не вставляется"
    find_p(dx, "[изображение: огромная]")
    assert any("big.png" in x for x in warnings_of(proc)), proc.stderr


def test_image_outside_document_dir(tmp_path):
    """Картинка выше каталога md или по абсолютному пути в другое место — вставляется, но
    с предупреждением; file: — как локальный файл (SPEC §3.6)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовые картинки")
    doc = tmp_path / "doc"
    doc.mkdir()
    for p in (tmp_path / "outside.png", tmp_path / "elsewhere.png", doc / "inside.png"):
        _png(p, 40, 20)
    md, out, proc, dx = run_md(doc, "pics", f"![выше](../outside.png)\n\n"
                               f"![абсолютный](<{tmp_path / 'elsewhere.png'}>)\n\n"
                               f"![внутри]({(doc / 'inside.png').as_uri()})\n")
    assert len(list(dx.body.iter(WP_INLINE))) == 3, proc.stderr
    warns = warnings_of(proc)
    assert len([x for x in warns if "вне каталога документа" in x]) == 2, warns
    assert not any("inside.png" in x for x in warns), warns
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_image_warnings_readable(tmp_path):
    """Предупреждения о картинках читаемы: имя файла раскодировано из %-кодировки, data: URI
    обрезан; alt с сущностями и экранированием — в заглушке как в тексте (SPEC §3.6)."""
    md, out, proc, dx = run_md(tmp_path, "img-warn", "![схема](нет.png)\n\n"
                               f"![данные](data:image/png;base64,{'A' * 5000})\n\n"
                               "![Схема &laquo;А&raquo; R&amp;D 10\\*20 a\\_b](missing.png)\n")
    warns = warnings_of(proc)
    assert any("нет.png" in x for x in warns) and "%D0" not in proc.stderr, proc.stderr
    data = [x for x in warns if "data:" in x]
    assert data and all(len(x) < 300 for x in data) and "A" * 100 not in proc.stderr, \
        f"data: URI в предупреждении не обрезан: длины строк {[len(x) for x in data]}"
    assert V.norm_ws(V.para_text(find_p(dx, "R&D"))) == "[изображение: Схема «А» R&D 10*20 a_b]"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


@pytest.mark.parametrize("entity", ["&#12;", "&#11;"], ids=["form-feed", "vertical-tab"])
def test_control_entities_become_spaces(tmp_path, entity):
    """Сущности управляющих символов-разделителей (&#12;, &#11;) — пробел, а не склейка слов
    (SPEC §2)."""
    e = entity
    md, out, proc, dx = run_md(tmp_path, "ctrl", f"# Титул{e}с FF\n\n| A{e}B | x |\n|---|---|\n"
                               f"| c{e}d | y |\n\nабзац{e}текст\n")
    assert V.core_props(dx)["title"] == "Титул с FF"
    t = dx.tables[0]
    assert [V.norm_ws(V.cell_text(r[0])) for r in t.rows] == ["A B", "c d"]
    find_p(dx, "абзац текст")
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


def test_inline_html_leftovers(tmp_path):
    """Незакрытый «<!--» в строке абзаца — текст с предупреждением; содержимое inline
    <script> и <style> не выводится; на месте комментария в конце выделения — один пробел
    (SPEC §2)."""
    md, out, proc, dx = run_md(tmp_path, "inline-html",
                               "Текст <!-- незакрытый в строке\n\nДальше.\n\n"
                               "Текст <script>alert('x')</script> и <style>p { color: red }</style>"
                               " дальше.\n\n*Слово <!-- c -->* слово.\n")
    assert V.para_text(find_p(dx, "незакрытый")) == "Текст <!-- незакрытый в строке"
    assert any("незакрыт" in x for x in warnings_of(proc)), proc.stderr
    text = V.docx_text(dx)
    assert "alert" not in text and "color" not in text, text
    assert V.norm_ws(V.para_text(find_p(dx, "дальше."))) == "Текст и дальше."
    assert V.para_text(find_p(dx, "Слово")) == "Слово слово."


def test_only_invisible_blocks_warns(tmp_path):
    """Документ только из невидимых блоков (в т. ч. <br>) — предупреждение «нет видимого
    содержимого» (SPEC §2)."""
    md, out, proc, dx = run_md(tmp_path, "invisible", "<!-- c -->\n\n<br>\n\n>\n")
    assert orient(dx) == ["P"]
    assert any("нет видимого содержимого" in x for x in warnings_of(proc)), proc.stderr


HN_FONT = Path("/System/Library/Fonts/HelveticaNeue.ttc")


def test_missing_glyph_in_arial(tmp_path):
    """Символ, которого нет в Helvetica Neue, но есть в Arial (→), — отдельным прогоном Arial
    с тем же начертанием; остальной текст и код (Menlo) не трогаются (SPEC §3.1)."""
    pytest.importorskip("PIL", reason="наличие глифа определяется через Pillow")
    if not HN_FONT.is_file():
        pytest.skip("нет Helvetica Neue — сравнивать не с чем")
    md, out, proc, dx = run_md(tmp_path, "arrow", "# Стрелка → вправо\n\nПереход А → Б в тексте.\n\n"
                               "| Было | Стало |\n|---|---|\n| старое → новое | **жирное → тоже** |\n\n"
                               "Абзац с `код → код` внутри.\n")
    arrows = [(p, r) for p in paragraphs(dx) for r in p.iter(q("r")) if "→" in V.para_text(r)]
    assert len(arrows) == 5, [V.para_text(r) for _, r in arrows]
    for p, r in arrows:
        t = V.para_text(r)
        if "код" in t:
            assert t == "код → код" and dx.eff_font(r, p) == "Menlo", f"код не трогается: {t!r}"
        else:
            fonts = (dx.eff_font(r, p), dx.eff_font(r, p, "hAnsi"))
            assert fonts == ("Arial", "Arial") and not re.search(r"\w", t), \
                f"«→» — отдельным прогоном Arial, а не {t!r} шрифтом {fonts}"
    p = find_p(dx, "жирное")
    assert V.is_on(dx.eff_rpr(run_with_text(p, "→"), p, "b")), "стрелка в полужирном — полужирная"
    p = find_p(dx, "Переход")
    assert dx.eff_font(run_with_text(p, "Переход"), p) != "Arial", "остальной текст — Helvetica Neue"
    rep = V.run_checks(md, out, ["P"])
    assert rep.ok, rep.render()


# --------------------------------------------------------------------------
# 7. Командная строка
# --------------------------------------------------------------------------


def test_cli_version():
    """--version печатает версию последнего раздела CHANGELOG (SPEC §1)."""
    m = re.search(r"^## (\d+\.\d+\.\d+)", (ROOT / "CHANGELOG.md").read_text(encoding="utf-8"), re.M)
    assert m, "в CHANGELOG нет раздела версии"
    proc = cli("--version")
    assert proc.returncode == 0 and re.search(rf"(?<![\d.]){re.escape(m.group(1))}(?![\d.])",
                                              proc.stdout + proc.stderr), \
        f"--version: {proc.stdout + proc.stderr!r}, в CHANGELOG — {m.group(1)}"


def test_cli_o_with_two_inputs(tmp_path):
    proc = cli(sample("01-headings"), sample("02-tables"), "-o", tmp_path / "x.docx")
    assert proc.returncode == 2, f"-o при двух входах — код 2, получено {proc.returncode}"
    assert proc.stderr.strip(), "нужно сообщение об ошибке"
    assert not (tmp_path / "x.docx").exists()


def test_cli_missing_file(tmp_path):
    proc = cli(tmp_path / "нет-такого.md", "-o", tmp_path / "x.docx")
    assert proc.returncode == 1, f"нет файла — код 1, получено {proc.returncode}"
    assert proc.stderr.strip() and "Traceback" not in proc.stderr, \
        f"нужно понятное сообщение, а не трассировка: {proc.stderr[-500:]}"


def test_cli_default_output_and_many_inputs(tmp_path):
    src, cwd = tmp_path / "src", tmp_path / "cwd"
    src.mkdir()
    cwd.mkdir()
    a, b = src / "doc.md", src / "второй.md"
    shutil.copy(sample("01-headings"), a)
    shutil.copy(sample("05-lists"), b)
    proc = cli(a, b, cwd=cwd)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert (src / "doc.docx").is_file() and (src / "второй.docx").is_file(), \
        "без -o результат пишется рядом с исходником"
    assert not list(cwd.iterdir()), "в текущий каталог ничего не пишется"
    lines = [x for x in proc.stdout.splitlines() if x.strip()]
    assert len(lines) == 2 and all(x.startswith("Готово:") for x in lines), proc.stdout


def test_cli_bad_inputs(tmp_path):
    """Битая картинка — заглушка с предупреждением; каталог вместо файла и -o на сам
    исходник — код 1 без трассировки, исходник цел, остальные файлы пакета конвертируются;
    пустой -o и значение у флага — код 2 с сообщением по-русски (SPEC §1)."""
    pytest.importorskip("PIL", reason="нужен Pillow, чтобы нарисовать тестовую картинку")
    _png(tmp_path / "full.png", 200, 100)
    (tmp_path / "bad.png").write_bytes((tmp_path / "full.png").read_bytes()[:20])
    bad = write_md(tmp_path, "bad-image", "Текст ![схема](bad.png) после.\n")
    good = write_md(tmp_path, "good", "# Заголовок\n\nТекст.\n")
    proc = cli(".", bad, good, cwd=tmp_path)
    assert proc.returncode == 1 and "Traceback" not in proc.stderr, proc.stderr[-2000:]
    find_p(V.Docx(tmp_path / "bad-image.docx"), "[изображение: схема]")
    assert (tmp_path / "good.docx").is_file(), "ошибка одного файла не обрывает пакет"
    before = good.read_bytes()
    proc = cli(good, "-o", good)
    assert proc.returncode == 1 and good.read_bytes() == before, proc.stderr
    assert cli(good, "-o", "").returncode == 2
    proc = cli(good, "--landscape=да")
    assert proc.returncode == 2 and "значение не предусмотрено" in proc.stderr, proc.stderr


ENC_TEXT = "# Отчёт о проверке\n\nТекст.\n"


@pytest.mark.parametrize("data, code, message", [
    (ENC_TEXT.encode("utf-8-sig"), 0, None),
    (b"\xff\xfe" + ENC_TEXT.encode("utf-16-le"), 0, None),
    (b"\xfe\xff" + ENC_TEXT.encode("utf-16-be"), 0, None),
    (ENC_TEXT.encode("utf-16-be"), 1, "похоже на UTF-16 без BOM или двоичный файл"),
    (ENC_TEXT.encode("cp1251"), 1, "файл не в кодировке UTF-8"),
], ids=["utf8-bom", "utf16le-bom", "utf16be-bom", "utf16-no-bom", "cp1251"])
def test_input_encodings(tmp_path, data, code, message):
    """BOM UTF-8 снимается, UTF-16 с BOM читается; без BOM байт 0 — ошибка с подсказкой, а не
    мусор в документе; не UTF-8 — ошибка (SPEC §1, §2)."""
    md = tmp_path / "enc.md"
    md.write_bytes(data)
    out = tmp_path / "enc.docx"
    proc = cli(md, "-o", out)
    assert proc.returncode == code and "Traceback" not in proc.stderr, proc.stderr[-2000:]
    if message:
        assert message in proc.stderr and not out.exists(), proc.stderr
    else:
        dx = V.Docx(out)
        assert V.core_props(dx)["title"] == "Отчёт о проверке"
        find_p(dx, "Текст.")


CP = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
DC = "{http://purl.org/dc/elements/1.1/}"


def _core_field(path: Path, tag: str):
    with zipfile.ZipFile(path) as z:
        el = etree.fromstring(z.read("docProps/core.xml")).find(tag)
    return None if el is None else (el.text or "")


def _edit_core(path: Path, fn):
    """Правка docProps/core.xml на месте — как сохранение файла в Word."""
    tmp = path.with_name(path.name + ".edit")
    mutate(path, tmp, "docProps/core.xml", fn)
    os.replace(tmp, path)


def _set_core(tag: str, value: str):
    def fn(root):
        el = root.find(tag)
        if el is None:
            el = etree.SubElement(root, tag)
        el.text = value
    return fn


def _drop_core(tag: str):
    def fn(root):
        for el in root.findall(tag):
            root.remove(el)
    return fn


def test_output_label_and_rewrite(tmp_path):
    """Результат помечен как файл md2docx (dc:identifier), revision 1, lastModifiedBy пуст;
    повторная конвертация в тот же файл — без вопросов (SPEC §1)."""
    md = write_md(tmp_path, "label", "# Метка\n\nТекст.\n")
    out = tmp_path / "label.docx"
    for _ in range(2):
        proc = cli(md, "-o", out)
        assert proc.returncode == 0, proc.stderr[-2000:]
    assert _core_field(out, DC + "identifier") == "md2docx"
    assert _core_field(out, CP + "revision") == "1"
    assert not _core_field(out, CP + "lastModifiedBy")


@pytest.mark.parametrize("spoil", ["revision", "editor", "no-label", "not-docx"])
def test_refuses_to_overwrite_foreign_file(tmp_path, spoil):
    """Файл результата правили после конвертации (revision > 1, lastModifiedBy) или он не
    от md2docx — отказ с кодом 1, файл цел; с --force — перезапись (SPEC §1)."""
    md = write_md(tmp_path, "foreign", "# Новый\n\nНовый текст.\n")
    out = tmp_path / "foreign.docx"
    if spoil == "not-docx":
        out.write_bytes("чужой файл, не docx".encode())
    else:
        assert cli(md, "-o", out).returncode == 0
        _edit_core(out, {"revision": _set_core(CP + "revision", "2"),
                         "editor": _set_core(CP + "lastModifiedBy", "Сотрудник"),
                         "no-label": _drop_core(DC + "identifier")}[spoil])
    before = out.read_bytes()
    proc = cli(md, "-o", out)
    assert proc.returncode == 1 and out.read_bytes() == before, \
        f"код {proc.returncode}, файл {'цел' if out.read_bytes() == before else 'перезаписан'}"
    assert "не перезаписываю" in proc.stderr and "--force" in proc.stderr, proc.stderr
    assert "Traceback" not in proc.stderr
    proc = cli(md, "-o", out, "--force")
    assert proc.returncode == 0, proc.stderr[-2000:]
    find_p(V.Docx(out), "Новый текст.")
    assert sorted(x.name for x in tmp_path.iterdir()) == ["foreign.docx", "foreign.md"], \
        "временных файлов не остаётся"


def test_batch_continues_after_refusal(tmp_path):
    """Отказ перезаписать чужой файл — ошибка этого входа (код 1), пакет идёт дальше."""
    a, b = write_md(tmp_path, "a", "Текст A.\n"), write_md(tmp_path, "b", "Текст B.\n")
    (tmp_path / "a.docx").write_bytes("чужой файл".encode())
    before = (tmp_path / "a.docx").read_bytes()
    proc = cli(a, b)
    assert proc.returncode == 1 and (tmp_path / "a.docx").read_bytes() == before, proc.stderr
    assert any("a.md" in x and "не перезаписываю" in x for x in proc.stderr.splitlines()), proc.stderr
    assert (tmp_path / "b.docx").is_file() and "Готово:" in proc.stdout


def test_batch_same_output_twice(tmp_path):
    """Два входа пакета с одним файлом результата: ранний результат остаётся, поздний вход —
    ошибка (код 1); тот же файл дважды — так же (SPEC §1)."""
    a = write_md(tmp_path, "notes", "Из md.\n")
    b = tmp_path / "notes.markdown"
    b.write_text("Из markdown.\n", encoding="utf-8")
    proc = cli(a, b)
    assert proc.returncode == 1, proc.stderr
    assert any(x.startswith("Ошибка:") and "notes.markdown" in x for x in proc.stderr.splitlines()), \
        proc.stderr
    find_p(V.Docx(tmp_path / "notes.docx"), "Из md.")
    assert proc.stdout.count("Готово:") == 1, proc.stdout
    proc = cli(a, a)
    assert proc.returncode == 1 and proc.stdout.count("Готово:") == 1, proc.stdout + proc.stderr


def api(script: str, *args, cwd: Path) -> subprocess.CompletedProcess:
    """Python-скрипт с md2docx из корня проекта (импорт, а не CLI); байткод не пишется."""
    return subprocess.run([sys.executable, "-c", script, str(ROOT), *map(str, args)], cwd=cwd,
                          capture_output=True, text=True, timeout=300,
                          env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})


API_SCRIPT = """
import json, sys
sys.path.insert(0, sys.argv[1])
from md2docx import convert
print(json.dumps(convert(sys.argv[2], sys.argv[3], landscape=True, auto_landscape=False),
                 ensure_ascii=False))
"""


def test_api_convert_keywords(tmp_path):
    """convert() из README: именованные landscape и auto_landscape, словарь итога, файл (SPEC §1)."""
    md = write_md(tmp_path, "api", "# API\n\n" + SMALL_TABLE)
    out = tmp_path / "api.docx"
    proc = api(API_SCRIPT, md, out, cwd=tmp_path)
    assert proc.returncode == 0, proc.stderr[-2000:]
    res = json.loads(proc.stdout)
    assert res == {"output": str(out), "sections": ["А"], "tables": 1, "auto_landscape_tables": 0,
                   "warnings": []}, res
    assert orient(V.Docx(out)) == ["L"]


FORCE_SCRIPT = """
import sys
sys.path.insert(0, sys.argv[1])
from md2docx import convert
try:
    convert(sys.argv[2], sys.argv[3])
except FileExistsError as exc:
    print("FileExistsError:", exc)
convert(sys.argv[2], sys.argv[3], force=True)
"""


def test_api_convert_force(tmp_path):
    """convert() не перезаписывает чужой файл (FileExistsError с русским текстом), с force=True —
    перезаписывает (SPEC §1)."""
    md = write_md(tmp_path, "force", "Текст.\n")
    out = tmp_path / "force.docx"
    out.write_bytes("чужой файл".encode())
    proc = api(FORCE_SCRIPT, md, out, cwd=tmp_path)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert "FileExistsError:" in proc.stdout and "не перезаписываю" in proc.stdout, proc.stdout
    find_p(V.Docx(out), "Текст.")


ATOMIC_SCRIPT = """
import errno, sys
sys.path.insert(0, sys.argv[1])
import docx.document, md2docx

def broken_save(self, target):  # запись оборвалась: место на диске кончилось
    data = b"PK" + bytes(998)
    if hasattr(target, "write"):
        target.write(data)
    else:
        with open(target, "wb") as f:
            f.write(data)
    raise OSError(errno.ENOSPC, "No space left on device")

docx.document.Document.save = broken_save
sys.exit(md2docx.main(sys.argv[2:]))
"""


def test_atomic_write_keeps_previous_result(tmp_path):
    """Сбой записи — код 1 без трассировки, прежний результат цел, временных файлов нет
    (SPEC §1: запись атомарная)."""
    md = write_md(tmp_path, "atomic", "# Атомарность\n\nТекст.\n")
    out = tmp_path / "atomic.docx"
    assert cli(md, "-o", out).returncode == 0
    before = out.read_bytes()
    proc = api(ATOMIC_SCRIPT, md, "-o", out, cwd=tmp_path)
    assert proc.returncode == 1 and "Ошибка:" in proc.stderr and "Traceback" not in proc.stderr, \
        proc.stderr[-2000:]
    assert out.read_bytes() == before, "прежний результат испорчен недописанным файлом"
    assert sorted(x.name for x in tmp_path.iterdir()) == ["atomic.docx", "atomic.md"]


def test_write_respects_readonly_and_symlink(tmp_path):
    """Замена файла не обходит защиту от записи: результат только для чтения — «нет доступа»,
    код 1, и с --force, файл цел; символьная ссылка остаётся ссылкой, заменяется её цель (SPEC §1)."""
    md = write_md(tmp_path, "ro", "# Только чтение\n\nТекст.\n")
    out = tmp_path / "ro.docx"
    assert cli(md, "-o", out).returncode == 0
    before = out.read_bytes()
    out.chmod(0o444)
    try:
        for flags in ((), ("--force",)):
            proc = cli(md, "-o", out, *flags)
            assert proc.returncode == 1 and "нет доступа" in proc.stderr, proc.stderr[-2000:]
        assert out.read_bytes() == before and not os.access(out, os.W_OK)
    finally:
        out.chmod(0o644)
    target, link = tmp_path / "real" / "ro.docx", tmp_path / "link.docx"
    target.parent.mkdir()
    out.rename(target)
    link.symlink_to(target)
    ino = target.stat().st_ino
    assert cli(md, "-o", link).returncode == 0
    assert link.is_symlink() and target.stat().st_ino != ino, "ссылка заменена файлом или цель не обновлена"
    assert sorted(x.name for x in target.parent.iterdir()) == ["ro.docx"]


def test_write_keeps_mode_and_names_directory(tmp_path):
    """Замена сохраняет права прежнего результата (600 не становится 644 по umask); в каталоге
    без права записи временный файл не создать — «нет доступа» называет каталог, а не
    записываемый файл в нём (SPEC §1)."""
    md = write_md(tmp_path, "mode", "# Права\n\nТекст.\n")
    out = tmp_path / "rodir" / "mode.docx"
    out.parent.mkdir()
    assert cli(md, "-o", out).returncode == 0
    out.chmod(0o600)
    assert cli(md, "-o", out).returncode == 0
    assert out.stat().st_mode & 0o777 == 0o600, f"права результата сброшены: {oct(out.stat().st_mode)}"
    out.parent.chmod(0o555)
    try:
        proc = cli(md, "-o", out)
    finally:
        out.parent.chmod(0o755)
    assert proc.returncode == 1 and proc.stderr.strip().endswith(f"нет доступа: {out.parent}"), \
        proc.stderr[-2000:]


@pytest.mark.parametrize("flags", [(), ("--force",)], ids=["plain", "force"])
def test_output_is_directory(tmp_path, flags):
    """-o на каталог — «это каталог, а не файл», а не подсказка про --force, который тоже не
    поможет (SPEC §1)."""
    md = write_md(tmp_path, "dir", "Текст.\n")
    (tmp_path / "outdir").mkdir()
    proc = cli(md, "-o", tmp_path / "outdir", *flags)
    assert proc.returncode == 1 and "это каталог, а не файл" in proc.stderr \
        and "--force" not in proc.stderr, proc.stderr[-2000:]


INTERNAL_SCRIPT = """
import sys
sys.path.insert(0, sys.argv[1])
import md2docx
real = md2docx.convert

def convert(src, *args, **kw):
    if str(src).endswith("a.md"):
        raise RuntimeError("проба")
    return real(src, *args, **kw)

md2docx.convert = convert
sys.exit(md2docx.main(sys.argv[2:]))
"""


def test_cli_internal_error_continues_batch(tmp_path):
    """Непредвиденная ошибка на одном файле — «внутренняя ошибка (класс: текст)», код 1,
    без трассировки; пакет идёт дальше (SPEC §1)."""
    a, b = write_md(tmp_path, "a", "Текст A.\n"), write_md(tmp_path, "b", "Текст B.\n")
    proc = api(INTERNAL_SCRIPT, a, b, cwd=tmp_path)
    assert proc.returncode == 1 and "Traceback" not in proc.stderr, proc.stderr[-2000:]
    assert "a.md: внутренняя ошибка (RuntimeError: проба)" in proc.stderr, proc.stderr
    assert (tmp_path / "b.docx").is_file() and "Готово:" in proc.stdout


def test_cli_os_error_in_russian(tmp_path):
    """Ошибка ОС при записи — «ошибка ввода-вывода (errno …)», код 1, без трассировки; номер
    errno у разных ОС разный — не проверяется (SPEC §1)."""
    md = write_md(tmp_path, "good", "Текст.\n")
    proc = cli(md, "-o", tmp_path / ("x" * 300 + ".docx"))
    assert proc.returncode == 1 and "ошибка ввода-вывода (errno" in proc.stderr, proc.stderr[-2000:]
    assert "Traceback" not in proc.stderr


def _nfd(s: str) -> str:
    return unicodedata.normalize("NFD", s)


# Имя файла → имя замка Word для Mac 16.114 (замер): основа длиной n символов в NFC; n ≤ 6 —
# «~$» + основа, n = 7 — без первого символа, n ≥ 8 — без двух первых. Имя замка Word пишет в NFD
WORD_LOCKS = [
    ("abcdef", "~$abcdef.docx"),
    ("abcdefg", "~$bcdefg.docx"),
    ("abcdefgh", "~$cdefgh.docx"),  # он же замок открытого cdefgh.docx — тоже отказ
    ("abcdefghijkl", "~$cdefghijkl.docx"),
    ("абвгдеж", "~$бвгдеж.docx"),
    ("Отчёт за год", _nfd("~$чёт за год.docx")),
    ("ab cdefgh", "~$ cdefgh.docx"),
    ("ёлка12", _nfd("~$ёлка12.docx")),
    (_nfd("йод123"), _nfd("~$йод123.docx")),  # в NFD 7 символов, в NFC — 6
    (_nfd("йогурт12"), "~$гурт12.docx"),
    ("Report", "~$report.docx"),  # имя набрано в другом регистре, чем на диске: APFS регистр не различает
]


@pytest.mark.parametrize("name, lock", WORD_LOCKS,
                         ids=["n6", "n7", "n8", "n12", "cyrillic", "space", "lock-space",
                              "nfd-lock", "nfd-name-n6", "nfd-name-n8", "case"])
def test_refuses_result_open_in_word(tmp_path, name, lock):
    """Рядом с результатом замок Word (правило имени — по замеру в Word для Mac): отказ с кодом 1,
    файл цел, даже если это нетронутый результат md2docx; с --force — перезапись; convert() —
    FileExistsError с тем же текстом. Длина имени — в символах NFC, имя замка на диске может
    быть в NFD и в другом регистре (SPEC §1)."""
    md = write_md(tmp_path, name, "Старый текст.\n")
    out = tmp_path / f"{name}.docx"
    assert cli(md).returncode == 0
    (tmp_path / lock).write_bytes(b"owner")
    md.write_text("Новый текст.\n", encoding="utf-8")
    before = out.read_bytes()
    proc = cli(md)
    assert proc.returncode == 1 and out.read_bytes() == before, proc.stdout + proc.stderr
    assert proc.stderr.strip() == (f"Ошибка: {md}: {out}: файл открыт в Word (рядом {lock}) — "
                                   f"закройте его или запустите с --force"), proc.stderr
    proc = cli(md, "--force")
    assert proc.returncode == 0, proc.stderr[-2000:]
    find_p(V.Docx(out), "Новый текст.")
    proc = api(FORCE_SCRIPT, md, out, cwd=tmp_path)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert f"FileExistsError: {out}: файл открыт в Word (рядом {lock})" in proc.stdout, proc.stdout


@pytest.mark.parametrize("name, other", [("abcdefg", "~$cdefg.docx"),
                                         ("abcdefgh", "~$abcdefgh.docx"),
                                         ("отчёт", "~$другой.docx")],
                         ids=["n7-two-cut", "n8-full-name", "other-document"])
def test_no_word_lock_rewrites_as_before(tmp_path, name, other):
    """Без замка своего файла — как раньше: нетронутый результат перезаписывается. При n = 7
    «~$» + имя без двух первых символов и при n ≥ 8 «~$» + полное имя — замки других файлов
    (cdefg.docx, xyabcdefgh.docx), не этого; замок другого документа — тоже (SPEC §1)."""
    md = write_md(tmp_path, name, "Старый текст.\n")
    assert cli(md).returncode == 0
    (tmp_path / other).write_bytes(b"owner")
    md.write_text("Новый текст.\n", encoding="utf-8")
    proc = cli(md)
    assert proc.returncode == 0, proc.stderr[-2000:]
    find_p(V.Docx(tmp_path / f"{name}.docx"), "Новый текст.")


def test_word_lock_in_unreadable_directory(tmp_path):
    """Каталог результата без права чтения: имена в нём не перечислить — замок Word ищется
    по имени, отказ тот же, файл цел (SPEC §1)."""
    md = write_md(tmp_path, "abcdefgh", "Старый текст.\n")
    out = tmp_path / "wx" / "abcdefgh.docx"
    out.parent.mkdir()
    assert cli(md, "-o", out).returncode == 0
    (out.parent / "~$cdefgh.docx").write_bytes(b"owner")
    md.write_text("Новый текст.\n", encoding="utf-8")
    before = out.read_bytes()
    out.parent.chmod(0o300)
    try:
        proc = cli(md, "-o", out)
    finally:
        out.parent.chmod(0o755)
    assert proc.returncode == 1 and "файл открыт в Word (рядом ~$cdefgh.docx)" in proc.stderr, \
        proc.stdout + proc.stderr
    assert out.read_bytes() == before


def _word_mac_resave(root):
    """fontTable после сохранения в Word для Mac (README «Helvetica Neue на Windows»): altName
    удалён, класс Helvetica Neue — auto вместо swiss."""
    for f in root.iter(q("font")):
        alt = f.find(q("altName"))
        if alt is not None:
            f.remove(alt)
        if wa(f, "name") == V.HN:
            f.find(q("family")).set(q("val"), "auto")


def _font_table(path: Path) -> list:
    """fontTable по смыслу: [(шрифт, [(поле, атрибуты)])] в порядке файла."""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/fontTable.xml"))
    return [(wa(f, "name"), [(c.tag, dict(c.attrib)) for c in f]) for f in root.iter(q("font"))]


def _other_members(path: Path) -> list:
    with zipfile.ZipFile(path) as z:
        return [(n, z.read(n)) for n in z.namelist() if n != "word/fontTable.xml"]


def test_fix_fonts_restores_hints(tmp_path, outdir):
    """--fix-fonts возвращает подсказки, которые стирает Word для Mac: fontTable — как у
    md2docx, прочие члены пакета — те же байты в том же порядке, права файла прежние,
    сверщик снова зелёный; второй прогон — «без изменений», файл не переписан (SPEC §1, §3.2)."""
    md = sample("06-blocks")
    orig, _ = convert(outdir, md)
    doc = mutate(orig, tmp_path / "resaved.docx", "word/fontTable.xml", _word_mac_resave)
    assert V.run_checks(md, doc, EXPECT_ORIENT["06-blocks"]).codes() == {"fonts"}, \
        "синтетический «пересохранённый» файл должен расходиться со сверщиком только в шрифтах"
    doc.chmod(0o600)
    before = _other_members(doc)
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert proc.stdout.strip() == (f"Готово: {doc} — восстановлены подсказки шрифтов: "
                                   f"Helvetica Neue (altName, family), Menlo (altName)"), proc.stdout
    assert _font_table(doc) == _font_table(orig)
    assert _other_members(doc) == before
    rep = V.run_checks(md, doc, EXPECT_ORIENT["06-blocks"])
    assert rep.ok, rep.render()
    assert doc.stat().st_mode & 0o777 == 0o600, f"права файла сброшены: {oct(doc.stat().st_mode)}"
    stamp = doc.stat().st_mtime_ns, doc.read_bytes()
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert proc.stdout.strip() == f"Без изменений: {doc} — подсказки шрифтов на месте", proc.stdout
    assert (doc.stat().st_mtime_ns, doc.read_bytes()) == stamp, "файл без изменений переписан"
    assert sorted(x.name for x in tmp_path.iterdir()) == ["resaved.docx"], "временных файлов не остаётся"


# Настоящие файлы после Word для Mac 16.114: выход md2docx 0.2.0 для samples/02 и samples/06,
# открытый и сохранённый в Word; в docProps/core.xml автор и «кем изменён» заменены на «Тест»
FIXTURES = TESTS / "fixtures"
HINTS = ("altName", "family", "panose1", "pitch")


def _font_fields(path: Path) -> dict:
    """{шрифт: {поле: w:val}} по fontTable."""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/fontTable.xml"))
    return {wa(f, "name"): {V.local(c): wa(c, "val") for c in f} for f in root.iter(q("font"))}


@pytest.mark.parametrize("smp", ["02-tables", "06-blocks"])
def test_fix_fonts_word_mac_fixture(tmp_path, outdir, smp):
    """Файл, сохранённый в Word для Mac: у Helvetica Neue и Menlo нет altName. --fix-fonts
    возвращает подсказки, как в выходе md2docx; сверщик больше не видит расхождений шрифтов,
    остальные (правки Word) те же; прочие члены zip — те же байты в том же порядке; повторный
    запуск — «без изменений», файл не переписан (SPEC §1, §3.2)."""
    md, doc = sample(smp), tmp_path / f"{smp}-resaved.docx"
    shutil.copy(FIXTURES / doc.name, doc)
    fonts = _font_fields(doc)
    assert all("altName" not in fonts[name] for name in (V.HN, "Menlo")), \
        f"фикстура не «после Word для Mac»: {fonts[V.HN]}, {fonts['Menlo']}"
    before = V.run_checks(md, doc, EXPECT_ORIENT[smp])
    assert before.has("fonts", "altName"), before.render()
    with zipfile.ZipFile(doc) as z:
        names = z.namelist()
    members = _other_members(doc)
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert proc.stdout.strip() == (f"Готово: {doc} — восстановлены подсказки шрифтов: "
                                   f"Helvetica Neue (altName, family), Menlo (altName)"), proc.stdout
    want, got = _font_fields(convert(outdir, md)[0]), _font_fields(doc)
    for name in (V.HN, "Menlo"):
        assert {k: got[name].get(k) for k in HINTS} == {k: want[name].get(k) for k in HINTS}, name
    with zipfile.ZipFile(doc) as z:
        assert z.namelist() == names, "порядок членов zip"
    assert _other_members(doc) == members, "прочие члены zip — те же байты"
    after = V.run_checks(md, doc, EXPECT_ORIENT[smp])
    assert after.errors == [e for e in before.errors if e.code != "fonts"], after.render()
    stamp = doc.stat().st_mtime_ns, doc.read_bytes()
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert proc.stdout.strip() == f"Без изменений: {doc} — подсказки шрифтов на месте", proc.stdout
    assert (doc.stat().st_mtime_ns, doc.read_bytes()) == stamp, "файл без изменений переписан"


REFUSE_SCRIPT = """
import sys
sys.path.insert(0, sys.argv[1])
from md2docx import convert
try:
    convert(sys.argv[2], sys.argv[3])
except FileExistsError as exc:
    print("FileExistsError:", exc)
"""


@pytest.mark.parametrize("smp", ["02-tables", "06-blocks"])
def test_word_mac_fixture_not_overwritten(tmp_path, smp):
    """Результат, сохранённый в Word для Mac (revision 2, автор заполнен), convert() без force
    не заменяет: FileExistsError «правился после конвертации», файл цел (SPEC §1)."""
    out = tmp_path / f"{smp}.docx"
    shutil.copy(FIXTURES / f"{smp}-resaved.docx", out)
    before = out.read_bytes()
    proc = api(REFUSE_SCRIPT, sample(smp), out, cwd=tmp_path)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert proc.stdout.strip() == (f"FileExistsError: {out} уже существует и правился после "
                                   f"конвертации или создан не md2docx — не перезаписываю; чтобы "
                                   f"перезаписать, добавьте --force"), proc.stdout
    assert out.read_bytes() == before


def _drop_menlo(root):
    for f in root.iter(q("font")):
        if wa(f, "name") == "Menlo":
            root.remove(f)
            return
    raise AssertionError("нет Menlo в fontTable")


@pytest.mark.parametrize("smp, added", [("06-blocks", True), ("05-lists", False)],
                         ids=["used", "unused"])
def test_fix_fonts_adds_missing_entry(tmp_path, outdir, smp, added):
    """Записи Menlo в fontTable нет: шрифт в тексте есть (код) — запись добавляется целиком,
    как у md2docx; нет — не добавляется, файл не тронут (SPEC §3.2)."""
    orig, _ = convert(outdir, sample(smp))
    doc = mutate(orig, tmp_path / "no-menlo.docx", "word/fontTable.xml", _drop_menlo)
    before = _font_table(doc)
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 0, proc.stderr[-2000:]
    if added:
        assert proc.stdout.strip().endswith("Menlo (altName, panose1, charset, family, pitch)"), \
            proc.stdout
        assert _font_table(doc) == _font_table(orig)
    else:
        assert proc.stdout.startswith("Без изменений:") and _font_table(doc) == before, proc.stdout


def test_fix_fonts_refuses_file_open_in_word(tmp_path, outdir):
    """Рядом замок Word (у имени длиннее 7 символов — без двух первых) — --fix-fonts файл
    не заменяет: код 1, файл цел; с --force — исправляет (SPEC §1)."""
    orig, _ = convert(outdir, sample("06-blocks"))
    doc = mutate(orig, tmp_path / "отчёт-за-год.docx", "word/fontTable.xml", _word_mac_resave)
    (tmp_path / "~$чёт-за-год.docx").write_bytes(b"owner")
    before = doc.read_bytes()
    proc = cli("--fix-fonts", doc)
    assert proc.returncode == 1 and doc.read_bytes() == before, proc.stdout + proc.stderr
    assert proc.stderr.strip() == (f"Ошибка: {doc}: файл открыт в Word (рядом ~$чёт-за-год.docx) — "
                                   f"закройте его или запустите с --force"), proc.stderr
    proc = cli("--fix-fonts", "--force", doc)
    assert proc.returncode == 0 and proc.stdout.startswith(f"Готово: {doc} —"), proc.stderr[-2000:]
    assert _font_table(doc) == _font_table(orig)


@pytest.mark.parametrize("mode", ["convert", "fix-fonts"])
def test_word_lock_next_to_symlink_target(tmp_path, outdir, mode):
    """Файл — символьная ссылка из другого каталога: замок Word рядом с файлом, на который
    она указывает, — отказ, файл цел; без замка ссылка остаётся ссылкой, заменяется её цель
    (SPEC §1)."""
    target, link = tmp_path / "real" / "отчёт.docx", tmp_path / "link" / "копия.docx"
    target.parent.mkdir()
    link.parent.mkdir()
    if mode == "convert":
        assert cli(write_md(target.parent, "отчёт", "Старый текст.\n")).returncode == 0
        args = (write_md(tmp_path, "новый", "Новый текст.\n"), "-o", link)
    else:
        orig, _ = convert(outdir, sample("06-blocks"))
        mutate(orig, target, "word/fontTable.xml", _word_mac_resave)
        args = ("--fix-fonts", link)
    link.symlink_to(target)
    lock = target.parent / "~$отчёт.docx"
    lock.write_bytes(b"owner")
    before = target.read_bytes()
    proc = cli(*args)
    assert proc.returncode == 1 and "файл открыт в Word (рядом ~$отчёт.docx)" in proc.stderr, \
        proc.stdout + proc.stderr
    assert target.read_bytes() == before
    lock.unlink()
    proc = cli(*args)
    assert proc.returncode == 0, proc.stderr[-2000:]
    assert link.is_symlink() and target.read_bytes() != before, "ссылка заменена файлом или цель не обновлена"


def test_fix_fonts_respects_readonly(tmp_path, outdir):
    """--fix-fonts не обходит защиту от записи: файл только для чтения — «нет доступа», код 1,
    и с --force, файл цел (SPEC §1)."""
    orig, _ = convert(outdir, sample("06-blocks"))
    doc = mutate(orig, tmp_path / "ro.docx", "word/fontTable.xml", _word_mac_resave)
    before = doc.read_bytes()
    doc.chmod(0o444)
    try:
        for flags in ((), ("--force",)):
            proc = cli("--fix-fonts", *flags, doc)
            assert proc.returncode == 1 and "нет доступа" in proc.stderr, proc.stderr[-2000:]
    finally:
        doc.chmod(0o644)
    assert doc.read_bytes() == before


FIX_SCRIPT = """
import json, sys
sys.path.insert(0, sys.argv[1])
from md2docx import fix_fonts
try:
    fix_fonts(sys.argv[2])
except FileExistsError as exc:
    print("FileExistsError:", exc)
for force in (True, False):
    print(json.dumps(fix_fonts(sys.argv[2], force=force), ensure_ascii=False))
"""


def test_api_fix_fonts(tmp_path, outdir):
    """fix_fonts() из README: при замке Word — FileExistsError с текстом CLI; force=True —
    словарь восстановленных полей; менять нечего — {} без отказа: файл не заменяется (SPEC §1)."""
    orig, _ = convert(outdir, sample("06-blocks"))
    doc = mutate(orig, tmp_path / "отчёт.docx", "word/fontTable.xml", _word_mac_resave)
    (tmp_path / "~$отчёт.docx").write_bytes(b"owner")
    proc = api(FIX_SCRIPT, doc, cwd=tmp_path)
    assert proc.returncode == 0, proc.stderr[-2000:]
    refused, fixed, again = proc.stdout.splitlines()
    assert refused == (f"FileExistsError: {doc}: файл открыт в Word (рядом ~$отчёт.docx) — "
                       f"закройте его или запустите с --force"), refused
    assert json.loads(fixed) == {"Helvetica Neue": ["altName", "family"], "Menlo": ["altName"]}
    assert json.loads(again) == {}


def test_fix_fonts_errors_continue_batch(tmp_path, outdir):
    """Не docx, битый zip (обрезан; испорчено сжатие; член зашифрован), оборванный XML части,
    docx без fontTable — ошибка этого файла по-русски, код 1, без трассировки; остальные файлы
    пакета исправляются (SPEC §1)."""
    orig, _ = convert(outdir, sample("06-blocks"))
    text = write_md(tmp_path, "текст", "Не docx.\n")
    cut = tmp_path / "обрезан.docx"
    cut.write_bytes(orig.read_bytes()[:2000])
    with zipfile.ZipFile(orig) as z:
        info = z.getinfo("word/document.xml")
    data = bytearray(orig.read_bytes())
    start = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra) + 10
    data[start:start + 30] = bytes(b ^ 0xFF for b in data[start:start + 30])
    spoiled = tmp_path / "сжатие.docx"
    spoiled.write_bytes(bytes(data))
    no_table = tmp_path / "без-шрифтов.docx"
    with zipfile.ZipFile(orig) as zin, zipfile.ZipFile(no_table, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename != "word/fontTable.xml":
                zout.writestr(item, zin.read(item.filename))
    torn = []  # архив цел, а XML части оборван на середине
    for name, part in (("обрыв-шрифтов.docx", "word/fontTable.xml"), ("обрыв-текста.docx", "word/document.xml")):
        torn.append(tmp_path / name)
        with zipfile.ZipFile(orig) as zin, zipfile.ZipFile(torn[-1], "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                zout.writestr(item, data[:len(data) // 2] if item.filename == part else data)
    locked = tmp_path / "пароль.docx"
    with zipfile.ZipFile(orig) as zin, zipfile.ZipFile(locked, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
            item.flag_bits |= 0x1  # «зашифрован»: бит попадёт в центральный каталог при закрытии
    good = mutate(orig, tmp_path / "хороший.docx", "word/fontTable.xml", _word_mac_resave)
    proc = cli("--fix-fonts", text, cut, spoiled, no_table, *torn, locked, good)
    assert proc.returncode == 1 and "Traceback" not in proc.stderr, proc.stderr[-2000:]
    assert proc.stderr.splitlines() == [
        f"Ошибка: {text}: не docx или файл повреждён",
        f"Ошибка: {cut}: не docx или файл повреждён",
        f"Ошибка: {spoiled}: не docx или файл повреждён",
        f"Ошибка: {no_table}: в файле нет word/fontTable.xml — это не docx",
        *(f"Ошибка: {p}: не docx или файл повреждён" for p in (*torn, locked)),
    ], proc.stderr
    assert proc.stdout.startswith(f"Готово: {good} —") and _font_table(good) == _font_table(orig)


def _repack(src: Path, dst: Path, part: str, data=None, size=None, method=zipfile.ZIP_DEFLATED) -> Path:
    """Копия docx, где член part записан с данными data и методом method, а в центральном
    каталоге заявлен размер size — его zipfile и считает размером члена."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if item.filename == part:
                item.compress_type = method
                body = body if data is None else data
            zout.writestr(item, body)
            if item.filename == part and size is not None:
                item.file_size = size  # попадёт в центральный каталог при закрытии
    return dst


BOMB_SCRIPT = """
import sys, tracemalloc
sys.path.insert(0, sys.argv[1])
from md2docx import fix_fonts
tracemalloc.start()
try:
    fix_fonts(sys.argv[2])
except ValueError as exc:
    print(exc)
print(tracemalloc.get_traced_memory()[1] >> 20)
"""


def test_fix_fonts_zip_bomb(tmp_path, outdir):
    """docx-бомба: по заявленным размерам архив больше 512 МБ или текст и стили больше 64 МБ —
    отказ до распаковки; член сжат не deflate (bzip2) — «не docx или файл повреждён»; член
    распаковывается дальше заявленного размера — читается не дальше него, память не растёт
    (SPEC §1)."""
    orig, _ = convert(outdir, sample("06-blocks"))
    total = _repack(orig, tmp_path / "архив.docx", "docProps/app.xml", size=600 << 20)
    markup = _repack(orig, tmp_path / "разметка.docx", "word/document.xml", size=65 << 20)
    bzip = _repack(orig, tmp_path / "bzip2.docx", "word/styles.xml", method=zipfile.ZIP_BZIP2)
    liar = _repack(orig, tmp_path / "нули.docx", "docProps/app.xml", data=bytes(64 << 20), size=10)
    before = {p: p.read_bytes() for p in (total, markup, bzip, liar)}
    proc = cli("--fix-fonts", total, markup, bzip, liar)
    assert proc.returncode == 1 and "Traceback" not in proc.stderr, proc.stderr[-2000:]
    assert proc.stderr.splitlines() == [
        f"Ошибка: {total}: в распакованном виде больше 512 МБ — слишком большой для docx",
        f"Ошибка: {markup}: текст и стили в распакованном виде больше 64 МБ — слишком большой для docx",
        f"Ошибка: {bzip}: не docx или файл повреждён",
        f"Ошибка: {liar}: не docx или файл повреждён",
    ], proc.stderr
    assert all(p.read_bytes() == data for p, data in before.items())
    proc = api(BOMB_SCRIPT, liar, cwd=tmp_path)  # 64 МБ нулей за заявленными 10 байтами
    assert proc.returncode == 0, proc.stderr[-2000:]
    error, peak_mb = proc.stdout.splitlines()
    assert error == "не docx или файл повреждён" and int(peak_mb) < 16, proc.stdout


@pytest.mark.parametrize("flags", [("-o", "x.docx"), ("-o", ""), ("--landscape",),
                                   ("--no-auto-landscape",)],
                         ids=["o", "empty-o", "landscape", "no-auto-landscape"])
def test_fix_fonts_incompatible_flags(tmp_path, flags):
    """--fix-fonts с -o, --landscape или --no-auto-landscape — ошибка аргументов, код 2,
    ничего не пишется (SPEC §1)."""
    proc = cli("--fix-fonts", "a.docx", *flags, cwd=tmp_path)
    assert proc.returncode == 2 and "--fix-fonts нельзя сочетать" in proc.stderr, proc.stderr
    assert not list(tmp_path.iterdir())


# --------------------------------------------------------------------------
# 8. Красная проба: испорченный docx сверщик обязан завалить
# --------------------------------------------------------------------------


def mutate(src: Path, dst: Path, part: str, fn) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        seen = False
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                root = etree.fromstring(data)
                fn(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                      standalone=True)
                seen = True
            zout.writestr(item, data)
    assert seen, f"в пакете нет {part}"
    return dst


def _drop_table_row(root):
    for tbl in root.iter(q("tbl")):
        rows = tbl.findall(q("tr"))
        if len(rows) >= 3:
            tbl.remove(rows[-1])
            return
    raise AssertionError("нет таблицы с 3+ строками")


def _replace_word(root):
    for t in root.iter(q("t")):
        m = re.search(r"\w{5,}", t.text or "")
        if m:
            t.text = t.text[:m.start()] + "ПОДМЕНА" + t.text[m.end():]
            return
    raise AssertionError("нет слова для подмены")


def _flip_orientation(root):
    for sp in root.iter(q("sectPr")):
        pg = sp.find(q("pgSz"))
        if wa(pg, "orient") == "landscape":
            w_, h_ = pg.get(q("w")), pg.get(q("h"))
            pg.set(q("w"), h_)
            pg.set(q("h"), w_)
            del pg.attrib[q("orient")]
            return
    raise AssertionError("нет альбомного раздела")


def _reorder_tblpr(root):
    tblpr = next(root.iter(q("tblPr")))
    layout, tblw = tblpr.find(q("tblLayout")), tblpr.find(q("tblW"))
    assert layout is not None and tblw is not None
    tblpr.remove(layout)
    tblw.addprevious(layout)


def _add_tbl_header(root):
    trpr = root.find(f".//{q('tr')}/{q('trPr')}")
    assert trpr is not None, "нет trPr у первой строки таблицы"
    etree.SubElement(trpr, q("tblHeader"))


def _compat_12(root):
    for cs in root.iter(q("compatSetting")):
        if wa(cs, "name") == "compatibilityMode":
            cs.set(q("val"), "12")
            return
    raise AssertionError("нет compatibilityMode")


def _drop_altname(root):
    for f in root.iter(q("font")):
        if wa(f, "name") == V.HN:
            alt = f.find(q("altName"))
            assert alt is not None
            f.remove(alt)
            return
    raise AssertionError("нет Helvetica Neue в fontTable")


def _inject_markup(root):
    for t in root.iter(q("t")):
        if (t.text or "").strip():
            t.text = "**" + t.text
            return


def _empty_section(root):
    body = root.find(q("body"))
    final = body.find(q("sectPr"))
    p = etree.SubElement(body, q("p"))
    ppr = etree.SubElement(p, q("pPr"))
    import copy
    ppr.append(copy.deepcopy(final))
    final.addprevious(p)


def _theme_heading_font(root):
    for st in root.iter(q("style")):
        if wa(st, "styleId") == "Heading1":
            rf = st.find("w:rPr/w:rFonts", NS)
            assert rf is not None
            rf.set(q("asciiTheme"), "majorHAnsi")
            return
    raise AssertionError("нет Heading1")


def _heading_to_normal(root):
    for ps in root.iter(q("pStyle")):
        if wa(ps, "val") == "Heading2":
            ps.getparent().remove(ps)
            return
    raise AssertionError("нет абзаца Heading2")


def _merge_lists(root):
    seen = []
    for np_ in root.iter(q("numPr")):
        nid = np_.find(q("numId"))
        if nid is not None and wa(nid, "val") not in seen:
            seen.append(wa(nid, "val"))
    assert len(seen) >= 2
    for nid in root.iter(q("numId")):
        if wa(nid, "val") == seen[1]:
            nid.set(q("val"), seen[0])


def _drop_content_type(root):
    for o in root:
        if (o.get("PartName") or "") == "/word/numbering.xml":
            root.remove(o)
            return
    raise AssertionError("нет Override для /word/numbering.xml")


def _undeclared_ignorable(root):
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
    root.set(mc, ((root.get(mc) or "") + " w99").strip())


def _broken_rel(root):
    for r in root:
        if (r.get("Target") or "") == "numbering.xml":
            r.set("Target", "numbering-missing.xml")
            return
    raise AssertionError("нет связи с numbering.xml")


RED_PROBES = [
    ("drop-table-row", "02-tables", "word/document.xml", _drop_table_row, {"tables", "words"}),
    ("replace-word", "01-headings", "word/document.xml", _replace_word, {"words"}),
    ("flip-orientation", "04-auto-landscape", "word/document.xml", _flip_orientation, {"orient"}),
    ("reorder-tblPr", "02-tables", "word/document.xml", _reorder_tblpr, {"order"}),
    ("add-tblHeader", "02-tables", "word/document.xml", _add_tbl_header, {"table-format"}),
    ("compat-12", "01-headings", "word/settings.xml", _compat_12, {"defaults"}),
    ("drop-altName", "01-headings", "word/fontTable.xml", _drop_altname, {"fonts"}),
    ("inject-markup", "01-headings", "word/document.xml", _inject_markup, {"markup"}),
    ("empty-section", "01-headings", "word/document.xml", _empty_section, {"sections"}),
    ("theme-heading-font", "01-headings", "word/styles.xml", _theme_heading_font,
     {"heading-styles"}),
    ("heading-to-normal", "01-headings", "word/document.xml", _heading_to_normal, {"headings"}),
    ("merge-lists", "05-lists", "word/document.xml", _merge_lists, {"lists"}),
    ("drop-content-type", "05-lists", "[Content_Types].xml", _drop_content_type, {"package"}),
    ("undeclared-ignorable", "01-headings", "word/document.xml", _undeclared_ignorable,
     {"package"}),
    ("broken-rel", "05-lists", "word/_rels/document.xml.rels", _broken_rel, {"package"}),
]


@pytest.mark.parametrize("name, smp, part, fn, codes", RED_PROBES, ids=[r[0] for r in RED_PROBES])
def test_red_probe(name, smp, part, fn, codes, outdir, tmp_path):
    md = sample(smp)
    good, _ = convert(outdir, md)
    exp = EXPECT_ORIENT[smp]
    base = V.run_checks(md, good, exp)
    assert base.ok, ("исходный docx уже красный — красная проба ничего не докажет:\n"
                     + base.render())
    bad = mutate(good, tmp_path / f"{name}.docx", part, fn)
    res = run_verify(md, bad, "--expect-orient", exp)
    assert res.returncode == 1, f"сверщик не заметил порчу «{name}»:\n{res.stdout}"
    rep = V.run_checks(md, bad, exp)
    missing = codes - rep.codes()
    assert not missing, f"порча «{name}»: нет расхождений категорий {missing}:\n{rep.render()}"


REFERENCE_DEVIATIONS = {
    "таблица шире текста": ("table-format", "шире текста"),
    "нет tblLayout fixed": ("table-format", "нет tblLayout type=fixed"),
    "вводный абзац со стилем заголовка": ("heading-styles", "абзац со стилем заголовка"),
    "pgSz 16820×11900": ("page", "pgSz 16820×11900"),
    "compatibilityMode 12": ("defaults", "compatibilityMode=12"),
}


def test_verify_catches_reference_deviations():
    ref = os.environ.get("MD2DOCX_REFERENCE_DOCX")
    if not ref or not Path(ref).is_file():
        pytest.skip("эталон не задан: MD2DOCX_REFERENCE_DOCX=<путь к эталонному docx>")
    res = run_verify("--invariants-only", ref)
    assert res.returncode == 1, res.stdout
    rep = V.run_checks(None, ref, invariants_only=True)
    lost = [k for k, (code, needle) in REFERENCE_DEVIATIONS.items() if not rep.has(code, needle)]
    assert not lost, f"сверщик не нашёл известные отклонения эталона: {lost}\n{rep.render()}"
    assert not rep.codes() & {"order", "structure", "internal"}, \
        "файл Word не должен давать ошибок порядка/структуры — иначе сверщик врёт:\n" + rep.render()


# --------------------------------------------------------------------------
# 9. Реальные документы (пути — в tests/real_docs.local.txt, не в git)
# --------------------------------------------------------------------------


def _real_docs():
    if not REAL_DOCS_LIST.is_file():
        return []
    return [Path(x.strip()) for x in REAL_DOCS_LIST.read_text(encoding="utf-8").splitlines()
            if x.strip() and not x.lstrip().startswith("#")]


REAL_DOCS = _real_docs()


@pytest.mark.skipif(not REAL_DOCS, reason="нет tests/real_docs.local.txt — регрессия не задана")
@pytest.mark.parametrize("idx", range(len(REAL_DOCS)), ids=[f"real-{i + 1}" for i in
                                                            range(len(REAL_DOCS))])
def test_real_document(idx, tmp_path):
    md = REAL_DOCS[idx]
    if not md.is_file():
        pytest.skip(f"real-{idx + 1}: файла из списка нет на диске")
    out = tmp_path / f"real-{idx + 1}.docx"
    try:
        proc = cli(md, "-o", out)
        assert proc.returncode == 0, f"real-{idx + 1}: код {proc.returncode}\n{proc.stderr[-2000:]}"
        res = run_verify(md, out)
        assert res.returncode == 0, f"real-{idx + 1}: сверщик нашёл расхождения:\n{res.stdout}"
    finally:  # результат — копия чужого документа: на диске не оставляем и при падении
        out.unlink(missing_ok=True)
        left = sorted(x.name for x in tmp_path.iterdir())
        for x in tmp_path.iterdir():
            shutil.rmtree(x) if x.is_dir() else x.unlink()
    assert not left, f"real-{idx + 1}: после сверки в каталоге теста осталось {left}"


# --------------------------------------------------------------------------
# 10. Рендер LibreOffice → PDF (SPEC §6): ориентация страниц, пустые страницы,
#     слова в ячейках, границы таблиц по полям, подмена шрифтов как на Windows
#     (Helvetica Neue → Arial, Menlo → Courier New)
# --------------------------------------------------------------------------


def _soffice():
    cand = [os.environ.get("MD2DOCX_SOFFICE"), shutil.which("soffice"),
            "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for c in cand:
        if c and Path(c).is_file():
            return c
    return None


SOFFICE = _soffice()
POPPLER = all(shutil.which(x) for x in ("pdfinfo", "pdftotext", "pdffonts", "pdftoppm"))
need_render = pytest.mark.skipif(not (SOFFICE and POPPLER),
                                 reason="нет LibreOffice или poppler — рендер не проверяется")
RENDER_SAMPLES = ("02-tables", "03-orientation-markers", "04-auto-landscape", "07-pagebreak")
# Синтетические входы рендера (A1.1, A1.5, A1.7, A1.9). Длина таблицы (tail-N) и число абзацев
# перед ней (fill-N) перебираются, чтобы граница страницы прошла через начало и конец таблицы при
# любых метриках шрифта: на 0.1.1 в этих диапазонах были и одинокая строка, и пустая страница
TAIL_ROWS, FILL_PARAS = range(34, 43), range(17, 29)
RENDER_SYNTH = {
    "codes": "# Коды и даты\n\n" + DATES_TABLE,
    "code-plates": ("# Код\n\nАбзац перед кодом.\n\n```bash\necho first\nls -la\n```\n\n"
                    "```python\nprint(\"second\")\n```\n\nАбзац между блоками.\n\n"
                    "```text\nодиночный блок\n```\n\nАбзац после кода.\n"),
    **{f"tail-{n}": "# Приложение\n\n| № | Пункт | Срок |\n|---|---|---|\n"
                    + "".join(f"| {i} | Пункт приложения номер {i} | 2026-10-22 |\n" for i in range(n))
       for n in TAIL_ROWS},
    **{f"fill-{n}": "# Контроль\n\n"
                    + "".join(f"Абзац-заполнитель номер {i}.\n\n" for i in range(1, n + 1))
                    + "| Строка | Описание |\n|---|---|\n"
                    + "".join(f"| Строка0{i} | первая<br>вторая<br>третья |\n" for i in range(1, 7))
       for n in FILL_PARAS},
    # \newpage посреди таблицы: части на разных страницах (§3.4); колонки разной ширины
    "split": ("# Разрез таблицы\n\n| Номер | Описание пункта |\n|---|---|\n"
              + "".join(f"| {i} | Пункт первой части {i}, текст пошире номера |\n" for i in (1, 2, 3))
              + "\\newpage\n"
              + "".join(f"| {i} | Пункт второй части {i}, текст пошире номера |\n" for i in (4, 5, 6))),
}


def _substitute_fonts(src: Path, dst: Path):
    """Копия docx, где Helvetica Neue → Arial и Menlo → Courier New — как на Windows."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                s = data.decode("utf-8")
                s = s.replace('"Helvetica Neue"', '"Arial"').replace('"Menlo"', '"Courier New"')
                data = s.encode("utf-8")
            zout.writestr(item, data)


@pytest.fixture(scope="session")
def rendered(outdir, tmp_path_factory):
    """{(образец, вариант): pdf} — один запуск soffice на все файлы."""
    if not (SOFFICE and POPPLER):
        pytest.skip("нет LibreOffice или poppler")
    pytest.importorskip("PIL", reason="нужен Pillow для замера границ таблиц на растре")
    src_dir = tmp_path_factory.mktemp("render-src")
    pdf_dir = tmp_path_factory.mktemp("render-pdf")
    profile = Path(os.environ.get("MD2DOCX_LO_PROFILE") or tmp_path_factory.mktemp("lo-profile"))
    md_dir = tmp_path_factory.mktemp("render-md")
    sources = [(name, sample(name), None) for name in RENDER_SAMPLES]
    for name, text in RENDER_SYNTH.items():
        (md_dir / f"{name}.md").write_text(text, encoding="utf-8")
        sources.append((name, md_dir / f"{name}.md", f"render-{name}"))
    files = {}
    for name, md, stem in sources:
        docx, _ = convert(outdir, md, stem=stem)
        a = src_dir / f"{name}.docx"
        shutil.copy(docx, a)
        b = src_dir / f"{name}.arial.docx"
        _substitute_fonts(docx, b)
        files[(name, "hn")], files[(name, "arial")] = a, b
    proc = subprocess.run([SOFFICE, "--headless", f"-env:UserInstallation={profile.as_uri()}",
                           "--convert-to", "pdf", "--outdir", str(pdf_dir),
                           *map(str, files.values())], capture_output=True, text=True, timeout=600)
    out = {}
    for key, f in files.items():
        pdf = pdf_dir / (f.stem + ".pdf")
        assert pdf.is_file(), f"LibreOffice не отрендерил {f.name}: {proc.stdout} {proc.stderr}"
        out[key] = pdf
    return out


def _pages(pdf: Path):
    info = subprocess.run(["pdfinfo", "-f", "1", "-l", "9999", str(pdf)], capture_output=True,
                          text=True).stdout
    sizes = re.findall(r"Page\s+\d+\s+size:\s+([\d.]+)\s+x\s+([\d.]+)", info)
    return ["L" if float(w_) > float(h_) else "P" for w_, h_ in sizes]


def _collapse(seq):
    out = []
    for x in seq:
        if not out or out[-1] != x:
            out.append(x)
    return out


def _page_text(pdf: Path, i: int) -> str:
    return subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), "-layout", str(pdf), "-"],
                          capture_output=True, text=True).stdout


MARGIN_PT = 1134 / 20  # поля 1134 twips = 56,7 pt
DPI = 144              # 2 px на пункт


def _border_x(pdf: Path, page: int):
    """(левая, правая) x в pt крайних вертикальных линий цвета границ таблицы (CCCCCC)
    и ширина страницы в pt; (None, None, w), если таких линий на странице нет."""
    from PIL import Image
    prefix = pdf.parent / f"{pdf.stem}-p{page}"
    subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-singlefile", "-r", str(DPI),
                    "-gray", "-png", str(pdf), str(prefix)], check=True, capture_output=True)
    im = Image.open(f"{prefix}.png").convert("L")
    width_px, height_px = im.size
    # серый 150…235: граница CCCCCC и её сглаживание; текст темнее, заливка шапки F2F2F2 светлее
    cols = im.point(lambda v: 255 if 150 <= v <= 235 else 0).transpose(Image.Transpose.TRANSPOSE)
    data, run = cols.tobytes(), re.compile(rb"\xff{40,}")  # 40 px = 20 pt — выше строки текста
    xs = [x for x in range(width_px) if run.search(data, x * height_px, (x + 1) * height_px)]
    k = DPI / 72
    if not xs:
        return None, None, width_px / k
    return min(xs) / k, max(xs) / k, width_px / k


def _occurrences(word: str, text: str) -> int:
    return len(re.findall(rf"(?<!\w){re.escape(word)}(?!\w)", text))


def _cell_words(md_path: Path):
    md = V.parse_md(md_path)
    words = set()
    for t in md.tables:
        for row in t.rows:
            for cell in row:
                for wd in re.split(r"\s+", cell):
                    for piece in re.findall(r"[^-‐–—]+", wd):
                        piece = piece.strip(".,:;!?()«»\"'|*`~")
                        if 2 <= len(piece) <= 30:
                            words.add(piece)
    return words


@need_render
@pytest.mark.parametrize("name", RENDER_SAMPLES)
@pytest.mark.parametrize("variant", ["hn", "arial"])
def test_render(rendered, name, variant):
    pdf = rendered[(name, variant)]
    pages = _pages(pdf)
    exp = V.parse_orient_list(EXPECT_ORIENT[name])
    assert _collapse(pages) == _collapse(exp), f"{name}/{variant}: страницы {pages}, разделы {exp}"
    for i in range(1, len(pages) + 1):
        assert _page_text(pdf, i).strip(), f"{name}/{variant}: страница {i} пустая"
    if name == "07-pagebreak":
        assert pages == ["P"] * 5 + ["L"], f"07: ожидалось 5 книжных и 1 альбомная, а {pages}"
    flat = re.sub(r"\s+", " ", subprocess.run(["pdftotext", "-raw", str(pdf), "-"],
                                              capture_output=True, text=True).stdout)
    # слово ищется не «есть ли где-то», а сколько раз: разорванное в одной таблице и
    # целое в другой иначе проходит незамеченным
    md_text = V.parse_md(sample(name)).text
    broken = sorted(w_ for w_ in _cell_words(sample(name))
                    if _occurrences(w_, flat) < _occurrences(w_, md_text))
    assert not broken, f"{name}/{variant}: слова из ячеек разорваны или потеряны: {broken[:10]}"
    with_tables = 0
    for i in range(1, len(pages) + 1):
        left, right, page_w = _border_x(pdf, i)
        if left is None:
            continue
        with_tables += 1
        assert abs(left - MARGIN_PT) <= 1.0 and abs(right - (page_w - MARGIN_PT)) <= 1.0, \
            (f"{name}/{variant}, стр. {i}: границы таблицы x={left:.1f}…{right:.1f} pt, "
             f"поля — {MARGIN_PT:.1f}…{page_w - MARGIN_PT:.1f} pt")
    if V.parse_md(sample(name)).tables:
        assert with_tables, f"{name}/{variant}: на страницах не найдено ни одной границы таблицы"
    if variant == "arial":
        fonts = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
        assert "Arial" in fonts and "HelveticaNeue" not in fonts.replace(" ", ""), fonts
        base = _pages(rendered[(name, "hn")])
        assert pages == base, f"{name}: с Arial страниц {pages}, с Helvetica Neue {base}"


def _vlines(pdf: Path, page: int) -> list:
    """x в pt вертикальных линий цвета границ таблицы (CCCCCC) на странице: соседние столбцы
    пикселей одной линии слиты в одну."""
    from PIL import Image
    prefix = pdf.parent / f"{pdf.stem}-v{page}"
    subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-singlefile", "-r", str(DPI),
                    "-gray", "-png", str(pdf), str(prefix)], check=True, capture_output=True)
    im = Image.open(f"{prefix}.png").convert("L")
    width_px, height_px = im.size
    cols = im.point(lambda v: 255 if 150 <= v <= 235 else 0).transpose(Image.Transpose.TRANSPOSE)
    data, run = cols.tobytes(), re.compile(rb"\xff{40,}")
    xs = [x for x in range(width_px) if run.search(data, x * height_px, (x + 1) * height_px)]
    lines = []
    for x in xs:
        if lines and x - lines[-1][-1] <= 2:
            lines[-1].append(x)
        else:
            lines.append([x])
    return [sum(g) / len(g) / (DPI / 72) for g in lines]


def _raw_text(pdf: Path, page: int | None = None) -> str:
    rng = ["-f", str(page), "-l", str(page)] if page else []
    return subprocess.run(["pdftotext", *rng, "-raw", str(pdf), "-"], capture_output=True,
                          text=True).stdout


@need_render
@pytest.mark.parametrize("variant", ["hn", "arial"])
def test_render_codes_not_split(rendered, variant):
    """A1.1: даты и коды с дефисом в узких колонках книжной таблицы, которая влезает, не рвутся."""
    pdf = rendered[("codes", variant)]
    assert set(_pages(pdf)) == {"P"}, f"codes/{variant}: страницы {_pages(pdf)}"
    flat = re.sub(r"\s+", " ", _raw_text(pdf))
    got = {tok: _occurrences(tok, flat) for tok in DATE_TOKENS}
    assert got == dict.fromkeys(DATE_TOKENS, 3), \
        f"codes/{variant}: целых вхождений {got}, в таблице каждого по 3 — остальные разорваны"


@need_render
@pytest.mark.parametrize("kind, row_re, rows", [
    ("tail", r"Пункт приложения номер \d+", None),
    ("fill", r"Строка0\d", 6),
], ids=["tail", "fill"])
@pytest.mark.parametrize("variant", ["hn", "arial"])
def test_render_table_rows_not_alone(rendered, variant, kind, row_re, rows):
    """A1.5, A1.7: таблица не начинается внизу страницы шапкой с одной строкой тела, последняя
    строка не уходит на новую страницу одна, пустой последней страницы нет."""
    bad = []
    for name in RENDER_SYNTH:
        if not name.startswith(kind + "-"):
            continue
        pdf = rendered[(name, variant)]
        n_pages = len(_pages(pdf))
        per_page = [len(re.findall(row_re, _raw_text(pdf, i))) for i in range(1, n_pages + 1)]
        empty = [i for i in range(1, n_pages + 1) if not _page_text(pdf, i).strip()]
        want = rows or int(name.rsplit("-", 1)[1])
        if empty or any(k == 1 for k in per_page) or sum(per_page) != want:
            bad.append(f"{name}: строк тела по страницам {per_page}"
                       + (f", пустые страницы {empty}" if empty else ""))
    assert not bad, f"{kind}/{variant}: " + "; ".join(bad)


@need_render
@pytest.mark.parametrize("variant", ["hn", "arial"])
def test_render_split_table(rendered, variant):
    """\\newpage посреди таблицы: части на разных страницах, сетка одна (вертикальные линии на тех
    же x), шапка (полоса заливки) — только у первой части, текст всех строк на месте (SPEC §3.4)."""
    pdf = rendered[("split", variant)]
    assert _pages(pdf) == ["P", "P"], f"split/{variant}: страницы {_pages(pdf)}"
    p1, p2 = (re.sub(r"\s+", " ", _raw_text(pdf, i)) for i in (1, 2))
    assert re.findall(r"Пункт (\w+) части (\d)", p1) == [("первой", str(i)) for i in (1, 2, 3)], p1
    assert re.findall(r"Пункт (\w+) части (\d)", p2) == [("второй", str(i)) for i in (4, 5, 6)], p2
    assert "Описание пункта" in p1 and "Описание пункта" not in p2, "шапка не повторяется"
    v1, v2 = _vlines(pdf, 1), _vlines(pdf, 2)
    assert len(v1) == 3 and len(v2) == 3 and all(abs(a - b) <= 1 for a, b in zip(v1, v2)), \
        f"split/{variant}: вертикальные линии {v1} и {v2} — сетка у частей одна"
    assert len(_plates(pdf, 1)[0]) == 1 and not _plates(pdf, 2)[0], "заливка шапки — только на стр. 1"


def _plates(pdf: Path, page: int = 1):
    """Серые плашки кода на странице: [(верх, низ, левый край, правый край, левый край текста)]
    в pt и ширина страницы в pt."""
    from PIL import Image
    prefix = pdf.parent / f"{pdf.stem}-plates-p{page}"
    subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-singlefile", "-r", str(DPI),
                    "-gray", "-png", str(pdf), str(prefix)], check=True, capture_output=True)
    im = Image.open(f"{prefix}.png").convert("L")
    w, h = im.size
    # заливка F2F2F2 — серый 242 (236…248 со сглаживанием); текст кода темнее 128
    fill = im.point(lambda v: 255 if 236 <= v <= 248 else 0).tobytes()
    dark = im.point(lambda v: 255 if v < 128 else 0).tobytes()
    wide = re.compile(rb"\xff{200,}")  # 200 px = 100 pt: строка плашки, а не просвет между буквами
    bands = []
    for y in range(h):
        row = fill[y * w:(y + 1) * w]
        if not wide.search(row):
            continue
        left, right = row.find(b"\xff"), row.rfind(b"\xff")
        text_x = dark.find(b"\xff", y * w, (y + 1) * w)
        text_x = text_x - y * w if text_x >= 0 else w
        if bands and bands[-1][1] == y - 1:
            b = bands[-1]
            b[1:] = [y, min(b[2], left), max(b[3], right), min(b[4], text_x)]
        else:
            bands.append([y, y, left, right, text_x])
    k = DPI / 72
    return [tuple(v / k for v in b) for b in bands], w / k


@need_render
@pytest.mark.parametrize("variant", ["hn", "arial"])
def test_render_code_plates(rendered, variant):
    """A1.9: три блока кода — три плашки (соседние не сливаются, строки одного блока — одна
    плашка), текст отстоит от края плашки, плашка не выходит за поля."""
    plates, page_w = _plates(rendered[("code-plates", variant)])
    where = [f"{top:.0f}…{bottom:.0f}" for top, bottom, *_ in plates]
    assert len(plates) == 3, f"code-plates/{variant}: плашек {len(plates)} (по высоте, pt: {where})"
    gaps = [round(b[0] - a[1], 1) for a, b in zip(plates, plates[1:])]
    assert min(gaps) >= 4, f"code-plates/{variant}: промежутки между плашками {gaps} pt"
    for top, _, left, right, text_x in plates:
        assert text_x - left >= 3, \
            f"code-plates/{variant}: плашка на {top:.0f} pt — текст в {text_x - left:.1f} pt от края"
        assert left >= MARGIN_PT - 1.5 and right <= page_w - MARGIN_PT + 1.5, \
            f"code-plates/{variant}: плашка x={left:.1f}…{right:.1f} pt выходит за поля"
