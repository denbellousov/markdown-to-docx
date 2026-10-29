#!/usr/bin/env python3
"""md2docx — локальный конвертер Markdown → DOCX.

Аккуратные заголовки и таблицы, книжная и альбомная ориентация страниц.
Контракт поведения — SPEC.md. Запуск: python3 md2docx.py doc.md [-o doc.docx]
"""
from __future__ import annotations

import argparse
import contextvars
import datetime as dt
import errno
import functools
import itertools
import json
import locale
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace
from typing import NamedTuple
from urllib.parse import unquote, urlparse

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.image.image import Image
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.oxml import serialize_part_xml
from docx.opc.part import PartFactory, XmlPart
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement, parse_xml
from docx.shared import Twips
from lxml import etree
from markdown_it import MarkdownIt
from markdown_it.common.utils import unescapeAll
from markdown_it.rules_block import reference
from markdown_it.tree import SyntaxTreeNode

try:  # Pillow нужен только для точного замера ширины текста в таблицах
    from PIL import ImageFont
except ImportError:  # pragma: no cover
    ImageFont = None

__version__ = "1.1.0"

# ---------------------------------------------------------------- константы

PAGE = {"portrait": (11906, 16838), "landscape": (16838, 11906)}  # A4, twips
MARGIN, HDR_FTR = 1134, 720
TEXT_W = {o: w - 2 * MARGIN for o, (w, _) in PAGE.items()}  # 9638 / 14570
TEXT_H = {o: h - 2 * MARGIN for o, (_, h) in PAGE.items()}  # 14570 / 9638
LINE_H = 262      # шаг строки ячейки: HN 11 pt при line=240 — 13,1 pt (у Arial меньше)
LETTER = {"portrait": "К", "landscape": "А"}  # convert()["sections"] — как в 1.0.0, на любом языке
STEP = 360       # отступ уровня списка и цитаты
MAX_INDENT = 8 * STEP  # глубже 8 уровней отступ не растёт: текст и таблица не уходят за правое поле
CELL_PAD = 216    # поля ячейки 108 + 108
HN, MONO = "Helvetica Neue", "Menlo"
HEADINGS = {1: (36, 360), 2: (32, 300), 3: (28, 260), 4: (24, 240), 5: (22, 220), 6: (22, 200)}
BULLETS = ("•", "◦", "▪")
SHADE = {"w:val": "clear", "w:color": "auto", "w:fill": "F2F2F2"}
BORDER = {"w:val": "single", "w:sz": 4, "w:space": 0, "w:color": "CCCCCC"}
TBL_LOOK = {"w:val": "04A0", "w:firstRow": 1, "w:lastRow": 0, "w:firstColumn": 1,
            "w:lastColumn": 0, "w:noHBand": 0, "w:noVBand": 1}
# Запись fontTable: Word на Windows, не найдя шрифт, берёт altName и метрики (§3.2)
FONT_ENTRIES = (
    (HN, (("altName", {"w:val": "Arial"}), ("panose1", {"w:val": "02000503000000020004"}),
          ("charset", {"w:val": "00"}), ("family", {"w:val": "swiss"}),
          ("pitch", {"w:val": "variable"}),
          ("sig", {"w:usb0": "E50002FF", "w:usb1": "500079DB", "w:usb2": "00000010",
                   "w:usb3": "00000000", "w:csb0": "00000001", "w:csb1": "00000000"}))),
    (MONO, (("altName", {"w:val": "Consolas"}), ("panose1", {"w:val": "020B0609030804020204"}),
            ("charset", {"w:val": "00"}), ("family", {"w:val": "modern"}),
            ("pitch", {"w:val": "fixed"}))),
)
FONT_ORDER = ("altName panose1 charset family notTrueType pitch sig embedRegular embedBold "
              "embedItalic embedBoldItalic").split()  # дети w:font по схеме ECMA-376
FONT_TABLE = "word/fontTable.xml"
MARKUP_PARTS = (FONT_TABLE, "word/document.xml", "word/styles.xml", "word/numbering.xml")  # их разбирает fix_fonts
UNPACKED_LIMIT = 512 * 1024 * 1024  # fix_fonts: предел распакованного docx (у файла Word — 512 МБ)
MARKUP_LIMIT = 64 * 1024 * 1024  # и частей MARKUP_PARTS: дерево lxml в разы больше самого XML

# Порядок детей по схеме ECMA-376: Word объявляет файл повреждённым при нарушении (§5)
PPR_ORDER = (
    "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers "
    "pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE "
    "autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents "
    "suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle "
    "rPr sectPr pPrChange").split()
RPR_ORDER = (
    "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint "
    "noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u "
    "effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath").split()


# ---------------------------------------------------------------- язык интерфейса

# Код → самоназвание для переключателя; en — исходный язык: английские строки кода и есть ключи
# переводов locales/<код>.json, строка без перевода выводится по-английски
LANGUAGES = {"ar": "العربية", "de": "Deutsch", "en": "English", "es": "Español", "fr": "Français",
             "hi": "हिन्दी", "ja": "日本語", "pt": "Português", "ru": "Русский", "zh": "简体中文"}
RTL = frozenset({"ar"})  # пишутся справа налево
LOCALES =Path(__file__).resolve().parent / "locales"
LANG = contextvars.ContextVar("md2docx_lang", default=None)  # None — язык системы; сервер задаёт свой на запрос


@functools.cache
def catalog(lang):
    """Переводы языка lang: {английская строка: перевод}; нет файла — пустой."""
    try:
        return json.loads((LOCALES / f"{lang}.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def pick_language(tags):
    """Первый поддерживаемый язык из тегов вида ru_RU.UTF-8, pt-BR, zh-Hans; иначе en."""
    for tag in tags:
        code = re.split(r"[-_.@]", tag.strip(), maxsplit=1)[0].lower()
        if code in LANGUAGES:
            return code
    return "en"


def os_languages():
    """Языки интерфейса ОС по предпочтению: у приложения, открытого из Finder или Проводника,
    переменных локали нет."""
    try:
        if sys.platform == "darwin":
            out = subprocess.run(["defaults", "read", "-g", "AppleLanguages"], capture_output=True,
                                 text=True, timeout=5).stdout
            return re.findall(r"[A-Za-z]{2,3}(?:[-_][A-Za-z0-9]+)*", out)
        if sys.platform == "win32":
            import ctypes
            return [locale.windows_locale.get(ctypes.windll.kernel32.GetUserDefaultUILanguage(), "")]
    except (OSError, subprocess.SubprocessError):
        pass
    return []


@functools.cache
def system_language():
    """MD2DOCX_LANG, затем LANGUAGE, LC_ALL, LC_MESSAGES, LANG, затем язык ОС."""
    tags = [os.environ.get("MD2DOCX_LANG", ""), *os.environ.get("LANGUAGE", "").split(":")]
    tags += [os.environ.get(name, "") for name in ("LC_ALL", "LC_MESSAGES", "LANG")]
    tags = [t for t in tags if t.strip() and t.split(".")[0] not in ("C", "POSIX")]
    return pick_language(tags or os_languages())


def language():
    """Текущий язык интерфейса: заданный для этого потока или системный."""
    return LANG.get() or system_language()


def tr(text, **values):
    """Строка интерфейса на текущем языке; values подставляются в её {имена}."""
    lang = language()
    if lang != "en":
        text = catalog(lang).get(text, text)
    if LANG.get() in RTL:  # страница справа налево: строки и пути изолированы, иначе «02-a.md» видно
        # как «a.md-02»; в терминале (LANG не задан) — без невидимых символов: они попали бы в скопированный путь
        values = {k: v if isinstance(v, (int, float)) else f"⁨{v}⁩" for k, v in values.items()}
    return text.format(**values) if values else text


def summary(stats):
    """Итог конвертации — для «Готово: …» в CLI и для списка результатов в браузере."""
    letters = ", ".join(tr("L") if s == LETTER["landscape"] else tr("P") for s in stats["sections"])
    return tr("sections: {sections} ({letters}), tables: {tables}, auto-landscape: {auto}",
              sections=len(stats["sections"]), letters=letters, tables=stats["tables"],
              auto=stats["auto_landscape_tables"])


def fonts_summary(changed):
    """Итог --fix-fonts: какие подсказки каким шрифтам возвращены."""
    fixed = ", ".join(f"{name} ({', '.join(tags)})" for name, tags in changed.items())
    return tr("font hints restored: {fonts}", fonts=fixed)
STYLE_ORDER = (
    "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed "
    "qFormat locked personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr "
    "tblStylePr").split()

# maxNesting: в пресете commonmark 20 — после списка глубины 10 остаток файла молча терялся
MD = MarkdownIt("commonmark", {"html": True, "maxNesting": 100}).enable(["table", "strikethrough"])
# схемы отбирает flatten(): по чёрному списку markdown-it [x](javascript:…) выходил разметкой в тексте
MD.validateLink = lambda url: True
NEST_OPEN = ("blockquote_open", "bullet_list_open", "ordered_list_open", "list_item_open")
SAFE_SCHEMES = ("http", "https", "mailto", "tel")  # гиперссылка — только для них (§3.6)
MAX_IMAGE_BYTES = 50 * 1024 * 1024  # больше — не картинка для документа: в память не читаем
MARK = "md2docx"  # dc:identifier наших файлов: без --force перезаписываются только они (§1)
FLAGS = {"strong": "b", "em": "i", "s": "strike"}
BOLD = frozenset({"b"})
MONO_FMT = frozenset({"mono"})
BREAKS = ("softbreak", "hardbreak")
PAGEBREAK_LINES = ("\\newpage", "\\pagebreak")
PAGEBREAK = SimpleNamespace(type="pagebreak", map=None, children=[])  # разрыв страницы в плане

BR_RE = re.compile(r"</?br\s*/?>", re.I)
COMMENT_RE = re.compile(r"<!--.*?(?:-->|\Z)", re.S)  # незакрытый — до конца блока, как в браузере
TAG_RE = re.compile(r"</?[A-Za-z][^>]*>|<![^>]*>|<\?.*?\?>", re.S)
HIDDEN_RE = re.compile(r"<(script|style|template)\b.*?</\1\s*>", re.S | re.I)  # не текст документа
BLOCK_TAG_RE = re.compile(r"</?(?:p|div|li|ul|ol|tr|td|th|h[1-6]|table|thead|tbody|blockquote|pre|hr|"
                          r"section|article|header|footer|details|summary|dl|dt|dd)\b[^>]*>", re.I)
INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
VT_FF = re.compile(r"&#(?:0*1[12]|[xX]0*[bBcC]);")  # сущности &#11; и &#12; в любой записи
WORD_GAP = re.compile(r"[ \t\u2002\u2003\u2009\u200b]+")
# перенос после дефиса и тире (§3.4); после «-» перед цифрой — только за цифрой: «2026-|10», а «Ф-123»
# LibreOffice не переносит
AFTER_DASH = re.compile(r"(?<=[‐–—])|(?<=-)(?!\d)|(?<=\d-)")
URL_TOKEN = re.compile(r"://|\bwww\.", re.I)  # URL по дефисам не делится: он целиком длиннее 30 (§3.4)
KEY_LINE = re.compile(r"[^\s:#][^:]*:(\s|$)")
FM_LINE = re.compile(r"\s|#|-\s|[^\s:#][^:]*:(\s|$)")


class Ctx(NamedTuple):
    """Контекст блока: отступ контейнера (twips), уровень цитаты, глубина списков, внутри пункта."""
    indent: int = 0
    quote: int = 0
    depth: int = 0
    item: bool = False


# ---------------------------------------------------------------- XML-помощники

def E(tag, attrs=None, *children):
    el = OxmlElement(tag)
    for key, val in (attrs or {}).items():
        el.set(qn(key), str(val))
    for child in children:
        el.append(child)
    return el


def local(el):
    return el.tag.rsplit("}", 1)[-1]


def ordered(tag, order, children):
    """Контейнер свойств, дети — в порядке схемы."""
    el = OxmlElement(tag)
    for child in sorted(children, key=lambda c: order.index(local(c))):
        el.append(child)
    return el


def insert_ordered(parent, child, order):
    """Вставить (или заменить) дочерний элемент на его место по схеме."""
    old = parent.find(child.tag)
    if old is not None:
        parent.remove(old)
    pos = order.index(local(child))
    for sib in parent:
        if local(sib) in order and order.index(local(sib)) > pos:
            sib.addprevious(child)
            return child
    parent.append(child)
    return child


def spacing(before=None, after=None, line=None, rule="auto"):
    attrs = {}
    if before is not None:
        attrs["w:before"] = before
    if after is not None:
        attrs["w:after"] = after
    if line is not None:
        attrs.update({"w:line": line, "w:lineRule": rule})
    return E("w:spacing", attrs)


def ind(left, hanging=None):
    return E("w:ind", {"w:left": left, **({"w:hanging": hanging} if hanging else {})})


def rfonts(name):
    return E("w:rFonts", {"w:ascii": name, "w:hAnsi": name, "w:cs": name})


def pstyle(style_id):
    return E("w:pStyle", {"w:val": style_id})


def tiny_rpr():
    """Знак абзаца в 1 pt — для абзаца-носителя разрыва."""
    return E("w:rPr", None, E("w:sz", {"w:val": 2}), E("w:szCs", {"w:val": 2}))


def make_p(ppr=()):
    p = OxmlElement("w:p")
    if ppr:
        p.append(ordered("w:pPr", PPR_ORDER, ppr))
    return p


def spacer():
    """Пустой абзац после таблицы: иначе Word склеит две таблицы подряд в одну."""
    return make_p([spacing(0, 0, 160, "exact")])


def run_props(fmt):
    kids = []
    if "link" in fmt:
        kids.append(E("w:rStyle", {"w:val": "Hyperlink"}))
    if "arial" in fmt:  # глифа нет в HN (→): Word для Mac иначе рисует его Times New Roman
        kids.append(rfonts("Arial"))
    if "code" in fmt or "mono" in fmt:
        kids.append(rfonts(MONO))
        if "head" not in fmt:  # в заголовке код — кеглем заголовка (из стиля)
            kids += [E("w:sz", {"w:val": 20}), E("w:szCs", {"w:val": 20})]
        if "code" in fmt:
            kids.append(E("w:shd", SHADE))
    for flag in ("b", "i", "strike"):
        if flag in fmt:
            kids.append(E("w:" + flag))
    if "gray" in fmt:
        kids.append(E("w:color", {"w:val": "808080"}))
    return ordered("w:rPr", RPR_ORDER, kids) if kids else None


def make_run(text, fmt):
    run = OxmlElement("w:r")
    rpr = run_props(fmt)
    if rpr is not None:
        run.append(rpr)
    for i, chunk in enumerate(INVALID_XML.sub("", text).split("\t")):
        if i:
            run.append(OxmlElement("w:tab"))
        if chunk:
            t = OxmlElement("w:t")
            t.text = chunk
            if chunk[0] == " " or chunk[-1] == " " or "  " in chunk:
                t.set(qn("xml:space"), "preserve")
            run.append(t)
    return run


@functools.lru_cache(maxsize=None)
def glyph_fonts():
    """HN и Arial кеглем 40 — проверять глифы отрисовкой; None — нет Pillow или шрифтов (не Mac)."""
    if ImageFont is None:
        return None
    try:
        return [ImageFont.truetype(path, 40, index=index)
                for path, index in (FONT_FILES["sans"][0][0], FONT_FILES["arial"][0][0])]
    except OSError:
        return None


@functools.lru_cache(maxsize=None)
def arial_only(ch):
    """Глифа нет в Helvetica Neue, но есть в Arial (→, ←). Нет глифа — шрифт рисует .notdef,
    как и заведомо отсутствующий U+E001."""
    fonts = glyph_fonts()
    if fonts is None:
        return False

    def drawn(font, c):
        mask = font.getmask(c)
        return mask.size, bytes(mask)

    hn, arial = fonts
    return drawn(hn, ch) == drawn(hn, "\ue001") and drawn(arial, ch) != drawn(arial, "\ue001")


# ---------------------------------------------------------------- разбор Markdown

def no_footnote(state, start, end, silent):
    """Определение сноски `[^метка]: текст` — абзац текста, а не определение ссылки: иначе текст
    пропадает, а [^метка] становится битой гиперссылкой (сноски не поддерживаются)."""
    if state.src.startswith("[^", state.bMarks[start] + state.tShift[start]):
        return False
    return reference(state, start, end, silent)


MD.block.ruler.at("reference", no_footnote)


def vt_ff_spaces(state):
    """&#11; и &#12; разделяют слова: пробел, а не склейка (&#11; markdown-it заменил бы на U+FFFD).
    До text_join, пока сущность — отдельный узел; и в alt картинки."""
    todo = [t.children for t in state.tokens if t.type == "inline"]
    while todo:
        for n in todo.pop() or []:
            if n.type == "text_special" and n.info == "entity" and VT_FF.fullmatch(n.markup):
                n.content = " "
            todo.append(n.children)


MD.core.ruler.before("text_join", "vt_ff_spaces", vt_ff_spaces)


def strip_front_matter(text):
    """YAML front matter → пустые строки (номера строк в предупреждениях сохраняются)."""
    lines = text.split("\n")
    if lines[0].rstrip() != "---" or len(lines) < 2 or not lines[1].strip():
        return text  # за --- пустая строка — это линия, а не метаданные (как в pandoc)
    for i in range(1, min(len(lines), 50)):
        if lines[i].rstrip() in ("---", "..."):  # «...» закрывает метаданные в pandoc и YAML
            inner = [line for line in lines[1:i] if line.strip()]
            if (inner and all(FM_LINE.match(line) for line in inner)
                    and any(KEY_LINE.match(line) for line in inner)):
                return "\n" * (i + 1) + "\n".join(lines[i + 1:])
            return text
    return text


def walk(node):
    yield node
    for child in getattr(node, "children", None) or []:
        yield from walk(child)


def inline_kids(node):
    """Inline-узлы абзаца, заголовка или ячейки."""
    return node.children[0].children if node.children else []


def flatten(nodes, fmt=frozenset()):
    """Inline-узлы → плоский список кусков (вид, значение, формат)."""
    hide = 0  # узлы до этого индекса — содержимое inline <script>/<style>/<template>
    for i, n in enumerate(nodes):
        t = n.type
        if i < hide:
            continue
        if t in ("text", "text_special"):  # text_special — сущность и \-экранирование в alt картинки
            yield ("text", n.content, fmt)
        elif t == "softbreak":
            yield ("text", " ", fmt)
        elif t == "hardbreak" or (t == "html_inline" and BR_RE.fullmatch(n.content.strip())):
            yield ("br", None, fmt)
        elif t == "code_inline":
            yield ("text", n.content, fmt | {"code"})
        elif t in FLAGS:
            yield from flatten(n.children, fmt | {FLAGS[t]})
        elif t == "link":
            href = str(n.attrs.get("href") or "")
            scheme = re.match(r"([A-Za-z][A-Za-z0-9+.-]+):", href)
            if scheme and scheme.group(1).lower() in SAFE_SCHEMES:
                yield ("link", href, fmt)
                yield from flatten(n.children, fmt | {"link"})
                yield ("endlink", None, fmt)
            else:  # относительный путь и якорь — просто текст, прочие схемы — текст с предупреждением
                if scheme:
                    yield ("badlink", href, fmt)
                yield from flatten(n.children, fmt)
        elif t == "image":
            yield ("image", n, fmt)
        elif t == "html_inline" and (m := re.match(r"<(script|style|template)\b", n.content, re.I)):
            # содержимое — не текст документа, если тег закрыт в том же абзаце (иначе — как прочие теги)
            hide = next((k + 1 for k in range(i + 1, len(nodes)) if nodes[k].type == "html_inline"
                         and re.fullmatch(rf"</{m.group(1)}\s*>", nodes[k].content, re.I)), 0)
        elif t != "html_inline" and n.children:
            yield from flatten(n.children, fmt)
        # прочие inline-теги и комментарии отбрасываются, текст между ними остаётся


def coalesce(pieces, bold=False):
    """Склеить соседние куски текста одного формата (меньше пробегов в XML)."""
    out = []
    for kind, val, fmt in pieces:
        if bold:
            fmt = fmt | BOLD
        if kind == "text" and not val:  # пустые куски остаются от разбора выделения
            continue
        if kind == "text" and out and out[-1][0] == "text" and out[-1][1].endswith(" ") and val.startswith(" ") \
                and "code" not in fmt | out[-1][2]:  # пробел кода — его содержимое
            val = val[1:]  # на месте вырезанного тега, и на границе форматов тоже
            if not val:
                continue
        if kind == "text" and out and out[-1][0] == "text" and out[-1][2] == fmt:
            out[-1] = ("text", out[-1][1] + val, fmt)
        else:
            out.append((kind, val, fmt))
    return out


def plain(pieces):
    return "".join(v if k == "text" else "\n" if k == "br" else "" for k, v, _ in pieces)


def one_line(text):
    return " ".join(text.split())


def marker_kind(html):
    """landscape / portrait / pagebreak для HTML-блока-маркера, иначе None."""
    m = re.fullmatch(r"<!--\s*(landscape|portrait|pagebreak)\s*-->", html.strip(), re.I)
    return m.group(1).lower() if m else None


def html_text(html):
    """Текст HTML-блока: комментарии, теги и <script>/<style> вырезаны, блочный тег разделяет
    слова переводом строки, <br> сохранён как перенос."""
    no_comments = HIDDEN_RE.sub("", COMMENT_RE.sub("", html))
    text = re.sub(r"\s*\n\s*", "\n", BLOCK_TAG_RE.sub("\n", no_comments))
    return TAG_RE.sub(lambda m: "<br>" if BR_RE.fullmatch(m.group(0)) else "", text)


def visible(node):
    """Выводит ли блок видимое — текст, картинку, таблицу или линию (иначе раздел с ним пуст, §4.4)."""
    t = node.type
    if t in ("table", "image", "hr"):
        return True
    if t in ("text", "code_inline", "fence", "code_block"):
        return bool(node.content.strip())
    if t == "html_block":  # как его выводит html_block()
        return not marker_kind(node.content) and \
            visible(SyntaxTreeNode(MD.parseInline(html_text(node.content))))
    return any(visible(child) for child in getattr(node, "children", None) or [])


def split_pagebreaks(node):
    """Строки \\newpage / \\pagebreak внутри абзаца → части абзаца; None, если таких строк нет.
    Строка сверяется с исходником абзаца: экранированное \\\\newpage — текст, а не разрыв."""
    kids = inline_kids(node)
    raw = [s.strip().lower() for s in node.children[0].content.split("\n")] if node.children else []
    last, cuts, line = len(kids) - 1, [], 0  # line — строка абзаца, с которой начинается kids[i]
    for i, k in enumerate(kids):
        if k.type == "text" and k.content.strip().lower() in PAGEBREAK_LINES \
                and (i == 0 or kids[i - 1].type in BREAKS) and (i == last or kids[i + 1].type in BREAKS):
            # счётчик может отстать (перевод строки внутри `кода` разбор не сохраняет): строка кандидата —
            # первая с этим текстом в любой записи (\\\\newpage, &#92;newpage), не выше счётчика
            line = next((r for r in range(line, len(raw)) if unescapeAll(raw[r]) in PAGEBREAK_LINES), line)
            if line < len(raw) and raw[line] in PAGEBREAK_LINES:
                cuts.append(i)
        line += sum(1 if n.type in BREAKS else n.content.count("\n")  # и внутри выделения, ссылки, тега
                    for n in walk(k) if n.type in BREAKS or n.type == "html_inline")
    if not cuts:
        return None
    parts, start = [], 0
    for i in cuts:
        parts.append(kids[start:max(i - 1, start)])
        start = i + 2
    parts.append(kids[start:])
    return parts


def pseudo_para(kids, map_):  # часть абзаца после разрывов: своих строк-разрывов в ней уже нет
    return SimpleNamespace(type="paragraph", map=map_,
                           children=[SimpleNamespace(type="inline", content="", children=kids)])


def break_line(line):
    """Строка таблицы — \\newpage / \\pagebreak (в списке и цитате — после отступа и «>»)?
    → отступ и «>» контейнера или None."""
    pre = re.match(r"[ \t>]*", line).group()
    return pre if line[len(pre):].strip().lower() in PAGEBREAK_LINES else None


def split_table_breaks(text):
    """Строка \\newpage / \\pagebreak под строкой таблицы по GFM стала бы строкой таблицы.
    Если под ней (и под серией таких строк) строка той же таблицы с «|», строка-разрыв остаётся в
    таблице — по ней таблица режется на части (Converter.break_rows). Иначе — таблица кончается,
    начинается новая таблица со своей шапкой или идёт строка без «|» (текст после таблицы) —
    строка становится маркером <!-- pagebreak -->: он таблицу закрывает, номера строк не
    сдвигаются. → (текст, номера оставленных строк-разрывов)."""
    lines, found = text.split("\n"), {}  # строка-разрыв → конец её таблицы
    for tok in MD.parse(text):
        if tok.type == "table_open" and tok.map:
            found.update((i, tok.map[1]) for i in range(tok.map[0] + 2, tok.map[1])
                         if break_line(lines[i]) is not None)
    if not found:
        return text, frozenset()
    marked = [break_line(line) + "<!-- pagebreak -->" if i in found else line
              for i, line in enumerate(lines)]
    # где начинаются таблицы, если каждая строка-разрыв таблицу закрывает
    starts = {tok.map[0] for tok in MD.parse("\n".join(marked)) if tok.type == "table_open" and tok.map}
    kept = set()
    for i, end in found.items():
        j = next(j for j in itertools.count(i + 1) if j not in found)  # строка под серией разрывов
        if j < end and j not in starts and "|" in lines[j]:
            kept.add(i)
            marked[i] = lines[i]
    return "\n".join(marked), frozenset(kept)


def tables_in(node, indent=0):
    """Таблицы блока с отступом их контейнера (цитата и список дают по STEP)."""
    if node.type == "table":
        yield node, indent
        return
    step = STEP if node.type in ("blockquote", "bullet_list", "ordered_list") else 0
    for child in getattr(node, "children", None) or []:
        if child.type != "inline":
            yield from tables_in(child, min(indent + step, MAX_INDENT))


# ---------------------------------------------------------------- замер ширины текста

_SUPP, _WINF = "/System/Library/Fonts/Supplemental/", "C:\\Windows\\Fonts\\"
_ARIAL = [
    [(_SUPP + n, 0) for n in ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf")],
    [(_WINF + n, 0) for n in ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")],
]
FONT_FILES = {  # варианты: обычный, полужирный, курсив, полужирный курсив
    "sans": [[("/System/Library/Fonts/HelveticaNeue.ttc", i) for i in range(4)]] + _ARIAL,
    "arial": _ARIAL,  # замена HN в Word на Windows: у части глифов шире HN (курсивная «т», скобки)
    "mono": [[("/System/Library/Fonts/Menlo.ttc", i) for i in range(4)]] + _ARIAL,
}


class Measurer:
    """Ширина текста в twips: Pillow по реальному шрифту, иначе эвристика 0,55/0,60 em (§3.4)."""

    def __init__(self):
        self.fonts, self.cache = {}, {}

    def font(self, family, variant):
        key = (family, variant)
        if key not in self.fonts:
            self.fonts[key] = None
            for files in FONT_FILES[family] if ImageFont is not None else []:
                path, index = files[variant]
                try:
                    self.fonts[key] = ImageFont.truetype(path, 1000, index=index)
                    break
                except OSError:
                    continue
        return self.fonts[key]

    def width(self, text, fmt):
        key = (text, fmt)
        if key not in self.cache:
            mono, bold, italic = "code" in fmt, "b" in fmt, "i" in fmt
            fonts = [self.font(fam, bold + 2 * italic) for fam in (("mono",) if mono else ("sans", "arial"))]
            fonts = [f for f in fonts if f is not None]
            if fonts:  # бо́льшая из ширин по HN и по Arial: вёрстка не зависит от подмены (§3.4)
                em = max(f.getlength(text) for f in fonts) / 1000
            else:
                em = len(text) * (0.60 if bold or mono else 0.55)
            self.cache[key] = em * (10 if mono else 11) * 20 * 1.03  # +3 % запаса
        return self.cache[key]

    def extent(self, pieces):
        """(min, max, мягкий min) ширины ячейки с полями: самое длинное слово, самая длинная строка
        и самое длинное слово без переноса по дефисам (даты и коды целиком)."""
        lines, boxes, lo, hi, soft = [[]], [0.0], 0.0, 0.0, 0.0
        for kind, val, fmt in pieces:
            if kind == "br":
                lines.append([])
                boxes.append(0.0)
            elif kind == "text":
                lines[-1].append((val, fmt))
            elif kind == "box":  # картинка: строка шире на её ширину, уже 3 см её не сжимать
                boxes[-1] += val[0]
                lo = max(lo, min(val[0], 1701))
        for line, box in zip(lines, boxes):
            hi = max(hi, sum(self.width(t, f) for t, f in line) + box)
            for word in split_words(line):
                if sum(len(t) for t, _ in word) <= 30:  # URL и хэши Word перенесёт посимвольно
                    lo = max(lo, sum(self.width(t, f) for t, f in word))
            for word in split_words(line, dashes=False):
                if sum(len(t) for t, _ in word) <= 30:
                    soft = max(soft, sum(self.width(t, f) for t, f in word))
        return lo + CELL_PAD, hi + CELL_PAD, max(lo, soft) + CELL_PAD

    def height(self, pieces, avail):
        """Оценка высоты ячейки в twips: жадный перенос слов, слово шире ячейки — посимвольно;
        картинка добавляет свою высоту (шире ячейки — уменьшенную, как в image())."""
        lines, boxes = [[]], 0.0
        for kind, val, fmt in pieces:
            if kind == "br":
                lines.append([])
            elif kind == "text":
                lines[-1].append((val, fmt))
            elif kind == "box":
                boxes += val[1] * min(1.0, avail / max(val[0], 1))
        count, space = 0, self.width(" ", frozenset())
        for line in lines:
            count, x = count + 1, 0.0
            for word in split_words(line):
                w = sum(self.width(t, f) for t, f in word) + space
                if x + w > avail:
                    count += (1 if x else 0) + int(w // avail)
                    x = w % avail
                else:
                    x += w
        return count * LINE_H + 80 + boxes  # + поля ячейки 40 + 40


def split_words(line, dashes=True):
    """Неразрывные слова строки: границы — пробелы, а также (dashes) место после дефиса и тире."""
    word = []
    for text, fmt in line:
        for i, chunk in enumerate(WORD_GAP.split(text)):
            if i and word:
                yield word
                word = []
            for j, part in enumerate(AFTER_DASH.split(chunk) if dashes and not URL_TOKEN.search(chunk)
                                     else [chunk]):
                if j and word:
                    yield word
                    word = []
                if part:
                    word.append((part, fmt))
    if word:
        yield word


def column_widths(mins, softs, maxs, areas, total):
    """Ширины колонок (§3.4) и признак «влезло без разрыва слов» — он по жёстким min. Нижняя
    граница колонки — мягкий min (даты и коды целиком), если его хватает на все колонки."""
    lows = softs if sum(softs) <= total else mins
    smin, smax = sum(lows), sum(maxs)
    if smax <= total:
        raw = [m + (total - smax) * m / smax for m in maxs]
    elif smin <= total:
        # по объёму текста тела в границах [low_j, max_j] (water-filling); какие колонки встают на
        # границу, решает знак суммарного выхода за границы. Колонки без текста в теле — на нижней
        # границе, а если остальные упёрлись в max, остаток — им по (max − low), как прежде
        n = len(lows)
        fixed = {j: lows[j] for j in range(n) if not areas[j]}
        while len(fixed) < n:
            free = [j for j in range(n) if j not in fixed]
            rest, mass = total - sum(fixed.values()), sum(areas[j] for j in free)
            share = {j: rest * areas[j] / mass for j in free}
            over = sum(min(max(s, lows[j]), maxs[j]) - s for j, s in share.items())
            bound = {j: lows[j] if over > 0 else maxs[j] for j, s in share.items()
                     if (s < lows[j] if over > 0 else s > maxs[j])}
            fixed.update(bound or share)
        raw = [fixed[j] for j in range(n)]
        idle = [j for j in range(n) if not areas[j]]
        slack, span = total - sum(raw), sum(maxs[j] - lows[j] for j in idle)
        for j in idle:
            raw[j] += slack * (maxs[j] - lows[j]) / span if span > 0 else 0
    else:  # не влезает: узкие колонки (номера, даты) — целиком, дефицит — с самых широких
        rest, left = total, len(mins)
        for m in sorted(mins):
            if m * left > rest:
                break
            rest, left = rest - m, left - 1
        raw = [min(a, rest / left) for a in mins]
    cols = [int(w) for w in raw]
    cols[raw.index(max(raw))] += total - sum(cols)  # остаток округления — самой широкой
    return cols, sum(mins) <= total


# ---------------------------------------------------------------- документ и стили

def restyle(style, ppr=None, rpr=None):
    """Полностью заменить свойства абзаца и текста у стиля."""
    for tag in ("w:pPr", "w:rPr"):
        old = style.find(qn(tag))
        if old is not None:
            style.remove(old)
    if ppr:
        insert_ordered(style, ordered("w:pPr", PPR_ORDER, ppr), STYLE_ORDER)
    if rpr:
        insert_ordered(style, ordered("w:rPr", RPR_ORDER, rpr), STYLE_ORDER)


def find_style(styles, style_id):
    found = styles.xpath(f'w:style[@w:styleId="{style_id}"]')
    return found[0] if found else None


def setup_styles(doc):
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rpr.clear()
    for child in (E("w:rFonts", {"w:ascii": HN, "w:eastAsia": "Times New Roman", "w:hAnsi": HN, "w:cs": HN}),
                  E("w:sz", {"w:val": 22}), E("w:szCs", {"w:val": 22}),
                  E("w:lang", {"w:val": "ru-RU", "w:eastAsia": "ru-RU", "w:bidi": "ar-SA"})):
        rpr.append(child)
    ppr = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr"))
    ppr.clear()
    ppr.append(E("w:widowControl"))
    ppr.append(spacing(after=160, line=278))

    def heading_rpr(level, size):
        kids = [rfonts(HN), E("w:b"), E("w:bCs"), E("w:sz", {"w:val": size}), E("w:szCs", {"w:val": size})]
        return kids + ([E("w:i"), E("w:iCs")] if level == 6 else [])

    for level, (size, before) in HEADINGS.items():
        restyle(find_style(styles, f"Heading{level}"),
                [E("w:keepNext"), E("w:keepLines"), spacing(before, 160),
                 E("w:outlineLvl", {"w:val": level - 1})],
                heading_rpr(level, size))
        linked = find_style(styles, f"Heading{level}Char")  # связанный символьный — тем же шрифтом
        if linked is not None:
            restyle(linked, None, heading_rpr(level, size))

    def quote_rpr():
        return [E("w:i"), E("w:iCs"), E("w:color", {"w:val": "555555"})]

    restyle(find_style(styles, "Quote"),
            [E("w:pBdr", None, E("w:left", {"w:val": "single", "w:sz": 18, "w:space": 8, "w:color": "DDDDDD"})),
             spacing(160, 160, 276), ind(STEP)],
            quote_rpr())
    linked = find_style(styles, "QuoteChar")
    if linked is not None:
        restyle(linked, None, quote_rpr())

    if find_style(styles, "Hyperlink") is None:
        styles.append(E("w:style", {"w:type": "character", "w:styleId": "Hyperlink"},
                        E("w:name", {"w:val": "Hyperlink"}),
                        E("w:basedOn", {"w:val": "DefaultParagraphFont"}),
                        E("w:uiPriority", {"w:val": 99}), E("w:unhideWhenUsed"),
                        E("w:rPr", None, E("w:color", {"w:val": "0563C1"}), E("w:u", {"w:val": "single"}))))


def new_document(title):
    # fontTable у python-docx — двоичная часть без API; регистрируем её как XML-часть
    PartFactory.part_type_for.setdefault(CT.WML_FONT_TABLE, XmlPart)
    doc = docx.Document()
    setup_styles(doc)

    compat = doc.settings.element.find(qn("w:compat"))
    for setting in compat.findall(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")  # 14 — таблица в Word уезжает в левое поле

    fonts = doc.part.part_related_by(RT.FONT_TABLE).element
    for name, props in FONT_ENTRIES:
        for old in fonts.findall(qn("w:font")):
            if old.get(qn("w:name")) == name:
                fonts.remove(old)
        fonts.append(E("w:font", {"w:name": name}, *(E("w:" + tag, attrs) for tag, attrs in props)))

    # Из шаблона python-docx убираем устаревшую копию стилей для Word 2010 (stylesWithEffects
    # со старыми Cambria/синими заголовками) и пустую миниатюру (белый лист в превью проводника)
    for rels, kind in ((doc.part.rels, "/stylesWithEffects"), (doc.part.package.rels, "/thumbnail")):
        for rid in [rel.rId for rel in rels.values() if rel.reltype.endswith(kind)]:
            del rels[rid]

    core = doc.core_properties
    core.title = title[:255]  # python-docx: свойство длиннее 255 символов — ValueError
    core.identifier = MARK  # Word метку не показывает; по ней узнаём свой файл при перезаписи
    core.author = core.last_modified_by = core.comments = ""
    core.subject = core.keywords = core.category = ""
    core.revision = 1
    core.created = core.modified = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    return doc


# ---------------------------------------------------------------- конвертер

class Converter:
    """Состояние одной конвертации: план разделов, текущая ширина, нумерация, разделители."""

    def __init__(self, md_path: Path, landscape: bool, auto_landscape: bool, breaks=frozenset()):
        self.md_path = md_path
        self.default = "landscape" if landscape else "portrait"
        self.auto = auto_landscape
        self.breaks = breaks  # номера строк md: \newpage между строками таблицы (split_table_breaks)
        self.measure = Measurer()
        self.tables = {}      # id(узла таблицы) → (ряды, выравнивания, min, max, мягкие min, объёмы)
        self.abstract = {}    # вид списка → abstractNumId
        self.warnings = []
        self.line = None
        self.width = TEXT_W[self.default]
        self.after_table = False   # последним в тело добавлена таблица
        self.page_break = False    # отложенный разрыв страницы перед следующим блоком
        self.pictures = 0          # счётчик id рисунков (wp:docPr)
        self.code_p = None         # абзац последнего блока кода

    def warn(self, msg):
        name = self.md_path.name
        self.warnings.append(tr("{file}, line {line}: {message}", file=name, line=self.line, message=msg)
                             if self.line else f"{name}: {msg}")

    # --- план разделов (§4)

    def top_items(self, root):
        """Блоки верхнего уровня и маркеры: [('block', узел) | ('marker', вид)]."""
        items = []
        for node in root.children:
            self.line = node.map[0] + 1 if node.map else None
            if node.type == "html_block":
                kind = marker_kind(node.content)
                if kind:
                    items.append(("marker", kind))
                    continue
                if not html_text(node.content).strip():  # комментарий или теги без текста
                    if COMMENT_RE.sub("", node.content).strip():
                        self.warn(tr("HTML block without text skipped"))
                    elif not node.content.rstrip().endswith("-->"):
                        self.warn(tr("unclosed HTML comment — the rest of the document is skipped"))
                    continue
            elif node.type == "paragraph":
                parts = split_pagebreaks(node)
                if parts is not None:
                    for i, part in enumerate(parts):
                        if i:
                            items.append(("marker", "pagebreak"))
                        if part:
                            items.append(("block", pseudo_para(part, node.map)))
                    continue
            elif node.type == "table" and any(row == 1 for row, _ in self.break_rows(node)):
                items.append(("marker", "pagebreak"))  # разрыв над первой строкой тела — перед таблицей
            items.append(("block", node))
        return items

    def wide_tables(self, block):
        """Таблицы блока, которые на книжной странице не укладываются без разрыва слов."""
        return [t for t, indent in tables_in(block)
                if sum(self.table_info(t)[2]) > TEXT_W["portrait"] - indent]

    def auto_groups(self, blocks):
        """Автоальбомные группы книжного участка: [начало, конец] по индексам блоков."""
        groups = []
        for i, block in enumerate(blocks):
            if block is PAGEBREAK or not self.wide_tables(block):
                continue
            start = self.group_start(blocks, i)
            if groups:
                gap = blocks[groups[-1][1] + 1:start]
                if sum(map(visible, gap)) <= 2 and not any(
                        g is PAGEBREAK or any(n.type == "table" for n in walk(g)) for g in gap):
                    groups[-1][1] = i  # склейка соседних групп: между ними не больше 2 видимых блоков
                    continue
            groups.append([start, i])
        return groups

    def group_start(self, blocks, i):
        """Начало группы с широкой таблицей blocks[i]: что уходит на альбомную вместе с ней (§4.3).
        Невидимые блоки прозрачны. Подпись с двоеточием — всегда; выше до 2 коротких простых абзацев
        и 1 линии — только вместе с заголовком за ними; так же выше, пока заголовки строго старше."""
        def back(j):  # предыдущий видимый блок, разрыв или -1 (начало участка)
            j -= 1
            while j >= 0 and blocks[j] is not PAGEBREAK and not visible(blocks[j]):
                j -= 1
            return j

        def typed(j, *types):
            return j >= 0 and blocks[j] is not PAGEBREAK and blocks[j].type in types

        def lead(j):  # подводка: абзац до 600 знаков без картинок или линия
            pieces = list(flatten(inline_kids(blocks[j]))) if typed(j, "paragraph") else []
            return typed(j, "paragraph", "hr") and len(one_line(plain(pieces))) <= 600 \
                and all(kind != "image" for kind, _, _ in pieces)

        start, j = i, back(i)
        if typed(j, "paragraph") and plain(flatten(inline_kids(blocks[j]))).rstrip().endswith(":"):
            start, j = j, back(j)  # подпись с двоеточием уходит вместе с таблицей
        top = 7  # уровень самого старшего из забранных заголовков
        while True:
            k, paras, rules = j, 0, 0
            while lead(k) and paras <= 2 and rules <= 1:
                paras += blocks[k].type == "paragraph"
                rules += blocks[k].type == "hr"
                k = back(k)
            if paras > 2 or rules > 1 or not typed(k, "heading") or int(blocks[k].tag[1]) >= top:
                return start
            while typed(k, "heading"):  # заголовок и заголовки прямо над ним — вместе с подводкой
                start, top, k = k, min(top, int(blocks[k].tag[1])), back(k)
            j = k

    def plan(self, items):
        """→ ([(ориентация, [блок | PAGEBREAK])], число автоальбомных таблиц)."""
        segments = [[self.default, []]]  # явно размеченные участки
        pending, brk, started = None, False, False
        for kind, val in items:
            if kind == "marker":  # серия подряд: ориентация — последняя, разрыв — если был хоть один
                if val == "pagebreak":
                    brk = True
                else:
                    pending = val
                continue
            if not visible(val):  # пустой блок не открывает участок: маркеры вокруг — «подряд»
                segments[-1][1].append(val)
                continue
            if started:
                if pending and pending != segments[-1][0]:
                    segments.append([pending, []])  # новый раздел и так с новой страницы
                elif brk:
                    segments[-1][1].append(PAGEBREAK)
            elif pending:
                segments[-1][0] = pending  # маркер до содержимого — ориентация первого раздела
            pending, brk, started = None, False, True
            segments[-1][1].append(val)

        flat, auto_tables = [], 0
        for orient, blocks in segments:
            marks = [orient] * len(blocks)
            if self.auto and orient == "portrait":
                for start, end in self.auto_groups(blocks):
                    marks[start:end + 1] = ["landscape"] * (end + 1 - start)
                    auto_tables += sum(len(self.wide_tables(b)) for b in blocks[start:end + 1]
                                       if b is not PAGEBREAK)
            if flat:  # маркер участка — «с новой страницы», даже если автоальбомная уравняла ориентацию
                flat.append([orient, PAGEBREAK])
            flat += [[m, b] for m, b in zip(marks, blocks)]
        # пустой блок не открывает раздел, а идёт в раздел видимого соседа (§4.4)
        shown = next((m for m, b in flat if visible(b)), None)
        for item in flat:
            if visible(item[1]):
                shown = item[0]
            elif shown:
                item[0] = shown
        # разрыв страницы — в раздел следующего блока: на стыке разделов он лишний
        for k in range(len(flat) - 2, -1, -1):
            if flat[k][1] is PAGEBREAK:
                flat[k][0] = flat[k + 1][0]
        sections = []
        for orient, block in flat:
            if not sections or sections[-1][0] != orient:
                sections.append((orient, []))
            blocks = sections[-1][1]
            if block is PAGEBREAK and (not blocks or blocks[-1] is PAGEBREAK):
                continue
            blocks.append(block)
        return sections or [(self.default, [])], auto_tables

    # --- рендер

    def run(self, root):
        sections, auto_tables = self.plan(self.top_items(root))
        title = next((one_line(plain(flatten(inline_kids(n)))) for n in walk(root)
                      if n.type == "heading" and n.tag == "h1"), "") or self.md_path.stem
        self.doc = new_document(title)
        self.body = self.doc.element.body
        self.numbering = self.doc.part.numbering_part.element
        for k, (orient, blocks) in enumerate(sections):
            if k:
                self.section_break()
            self.setup_section(orient)
            prev = None  # предыдущий видимый блок раздела после последнего разрыва
            for block in blocks:
                if block is PAGEBREAK:
                    self.page_break, prev = True, None
                    continue
                if block.type == "table" and prev is not None and prev.type == "heading":
                    # отбивка 6 pt между заголовком и таблицей — как в эталоне (SPEC §3.4)
                    self.add(make_p([E("w:keepNext"), spacing(0, 0, 120, "exact"), tiny_rpr()]))
                self.block(block, Ctx())
                if visible(block):
                    prev = block
        if not any(visible(b) for _, bs in sections for b in bs):  # и абзац из одного <br>
            self.warn(tr("the document has no visible content — the result is a single empty page"))
        if self.after_table or len(self.body) == 1:
            # таблица в конце — носитель 1 pt, как у разрыва раздела (разделитель 8 pt под таблицей,
            # дошедшей до низа страницы, дал бы пустую последнюю страницу); пустой md — пустой абзац
            last = [spacing(0, 0, 20, "exact"), tiny_rpr()] if self.after_table else []
            self.body.sectPr.addprevious(make_p(last))
        return sections, auto_tables

    def setup_section(self, orient):
        """Каждый раздел явно получает A4, ориентацию и поля (шаблон python-docx — Letter)."""
        section = self.doc.sections[-1]
        width, height = PAGE[orient]
        section.orientation = WD_ORIENT.LANDSCAPE if orient == "landscape" else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Twips(width), Twips(height)
        section.top_margin = section.bottom_margin = Twips(MARGIN)
        section.left_margin = section.right_margin = Twips(MARGIN)
        section.header_distance = section.footer_distance = Twips(HDR_FTR)
        section.gutter = Twips(0)
        self.body.sectPr.get_or_add_type().set(qn("w:val"), "nextPage")
        self.width = TEXT_W[orient]

    def section_break(self):
        self.after_table = False  # носитель разрыва сам служит разделителем после таблицы
        self.doc.add_section(WD_SECTION.NEW_PAGE)
        ppr = self.body.sectPr.getprevious().pPr  # абзац-носитель sectPr прошлого раздела
        insert_ordered(ppr, spacing(0, 0, 20, "exact"), PPR_ORDER)
        insert_ordered(ppr, tiny_rpr(), PPR_ORDER)

    def add(self, el):
        """Добавить блок в тело: разделитель после таблицы и отложенный разрыв страницы."""
        if self.after_table and not self.page_break:  # абзац с разрывом сам отделяет таблицу:
            self.body.sectPr.addprevious(spacer())    # разделитель под таблицей у низа страницы дал бы пустую
        if self.page_break:
            self.page_break = False
            if el.tag == qn("w:tbl"):  # у таблицы нет pageBreakBefore — нужен абзац перед ней
                self.body.sectPr.addprevious(
                    make_p([E("w:pageBreakBefore"), spacing(0, 0, 20, "exact"), tiny_rpr()]))
            else:
                insert_ordered(el.get_or_add_pPr(), E("w:pageBreakBefore"), PPR_ORDER)
        self.body.sectPr.addprevious(el)
        self.after_table = el.tag == qn("w:tbl")

    def block(self, node, ctx):
        if node.map:
            self.line = node.map[0] + 1
        t = node.type
        if t == "heading":
            ppr = [pstyle(f"Heading{node.tag[1]}")] + ([ind(ctx.indent)] if ctx.indent else [])
            self.add_par(ppr, flatten(inline_kids(node), frozenset({"head"})), ctx)
        elif t == "paragraph":
            self.paragraph(node, ctx)
        elif t in ("bullet_list", "ordered_list"):
            self.list_block(node, ctx)
        elif t == "blockquote":
            inner = Ctx(min(ctx.indent + STEP, MAX_INDENT), ctx.quote + 1, ctx.depth, False)
            for child in node.children:
                self.block(child, inner)
        elif t in ("fence", "code_block"):
            self.code(node, ctx)
        elif t == "table":
            self.table(node, ctx)
        elif t == "hr":
            self.add(make_p([E("w:pBdr", None, E("w:bottom", {"w:val": "single", "w:sz": 6, "w:space": 1,
                                                              "w:color": "BBBBBB"}))]
                            + ([ind(ctx.indent)] if ctx.indent else [])))
        elif t == "html_block":
            self.html_block(node, ctx)

    def add_par(self, ppr, pieces, ctx):
        p = make_p(ppr)
        self.inline(p, pieces, max_w=self.width - ctx.indent)
        self.add(p)

    def paragraph(self, node, ctx, num=None):
        kids = inline_kids(node)
        parts = split_pagebreaks(node)
        if parts is None:
            parts = [kids]
        else:  # сюда доходят только абзацы внутри списков и цитат
            self.warn(tr("page break marker inside a list or quote is ignored"))
            parts = [p for p in parts if p] or ([[]] if num else [])
        for part in parts:
            self.text_par(part, ctx, num)
            num = None

    def text_par(self, kids, ctx, num=None):
        ppr = [pstyle("Quote")] if ctx.quote else []
        if ctx.item:
            ppr.append(spacing(before=0 if ctx.quote else None, after=80))
        if num:
            num_id, ilvl = num
            ppr.append(E("w:numPr", None, E("w:ilvl", {"w:val": ilvl}), E("w:numId", {"w:val": num_id})))
            if ctx.indent != STEP * (ilvl + 1):  # уровень нумерации даёт иной отступ (цитата и т. п.)
                ppr.append(ind(ctx.indent, STEP))
        elif ctx.indent:
            ppr.append(ind(ctx.indent))
        self.add_par(ppr, flatten(kids), ctx)

    def list_block(self, node, ctx):
        kind = node.markup if node.type == "ordered_list" else "bullet"
        ilvl = min(ctx.depth, 8)
        num_id = self.new_num(kind, ilvl, int(node.attrs.get("start", 1)))
        inner = Ctx(min(ctx.indent + STEP, MAX_INDENT), ctx.quote, ctx.depth + 1, True)
        for item in node.children:
            self.line = item.map[0] + 1 if item.map else self.line
            kids = list(item.children)
            if kids and kids[0].type == "paragraph":
                self.paragraph(kids.pop(0), inner, (num_id, ilvl))
            else:  # пункт начинается не с абзаца — номер на пустом абзаце
                self.text_par([], inner, (num_id, ilvl))
            for child in kids:
                self.block(child, inner)
        if ctx.depth == 0:  # у последнего пункта внешнего списка интервал после — 160
            last = self.body.sectPr.getprevious()
            sp = last.find(f"{qn('w:pPr')}/{qn('w:spacing')}") if last is not None else None
            if sp is not None and last.tag == qn("w:p") and sp.get(qn("w:after")) == "80":
                sp.set(qn("w:after"), "160")

    def new_num(self, kind, ilvl, start):
        """Свой w:num на каждый список: соседние списки не продолжают нумерацию друг друга."""
        if kind not in self.abstract:
            self.abstract[kind] = self.add_abstract(kind)
        num = self.numbering.add_num(self.abstract[kind])
        num.add_lvlOverride(ilvl).add_startOverride(start)
        return num.numId

    def add_abstract(self, kind):
        ids = [int(x) for x in self.numbering.xpath("w:abstractNum/@w:abstractNumId")]
        aid = max(ids, default=-1) + 1
        abstract = E("w:abstractNum", {"w:abstractNumId": aid},
                     E("w:nsid", {"w:val": f"{0x6D640000 + aid:08X}"}),
                     E("w:multiLevelType", {"w:val": "hybridMultilevel"}))
        for i in range(9):
            abstract.append(E(
                "w:lvl", {"w:ilvl": i},
                E("w:start", {"w:val": 1}),
                E("w:numFmt", {"w:val": "bullet" if kind == "bullet" else "decimal"}),
                E("w:lvlText", {"w:val": BULLETS[i % 3] if kind == "bullet" else f"%{i + 1}{kind}"}),
                E("w:lvlJc", {"w:val": "left"}),
                E("w:pPr", None, ind(STEP * (i + 1), STEP)),
                # ◦ и ▪ в HN нет: маркеры — Arial (есть и на Mac, и на Windows), иначе запасной шрифт
                *([E("w:rPr", None, rfonts("Arial"))] if kind == "bullet" else [])))
        first_num = self.numbering.find(qn("w:num"))  # abstractNum — строго перед первым num
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            self.numbering.append(abstract)
        return aid

    def code(self, node, ctx):
        text = node.content[:-1] if node.content.endswith("\n") else node.content
        # рамка цвета заливки (0,5 pt с отступом 4 pt) — внутренние поля плашки; отступы абзаца 90 —
        # на рамку с отступом: плашка встаёт ровно по краям колонки и в Word, и в LibreOffice
        pad = {**BORDER, "w:space": 4, "w:color": SHADE["w:fill"]}
        p = make_p([E("w:pBdr", None, *(E("w:" + s, pad) for s in ("top", "left", "bottom", "right"))),
                    E("w:shd", SHADE), spacing(0, 160, 240),
                    E("w:ind", {"w:left": ctx.indent + 90, "w:right": 90}), E("w:jc", {"w:val": "left"})])
        for i, line in enumerate(text.split("\n")):
            if i:
                p.append(E("w:r", None, run_props(MONO_FMT), E("w:br")))
            if line:
                p.append(make_run(line, MONO_FMT))
        if self.code_p is not None and self.body.sectPr.getprevious() is self.code_p and not self.page_break:
            self.body.sectPr.addprevious(spacer())  # соседние блоки: Word слил бы одинаковые рамки в одну
        self.add(p)
        self.code_p = p

    def html_block(self, node, ctx):
        kind = marker_kind(node.content)
        if kind:
            self.warn(tr("page break marker inside a list or quote is ignored") if kind == "pagebreak"
                      else tr("orientation marker inside a list or quote is ignored"))
            return
        text = html_text(node.content).strip()
        if not text:
            if COMMENT_RE.sub("", node.content).strip():
                self.warn(tr("HTML block without text skipped"))
            return
        self.warn(tr("HTML block: tags removed, the text is kept as a plain paragraph"))
        self.text_par(inline_kids(SyntaxTreeNode(MD.parseInline(text))), ctx)

    # --- таблицы (§3.4)

    def break_rows(self, node):
        """Строки \\newpage в таблице: [(номер строки таблицы, перед которой разрыв, — 0 у шапки;
        номер строки md)]. У строк-разрывов подряд номер строки таблицы один (серия)."""
        out, row = [], 0
        for part in node.children:  # thead, tbody
            for tr in part.children:
                if tr.map and tr.map[0] in self.breaks:
                    out.append((row, tr.map[0]))
                else:
                    row += 1
        return out

    def table_info(self, node):
        key = id(node)
        if key not in self.tables:
            rows, aligns = [], []
            for part in node.children:  # thead, tbody
                for tr in part.children:
                    if tr.map and tr.map[0] in self.breaks:  # строка-разрыв — не строка таблицы
                        continue
                    rows.append([list(flatten(inline_kids(cell))) for cell in tr.children])
                    if part.type == "thead":
                        for cell in tr.children:
                            m = re.search(r"text-align:\s*(left|center|right)", str(cell.attrs.get("style", "")))
                            aligns.append(m.group(1) if m else "left")
            n = len(aligns) or max(len(r) for r in rows)
            aligns += ["left"] * (n - len(aligns))
            rows = [(r + [[] for _ in range(n)])[:n] for r in rows]  # как в GFM: дополнить / отбросить
            mins, maxs, softs, areas = [0.0] * n, [0.0] * n, [0.0] * n, [0.0] * n
            for i, row in enumerate(rows):
                for j, pieces in enumerate(row):
                    lo, hi, soft = self.measure.extent([self.sized(p) for p in coalesce(pieces, bold=i == 0)])
                    mins[j], maxs[j], softs[j] = max(mins[j], lo), max(maxs[j], hi), max(softs[j], soft)
                    if i:  # объём текста тела: ширина ячейки в одну строку, шапка не считается
                        areas[j] += hi - CELL_PAD
            self.tables[key] = (rows, aligns, mins, maxs, softs, areas)
        return self.tables[key]

    def sized(self, piece):
        """Картинка ячейки для замера: файл — ("box", (ширина, высота) в twips), иначе текст заглушки."""
        kind, node, fmt = piece
        if kind != "image":
            return piece
        path = self.local_file(str(node.attrs.get("src") or ""))
        if path is not None and path.stat().st_size <= MAX_IMAGE_BYTES:
            try:
                info = Image.from_file(str(path))
                width, height = info.px_width * 15, info.px_height * 15  # 96 dpi: 1 px = 15 twips
                orient = "landscape" if self.width == TEXT_W["landscape"] else "portrait"
                limit = TEXT_H[orient] * 8 // 10  # как в image(): не выше 0,8 области текста
                return ("box", (width, height) if height <= limit else (width * limit // height, limit), fmt)
            except Exception:  # битый файл — в ячейке будет заглушка
                pass
        alt = one_line(plain(flatten(node.children)))
        return ("text", tr("[image: {alt}]").format(alt=alt), fmt | {"i"})  # как в image()

    def table(self, node, ctx):
        rows, aligns, mins, maxs, softs, areas = self.table_info(node)
        total = self.width - ctx.indent
        widths, fits = column_widths(mins, softs, maxs, areas, total)
        orient = "landscape" if self.width == TEXT_W["landscape"] else "portrait"
        if not fits:
            self.warn(tr("the table does not fit a landscape page without breaking words") if orient == "landscape"
                      else tr("the table does not fit a portrait page without breaking words"))
        # шапка без текста и картинок не выводится: пустой серой полосы нет
        headed = len(rows) == 1 or any(kind == "image" or kind == "text" and val.strip()
                                       for cell in rows[0] for kind, val, _ in cell)
        # \newpage между строками тела режет таблицу: части с новой страницы, с той же сеткой, шапка —
        # только у первой (§3.4). Разрыв над первой строкой тела стоит перед таблицей (top_items);
        # внутри списка или цитаты разрыв игнорируется, как маркер
        top = ctx == Ctx()
        cuts = sorted({row for row, _ in self.break_rows(node) if row > 1}) if top else []
        for lo, hi in zip([0] + cuts, cuts + [len(rows)]):
            if lo:
                self.page_break = True
            part = rows[lo:hi] if lo or headed else rows[1:hi]
            self.add(self.table_el(part, headed and not lo, aligns, widths, total, ctx.indent, orient))
        if not top:
            for _, line in self.break_rows(node):
                self.line = line + 1
                self.warn(tr("page break marker inside a list or quote is ignored"))

    def table_el(self, rows, headed, aligns, widths, total, indent, orient):
        """w:tbl из строк rows по готовой сетке widths; headed — первая строка шапка."""
        heights = [max(self.measure.height([self.sized(p) for p in coalesce(pieces, bold=headed and i == 0)],
                                           max(w - CELL_PAD, 1))
                       for pieces, w in zip(row, widths)) for i, row in enumerate(rows)]
        room = TEXT_H[orient] - (heights[0] if headed else 0)  # область страницы под шапкой
        sides = ("top", "left", "bottom", "right", "insideH", "insideV")
        margins = (("top", 40), ("left", 108), ("bottom", 40), ("right", 108))
        tbl = E("w:tbl", None,
                E("w:tblPr", None,
                  E("w:tblW", {"w:w": total, "w:type": "dxa"}),
                  E("w:tblInd", {"w:w": indent, "w:type": "dxa"}),
                  E("w:tblBorders", None, *(E("w:" + s, BORDER) for s in sides)),
                  E("w:tblLayout", {"w:type": "fixed"}),
                  E("w:tblCellMar", None, *(E("w:" + s, {"w:w": v, "w:type": "dxa"}) for s, v in margins)),
                  E("w:tblLook", TBL_LOOK if headed else {**TBL_LOOK, "w:val": "0480", "w:firstRow": 0})),
                E("w:tblGrid", None, *(E("w:gridCol", {"w:w": w}) for w in widths)))
        last, first = len(rows) - 1, 0  # first — строка, с которой начинается цепочка keepNext
        for i, row in enumerate(rows):
            head = headed and i == 0
            # строку выше страницы с cantSplit LibreOffice выводит за поле, а Word для Mac уносит
            # на новую страницу и всё равно делит (SPEC §3.4, «Строки»);
            # запас 10 %: оценка не знает, где рендер перенесёт URL (по «/» и «.»)
            whole = head or heights[i] <= 0.9 * room
            keep = [E("w:cantSplit")] if whole else []
            # шапка и первая строка тела (без шапки — первая строка) держатся со следующей строкой:
            # начало таблицы не остаётся внизу страницы одной строкой; предпоследняя — с последней.
            # Строку без cantSplit разрешено рвать — её не держим, а от следующей такой в цепочку идёт
            # только начало. Цепочка держится, только если влезает: от начала таблицы — в ¾ страницы
            # (над ней заголовок, а не влезшую с ним цепочку LibreOffice уносит, оставляя заголовок
            # на странице одного), хвост — в страницу (шапка на страницах не повторяется)
            end = i + 2 if i < last and heights[i + 1] <= 0.9 * room else i + 1
            cap = 0.75 * TEXT_H[orient] if first == 0 else TEXT_H[orient]
            hold = whole and i < last and (i < (2 if headed else 1) or i == last - 1) \
                and sum(heights[first:end]) <= cap
            first = first if hold else i + 1
            tr = E("w:tr", None, E("w:trPr", None, *keep))  # без tblHeader: шапка не повторяется (§3.4)
            for j, pieces in enumerate(row):
                tcpr = E("w:tcPr", None, E("w:tcW", {"w:w": widths[j], "w:type": "dxa"}),
                         *([E("w:shd", SHADE)] if head else []))
                ppr = [spacing(0, 0, 240), E("w:jc", {"w:val": aligns[j]})]
                p = make_p(([E("w:keepNext")] if hold else []) + ppr)
                self.inline(p, pieces, bold=head, max_w=widths[j] - CELL_PAD)
                tr.append(E("w:tc", None, tcpr, p))
            tbl.append(tr)
        return tbl

    # --- inline (§3.6, §3.7)

    def inline(self, p, pieces, bold=False, max_w=TEXT_W["portrait"]):
        parent = p
        for kind, val, fmt in coalesce(pieces, bold):
            if kind == "text":
                if "<!--" in val and "code" not in fmt:  # markdown-it оставил его текстом
                    self.warn(tr("unclosed inline HTML comment — shown as text"))
                # символ без глифа в HN, но с глифом в Arial — своим прогоном Arial; код не трогаем
                for arial, part in itertools.groupby(val, lambda ch: "code" not in fmt and arial_only(ch)):
                    parent.append(make_run("".join(part), fmt | {"arial"} if arial else fmt))
            elif kind == "br":  # свойства — как у текста вокруг (жирность шапки и т. п.)
                rpr = run_props(fmt)
                parent.append(E("w:r", None, *([rpr] if rpr is not None else []), E("w:br")))
            elif kind == "link":
                rid = self.doc.part.relate_to(val, RT.HYPERLINK, is_external=True)
                parent = E("w:hyperlink", {"r:id": rid, "w:history": 1})
                p.append(parent)
            elif kind == "endlink":
                parent = p
            elif kind == "badlink":
                shown = unquote(val)
                shown = shown if len(shown) <= 60 else shown[:60] + "…"
                self.warn(tr("link is not http(s), mailto or tel: {link} — shown as text without a hyperlink",
                             link=shown))
            elif kind == "image":
                parent.extend(self.image(val, fmt, max_w))

    def image(self, node, fmt, max_w):
        src = str(node.attrs.get("src") or "")
        alt = one_line(plain(flatten(node.children)))
        # в предупреждениях — имя без %-кодировки, data: URI — коротко
        shown = src[:40] + "…" if src.lower().startswith("data:") and len(src) > 40 else unquote(src)
        path = self.local_file(src)
        if path is None:
            self.warn(tr("image not found or remote: {image} — a placeholder is inserted", image=shown))
        elif path.stat().st_size > MAX_IMAGE_BYTES:
            self.warn(tr("image larger than 50 MB: {image} — a placeholder is inserted", image=shown))
        else:
            if not path.resolve().is_relative_to(self.md_path.parent.resolve()):
                self.warn(tr("image outside the document's folder: {image}", image=shown))
            try:
                info = Image.from_file(str(path))
                width, height = info.px_width * 914400 // 96, info.px_height * 914400 // 96  # 96 dpi
                limit = max(max_w, 1) * 635  # twips → EMU
                if width > limit:
                    width, height = limit, height * limit // width
                orient = "landscape" if self.width == TEXT_W["landscape"] else "portrait"
                limit = TEXT_H[orient] * 8 // 10 * 635  # и не выше 0,8 области текста раздела
                if height > limit:
                    width, height = max(width * limit // height, 1), limit
                inline = self.doc.part.new_pic_inline(str(path), width, height)
                # python-docx берёт id по уже вставленным в тело рисункам, а абзац ещё не в теле:
                # две картинки одного абзаца получили бы одинаковый id, Word это не любит
                self.pictures += 1
                inline.docPr.set("id", str(self.pictures))
                inline.docPr.set("name", f"Picture {self.pictures}")
                if alt:
                    inline.docPr.set("descr", alt)
                return [E("w:r", None, E("w:drawing", None, inline))]
            except Exception as exc:  # битый файл: python-docx бросает и голый Exception
                self.warn(tr("image {image} not inserted ({error}) — a placeholder is inserted",
                             image=shown, error=exc.__class__.__name__))
        # format, а не tr(…, alt=): изолирующие символы страницы справа налево — не текст документа;
        # буквы без глифа в HN (арабская заглушка) — прогоном Arial, как в тексте
        text, fmt = tr("[image: {alt}]").format(alt=alt), fmt | {"i", "gray"}
        return [make_run("".join(part), fmt | {"arial"} if arial else fmt)
                for arial, part in itertools.groupby(text, arial_only)]

    def local_file(self, src):
        if re.match(r"[A-Za-z][A-Za-z0-9+.-]+:", src):  # http:, https:, data:, file:
            if not src.lower().startswith("file:"):
                return None
            src = urlparse(src).path
        path = Path(unquote(src))
        if not path.is_absolute():
            path = self.md_path.parent / path
        return path if path.is_file() else None


# ---------------------------------------------------------------- точка входа

NOT_TEXT = "looks like UTF-16 without BOM or a binary file — save it as UTF-8"


def overwritable(path):
    """Наш docx, не правленный после конвертации: метка, revision 1, пустой lastModifiedBy
    (Word при сохранении меняет оба). Не docx, файл 0.1.0/0.1.1 без метки — нет."""
    try:
        with zipfile.ZipFile(path) as z:
            core = parse_xml(z.read("docProps/core.xml"))
    except Exception:  # не zip, нет core.xml, каталог и т. п.
        return False
    return (core.findtext(qn("dc:identifier")) == MARK and core.findtext(qn("cp:revision")) == "1"
            and not core.findtext(qn("cp:lastModifiedBy")))


def refuse_if_open_in_word(path):
    """Рядом замок Word — файл открыт в Word: не заменять (§1). Имя замка (замер в Word для Mac
    16.114): основа — имя без последнего расширения, n — её длина в символах NFC; n ≤ 6 — «~$» +
    основа, n = 7 — без первого символа, n ≥ 8 — без двух первых; расширение то же. Word пишет
    имя замка в NFD, поэтому имена каталога сравниваются после приведения к NFC; регистр не
    различается, как в APFS. У символьной ссылки замок ищется и рядом с файлом, на который она
    указывает: заменяется он."""
    if path.is_symlink():
        refuse_if_open_in_word(path.resolve())
    stem = unicodedata.normalize("NFC", path.stem)
    n = len(stem)
    lock = "~$" + stem[0 if n <= 6 else 1 if n == 7 else 2:] + unicodedata.normalize("NFC", path.suffix)
    try:
        names = os.listdir(path.parent)
    except OSError:  # каталог без права чтения: спросить файловую систему по имени
        names = [lock] if (path.parent / lock).exists() else []
    found = next((name for name in names if unicodedata.normalize("NFC", name).casefold() == lock.casefold()),
                 None)
    if found:
        raise FileExistsError(tr("{path}: the file is open in Word ({lock} is next to it) — close it "
                                 "or run with --force", path=path, lock=found))


def convert(md_path, out_path, *, landscape=False, auto_landscape=True, force=False) -> dict:
    """Сконвертировать один файл; вернуть статистику для строки итога."""
    md_path, out_path = Path(md_path), Path(out_path)
    raw = md_path.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):  # BOM UTF-16 LE/BE
        text = raw.decode("utf-16")
    elif b"\x00" in raw and not raw.startswith(b"\xef\xbb\xbf"):  # UTF-16 без BOM «читается» как UTF-8
        pos = raw.index(b"\x00")
        raise UnicodeDecodeError("utf-8", raw, pos, pos + 1, NOT_TEXT)
    else:
        text = raw.decode("utf-8-sig")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x0b", " ").replace("\x0c", " ")  # VT и FF разделяют слова — пробел, а не склейка
    text = INVALID_XML.sub("", text)  # иначе lxml падает на заголовке (title) и alt картинки
    text, breaks = split_table_breaks(strip_front_matter(text))
    tokens = MD.parse(text)
    root = SyntaxTreeNode(tokens)
    conv = Converter(md_path, landscape, auto_landscape, breaks)
    deep = next((k for k, t in enumerate(tokens)
                 if t.type in NEST_OPEN and t.level >= MD.options.maxNesting - 1), None)
    if deep is not None:  # разборщик молча отбросил содержимое контейнера (в списке — и хвост документа)
        depth = sum(t.nesting for t in tokens[:deep + 1]
                    if t.type.startswith(("blockquote", "bullet_list", "ordered_list")))
        conv.line = tokens[deep].map[0] + 1 if tokens[deep].map else None
        conv.warn(tr("content nested deeper than {depth} levels is skipped", depth=depth - 1))
    sections, auto_tables = conv.run(root)
    if out_path.is_dir():  # иначе отказ ниже отправил бы к --force, который тоже не поможет
        raise IsADirectoryError(errno.EISDIR, os.strerror(errno.EISDIR), str(out_path))
    if out_path.exists() and not force:
        refuse_if_open_in_word(out_path)
    if out_path.exists() and not force and not overwritable(out_path):
        raise FileExistsError(tr("{path} already exists and was edited after conversion or not created "
                                 "by md2docx — not overwriting; add --force to overwrite", path=out_path))
    # символьная ссылка остаётся ссылкой — заменяется файл, на который она указывает (как в 0.1.1)
    dest = out_path.resolve() if out_path.is_symlink() else out_path
    if dest.is_file() and not os.access(dest, os.W_OK):  # замена файла молча обошла бы защиту от записи
        raise PermissionError(errno.EACCES, os.strerror(errno.EACCES), str(out_path))
    tmp = dest.parent / f".{dest.name}.{os.getpid()}.tmp"  # файл результата — целиком или никак
    try:
        conv.doc.save(str(tmp))
        if dest.exists():  # права прежнего результата, а не по umask (600 не становится 644)
            shutil.copymode(dest, tmp)
        tmp.replace(dest)
    except OSError as exc:
        # в сообщении — файл результата, а не временный; нет доступа — значит, временный файл
        # не создать в каталоге (сам результат проверен выше): называем каталог
        exc.filename = str(dest.absolute().parent if isinstance(exc, PermissionError) else out_path)
        raise
    finally:
        tmp.unlink(missing_ok=True)
    return {
        "output": str(out_path),
        "sections": [LETTER[orient] for orient, _ in sections],
        "tables": sum(1 for n in walk(root) if n.type == "table"),
        "auto_landscape_tables": auto_tables,
        "warnings": conv.warnings,
    }


class DocxError(ValueError):
    """Не docx или слишком велик. Аргумент — функция, которая возвращает текст: он переводится
    при показе, а не при создании (браузер хранит ошибку .docx и показывает её на языке каждого
    следующего запроса)."""

    def __str__(self):
        return self.args[0]()


def fix_fonts(path, *, force=False) -> dict:
    """Вернуть в fontTable готового docx подсказки FONT_ENTRIES, которые Word для Mac стирает
    при сохранении (§3.2); → {шрифт: [восстановленные поля]}, пустой — файл не тронут."""
    path = Path(path)
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            # docx-бомба: размеры проверяются по заявленным, а читается не дальше заявленного —
            # read() без предела распаковывает до гигабайта за вызов, bzip2 и lzma — без предела вовсе
            if any(i.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED) for i in infos):
                raise DocxError(lambda: tr("not a docx file or the file is damaged"))  # Word сжимает только deflate
            if sum(i.file_size for i in infos) > UNPACKED_LIMIT:
                raise DocxError(lambda: tr("more than {size} MB unpacked — too large for a docx",
                                           size=UNPACKED_LIMIT >> 20))
            if sum(i.file_size for i in infos if i.filename in MARKUP_PARTS) > MARKUP_LIMIT:
                raise DocxError(lambda: tr("text and styles take more than {size} MB unpacked — too large for a docx",
                                           size=MARKUP_LIMIT >> 20))
            items = []
            for info in infos:
                with z.open(info) as f:
                    items.append((info, f.read(info.file_size)))
    except (zipfile.BadZipFile, zlib.error,  # не zip, обрезан, испорчено сжатие
            RuntimeError) as exc:  # член зашифрован или сжат неизвестным методом
        raise DocxError(lambda: tr("not a docx file or the file is damaged")) from exc
    parts = {info.filename: data for info, data in items}
    if FONT_TABLE not in parts:
        raise DocxError(lambda: tr("the file has no {part} — not a docx", part=FONT_TABLE))
    try:
        fonts = parse_xml(parts[FONT_TABLE])
        used = {rf.get(qn("w:" + attr))
                for part in ("word/document.xml", "word/styles.xml", "word/numbering.xml") if part in parts
                for rf in parse_xml(parts[part]).iter(qn("w:rFonts"))
                for attr in ("ascii", "hAnsi", "cs", "eastAsia")}
    except etree.XMLSyntaxError as exc:  # архив цел, а XML части оборван или испорчен
        raise DocxError(lambda: tr("not a docx file or the file is damaged")) from exc
    changed = {}
    for name, props in FONT_ENTRIES:
        font = next((f for f in fonts.findall(qn("w:font")) if f.get(qn("w:name")) == name), None)
        if font is None:
            if name not in used:
                continue
            font = E("w:font", {"w:name": name})  # шрифт в тексте есть, а записи о нём нет
            fonts.append(font)
        for tag, attrs in props:  # прочие поля записи (notTrueType, embed*) не трогаем
            want = E("w:" + tag, attrs)
            have = font.find(want.tag)
            if have is None or dict(have.attrib) != dict(want.attrib):
                insert_ordered(font, want, FONT_ORDER)
                changed.setdefault(name, []).append(tag)
    if not changed:
        return changed
    if not force:
        refuse_if_open_in_word(path)
    # дальше — как запись в convert(): ссылка остаётся ссылкой, права прежние, целиком или никак
    dest = path.resolve() if path.is_symlink() else path
    if not os.access(dest, os.W_OK):
        raise PermissionError(errno.EACCES, os.strerror(errno.EACCES), str(path))
    table = serialize_part_xml(fonts)
    tmp = dest.parent / f".{dest.name}.{os.getpid()}.tmp"
    try:
        with zipfile.ZipFile(tmp, "w") as z:  # прочие члены — те же байты, в том же порядке
            for info, data in items:
                z.writestr(info, table if info.filename == FONT_TABLE else data)
        shutil.copymode(dest, tmp)
        tmp.replace(dest)
    except OSError as exc:
        exc.filename = str(dest.absolute().parent if isinstance(exc, PermissionError) else path)
        raise
    finally:
        tmp.unlink(missing_ok=True)
    return changed


def describe_error(exc, src):
    name = str(getattr(exc, "filename", None) or "")
    where = f": {name}" if name and name != str(src) else ""  # путь называем, если это не сам исходник
    if isinstance(exc, FileNotFoundError):
        return tr("file or folder not found") + where
    if isinstance(exc, IsADirectoryError):
        return tr("this is a folder, not a file") + where
    if isinstance(exc, PermissionError):
        return tr("permission denied") + where
    if isinstance(exc, UnicodeDecodeError):
        return tr(NOT_TEXT) if exc.reason == NOT_TEXT else tr("the file is not UTF-8 encoded")
    if getattr(exc, "errno", None):  # нет места, только чтение и т. п. — текст ОС английский
        return tr("I/O error (errno {errno}: {message})", errno=exc.errno, message=exc.strerror) + where
    return str(exc)


def internal_error(exc):
    """Непредвиденное исключение — одной строкой для пользователя."""
    return tr("internal error ({error})", error=f"{exc.__class__.__name__}: {exc}")


class Parser(argparse.ArgumentParser):
    """argparse с сообщениями на языке интерфейса: своего перевода у argparse нет."""
    MESSAGES = ("the following arguments are required", "unrecognized arguments", "expected one argument",
                "ambiguous option", "ignored explicit argument")

    def format_usage(self):
        return super().format_usage().replace("usage:", tr("usage:"), 1)

    def format_help(self):
        return super().format_help().replace("usage:", tr("usage:"), 1)

    def error(self, message):
        for en in self.MESSAGES:
            message = message.replace(en, tr(en))
        if message.startswith("argument "):  # «argument --port: …»
            message = tr("argument") + message[len("argument"):]
        self.print_usage(sys.stderr)
        self.exit(2, tr("{prog}: error: {message}", prog=self.prog, message=message) + "\n")


def main(argv=None):
    parser = Parser(
        prog="md2docx", add_help=False,
        description=tr("Markdown → DOCX converter: headings, tables, portrait and landscape pages."))
    parser._positionals.title, parser._optionals.title = tr("positional arguments"), tr("options")
    parser.add_argument("-h", "--help", action="help", help=tr("show this help message and exit"))
    parser.add_argument("inputs", nargs="+", metavar="INPUT.md",
                        help=tr("source Markdown files; with --fix-fonts, docx files"))
    parser.add_argument("-o", dest="output", metavar="OUT.docx",
                        help=tr("output file (only with a single input file); by default, next to the source"))
    parser.add_argument("--landscape", action="store_true", help=tr("make landscape the default orientation"))
    parser.add_argument("--no-auto-landscape", action="store_true",
                        help=tr("do not move wide tables to landscape pages"))
    parser.add_argument("--force", action="store_true",
                        help=tr("overwrite the output file even if it was edited after conversion, was not "
                                "created by md2docx or is open in Word"))
    parser.add_argument("--fix-fonts", action="store_true",
                        help=tr("instead of converting, restore the Windows font hints that Word for Mac "
                                "strips from finished docx files"))
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}",
                        help=tr("show the version and exit"))
    args = parser.parse_args(argv)
    if args.fix_fonts and (args.output is not None or args.landscape or args.no_auto_landscape):
        parser.error(tr("--fix-fonts cannot be combined with -o, --landscape or --no-auto-landscape"))
    if args.fix_fonts:
        code = 0
        for src in args.inputs:
            try:
                changed = fix_fonts(src, force=args.force)
            except FileExistsError as exc:  # открыт в Word: путь уже в тексте
                print(tr("Error: {message}", message=exc), file=sys.stderr)
                code = 1
                continue
            except (OSError, ValueError) as exc:
                print(tr("Error: {file}: {message}", file=src, message=describe_error(exc, src)), file=sys.stderr)
                code = 1
                continue
            except Exception as exc:  # непредвиденное — сообщить и перейти к следующему файлу пакета
                print(tr("Error: {file}: {message}", file=src, message=internal_error(exc)), file=sys.stderr)
                code = 1
                continue
            if changed:
                print(tr("Done: {file} — {summary}", file=src, summary=fonts_summary(changed)))
            else:
                print(tr("No changes: {file} — the font hints are already in place", file=src))
        return code
    if args.output and len(args.inputs) > 1:
        parser.error(tr("-o is allowed only with a single input file"))
    if args.output == "":
        parser.error(tr("empty path in -o"))

    code = 0
    written = []  # (результат, вход) — уже записанные в этом запуске
    for src in args.inputs:
        md_path = Path(src)
        try:
            # не with_suffix: на «.» и «/» он бросает ValueError раньше, чем чтение скажет «каталог»
            out_path = Path(args.output) if args.output else md_path.parent / (md_path.stem + ".docx")
            if out_path.exists() and out_path.samefile(md_path):  # -o на сам исходник
                print(tr("Error: {file}: {message}", file=src, message=tr(
                    "the output file is the source itself — the source is left untouched")), file=sys.stderr)
                code = 1
                continue
            first = next((s for p, s in written if out_path.exists() and out_path.samefile(p)), None)
            if first is not None:  # notes.md и notes.markdown, один вход дважды: поздний затёр бы ранний
                print(tr("Error: {file}: {message}", file=src, message=tr(
                    "output {output} is the same as the output of {first} — skipped", output=out_path,
                    first=first)), file=sys.stderr)
                code = 1
                continue
            stats = convert(md_path, out_path, landscape=args.landscape,
                            auto_landscape=not args.no_auto_landscape, force=args.force)
            written.append((out_path, src))
        except (OSError, UnicodeDecodeError) as exc:
            print(tr("Error: {file}: {message}", file=src, message=describe_error(exc, src)), file=sys.stderr)
            code = 1
            continue
        except Exception as exc:  # непредвиденное — сообщить и перейти к следующему файлу пакета
            print(tr("Error: {file}: {message}", file=src, message=internal_error(exc)), file=sys.stderr)
            code = 1
            continue
        for warning in stats["warnings"]:
            print(tr("Warning: {message}", message=warning), file=sys.stderr)
        print(tr("Done: {file} — {summary}", file=stats["output"], summary=summary(stats)))
    return code


if __name__ == "__main__":
    sys.exit(main())
